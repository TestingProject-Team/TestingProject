# -*- coding: utf-8 -*-
import docx
import sys
import io

if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

def inspect_checklist_table(file_path):
    doc = docx.Document(file_path)
    print("=== Table 94 (Checklist Table) ===")
    table = doc.tables[94]
    for r_idx, row in enumerate(table.rows):
        vals = [c.text.strip().replace('\n', ' ') for c in row.cells]
        print(f"Row {r_idx}: {' | '.join(vals)}")

    print("\n=== Table 95 (Work Allocation Table) ===")
    table = doc.tables[95]
    for r_idx, row in enumerate(table.rows):
        vals = [c.text.strip().replace('\n', ' ') for c in row.cells]
        print(f"Row {r_idx}: {' | '.join(vals)}")

if __name__ == "__main__":
    inspect_checklist_table("YiYi_Book_Capstone_Project_Report_Tieng_Viet_Updated.docx")
