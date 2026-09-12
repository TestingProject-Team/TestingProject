import sys
sys.path.append(r'E:\TestingProject\scripts')
from optimize_docx_helper import optimize_docx
import docx, os

target = r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Content_Final.docx'
temp = target + '.tmp'
optimize_docx(target, temp)
doc = docx.Document(temp)
print(f"Validation successful! Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")
os.replace(temp, target)
print("Optimized YiYi_Book_Capstone_Project_Report_Content_Final.docx successfully.")
