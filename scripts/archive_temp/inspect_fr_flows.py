import docx

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

with open(r'E:\TestingProject\chap3_fr_flows.txt', 'w', encoding='utf-8') as f:
    f.write("=== TABLE 29: FUNCTIONAL REQUIREMENTS ===\n")
    t29 = doc.tables[29]
    for r_idx, r in enumerate(t29.rows):
        f.write(f"Row {r_idx}: {[c.text.strip().replace('\n', ' ') for c in r.cells]}\n")
    
    f.write("\n=== FLOW PARAGRAPHS IN CHAPTER III ===\n")
    for idx in range(180, 360):
        t = doc.paragraphs[idx].text.strip()
        if any(t.startswith(x) for x in ['FR-', 'Mục tiêu:', 'Actor(s):', 'Luồng chính', 'Luồng thay thế', 'Tiêu chí nghiệm thu', 'Hình 3.']):
            f.write(f"[P{idx}] {t}\n")

print("Dumped FRs and flows.")
