# -*- coding: utf-8 -*-
import os
import sys
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

DOCX_FILES = [
    "YiYi_Book_Capstone_Project_Report_Tieng_Viet_Updated.docx",
    "YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full.docx"
]

IMG_31 = "Hinh_3.1_System_Context_Diagram_YiYi_Book.png"
IMG_32 = "Hinh_3.2_Overall_Use_Case_Diagram_YiYi_Book.png"

for doc_path in DOCX_FILES:
    if not os.path.exists(doc_path):
        continue
    try:
        print(f"Processing {doc_path}...")
        doc = docx.Document(doc_path)
        body = doc._body._element
        elements = list(body)
        
        # 1. Tìm vị trí chèn Hình 3.1
        found_31 = False
        for idx, elem in enumerate(elements):
            tag = elem.tag.split("}")[-1]
            text = "".join(elem.itertext())
            if tag == "tbl" and "Sơ đồ System Context" in text and "System Context Diagram" in text:
                # Kiểm tra paragraph sau
                if idx + 1 < len(elements) and "Hình 3.1" in "".join(elements[idx+1].itertext()):
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(8)
                    p_img.paragraph_format.space_after = Pt(4)
                    run_img = p_img.add_run()
                    run_img.add_picture(IMG_31, width=Inches(6.4))
                    
                    elem.addprevious(p_img._element)
                    body.remove(elem)
                    found_31 = True
                    print("  Replaced placeholder for Hình 3.1 successfully.")
                    break
                    
        # 2. Tìm vị trí chèn Hình 3.2
        elements = list(body) # Cập nhật lại danh sách elements
        found_32 = False
        for idx, elem in enumerate(elements):
            tag = elem.tag.split("}")[-1]
            text = "".join(elem.itertext())
            if tag == "tbl" and "Sơ đồ Use Case tổng quát" in text and "Overall Use Case Diagram" in text:
                if idx + 1 < len(elements) and "Hình 3.2" in "".join(elements[idx+1].itertext()):
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(8)
                    p_img.paragraph_format.space_after = Pt(4)
                    run_img = p_img.add_run()
                    run_img.add_picture(IMG_32, width=Inches(6.4))
                    
                    elem.addprevious(p_img._element)
                    body.remove(elem)
                    found_32 = True
                    print("  Replaced placeholder for Hình 3.2 successfully.")
                    break

        # Chuẩn hóa caption cho cả 2 hình
        for p in doc.paragraphs:
            p_text = p.text.strip()
            if p_text.startswith("Hình 3.1.") or p_text.startswith("Hình 3.2."):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(8)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)
                    run.font.italic = True
                    run.font.color.rgb = RGBColor(68, 68, 68)

        if found_31 or found_32:
            doc.save(doc_path)
            print(f"Saved {doc_path} successfully.")
    except Exception as e:
        print(f"Could not update {doc_path}: {e}")
