import docx

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

# 1. Update Paragraphs
for idx, p in enumerate(doc.paragraphs):
    # Benchmark in Chap 1 (P80)
    if 'Nhóm nên benchmark YiYi Book với trải nghiệm' in p.text:
        p.text = "YiYi Book được đối chiếu với các hệ thống thương mại điện tử và nhà sách trực tuyến hiện có nhằm đánh giá cấu trúc danh mục, tìm kiếm sách, quy trình thanh toán, chính sách khuyến mãi và tính năng quản trị. Thông tin benchmark chi tiết sẽ được bổ sung sau khi nhóm hoàn tất quá trình đối chiếu thực tế."
    
    # Limitations in Chap 1 (P102)
    if 'Archive được cung cấp không có tên sinh viên chính thức' in p.text:
        p.text = "Dự án đã chuẩn hóa đầy đủ thông tin nhóm gồm 5 thành viên, Giảng viên hướng dẫn Nguyễn Văn Chiến và cấu trúc chức năng cốt lõi. Tuy nhiên, các biên bản nghiệm thu người dùng thực tế (UAT) và bằng chứng triển khai trên môi trường Production công khai vẫn cần được nhóm tiếp tục hoàn thiện trước khi phát hành tài liệu chính thức."

# 2. Update Table 3
t3 = doc.tables[3]
for row in t3.rows:
    for cell in row.cells:
        if 'Minh chứng từ archive được cung cấp' in cell.text:
            cell.text = cell.text.replace('Minh chứng từ archive được cung cấp', 'Minh chứng thành phần hệ thống triển khai')
        if 'archive được cung cấp' in cell.text:
            cell.text = cell.text.replace('archive được cung cấp', 'mã nguồn và thành phần hệ thống')

# 3. Update Table 6 (Benchmark)
t6 = doc.tables[6]
for row in t6.rows:
    for cell in row.cells:
        if '[CẦN BỔ SUNG EVIDENCE: Tên hệ thống, URL và ảnh chụp so sánh]' in cell.text:
            cell.text = "[CẦN BỔ SUNG: Benchmark system / URL / Evidence]"
        if '[CẦN BỔ SUNG EVIDENCE: Minh chứng luồng thanh toán Sandbox]' in cell.text:
            cell.text = "[CẦN BỔ SUNG: Benchmark payment gateway / Evidence]"
        if '[CẦN BỔ SUNG EVIDENCE: Bảng ghi chú đánh giá Benchmark]' in cell.text:
            cell.text = "[CẦN BỔ SUNG: Benchmark AI assistant / Evidence]"

doc.save(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')
print("Chapter 1 normalized successfully.")
