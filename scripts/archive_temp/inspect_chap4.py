import docx

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

with open(r'E:\TestingProject\chap4_inspection.txt', 'w', encoding='utf-8') as f:
    f.write("=== PARAGRAPHS IN CHAPTER IV ===\n")
    for idx in range(377, 592):
        t = doc.paragraphs[idx].text.strip()
        if doc.paragraphs[idx].style.name.startswith('Heading') or t.startswith('Hình 4.') or 'table' in t.lower() or 'entity' in t.lower() or 'bảng' in t.lower() or 'port' in t.lower() or '808' in t:
            f.write(f"[P{idx}][{doc.paragraphs[idx].style.name}] {t}\n")

print("Dumped Chap 4 inspection.")
