import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
docx_path = "YiYi_Book_Capstone_Project_Report_Tieng_Viet (1).docx"
doc = docx.Document(docx_path)

print("Paragraph count:", len(doc.paragraphs))
print("Table count:", len(doc.tables))

# Print Chapter IV, V, VI paragraphs
for i in range(360, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    if p.text.strip():
        print(f"[{i}] ({p.style.name}): {p.text.strip()[:80]}")
