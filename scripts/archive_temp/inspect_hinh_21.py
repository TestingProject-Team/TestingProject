# -*- coding: utf-8 -*-
import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document("YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full.docx")
print(f"Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}")

found_h21 = False
for idx, p in enumerate(doc.paragraphs):
    if "Hình 2.1" in p.text:
        print(f"P[{idx}]: {p.text}")
        found_h21 = True
        # Print surrounding paragraphs
        for sub in range(max(0, idx-8), min(len(doc.paragraphs), idx+8)):
            print(f"  [{sub}] {doc.paragraphs[sub].text[:120]}")
        break

# Also inspect body elements around Section 2.1
print("\n--- Body elements search ---")
for b_idx, elem in enumerate(doc._body._element):
    tag = elem.tag.split("}")[-1]
    text = elem.text or ""
    if "Hình 2.1" in text or "2.1 Quy trình dự án" in text:
        print(f"Body[{b_idx}] <{tag}>: {text[:100]}")
        for neighbor in range(max(0, b_idx-3), min(len(doc._body._element), b_idx+6)):
            n_elem = doc._body._element[neighbor]
            n_tag = n_elem.tag.split("}")[-1]
            n_text = "".join(n_elem.itertext())
            print(f"   [{neighbor}] <{n_tag}>: {n_text[:100]}")
        break
