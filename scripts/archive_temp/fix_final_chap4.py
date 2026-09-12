import docx

doc_path = r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Content_Final.docx'
doc = docx.Document(doc_path)

# Chapter IV Detailed Design Placeholders
# Ensure no merged "Class/Sequence Diagram" text remains in placeholders or captions

diagram_placeholders = {
    55: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ thực thể quan hệ cơ sở dữ liệu ERD 21 bảng PostgreSQL]",
    50: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ kiến trúc phân tầng Backend Spring Boot]",
    51: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ luồng xử lý AI RAG và Phân loại Intent]",
    53: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ phân rã gói phần mềm hệ thống (Package Diagram)]",
    57: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ tuần tự xác thực JWT (Login Sequence Diagram)]",
    58: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ tuần tự xem danh sách sách và danh mục (Product List Sequence Diagram)]",
    59: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ tuần tự tìm kiếm và lọc sách nâng cao (Search & Filter Sequence Diagram)]",
    60: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ tuần tự xem chi tiết sách và đánh giá (Product Detail Sequence Diagram)]",
    61: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ tuần tự thao tác giỏ hàng (Cart Management Sequence Diagram)]",
    62: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ tuần tự đặt hàng COD và áp dụng mã giảm giá (Checkout Sequence Diagram)]",
    63: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ tuần tự tích hợp cổng thanh toán trực tuyến VNPAY (Online Payment Sequence Diagram)]",
    64: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ tuần tự xử lý đơn hàng và yêu cầu đổi trả (Order & Return Sequence Diagram)]",
    65: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ tuần tự tích lũy và sử dụng điểm Y-Point (Y-Point & Reward Sequence Diagram)]",
    66: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ tuần tự tương tác Trợ lý AI và Client Mini-RAG (YiYi AI Sequence Diagram)]",
    67: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ tuần tự quản trị danh mục và kho sách (Admin Catalogue Sequence Diagram)]",
    68: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ tuần tự xử lý đơn hàng quản trị (Admin Order Sequence Diagram)]",
    69: "[CẦN BỔ SUNG DIAGRAM: Sơ đồ tuần tự quản trị khuyến mãi và nội dung (Admin Promotion Sequence Diagram)]"
}

for t_idx, txt in diagram_placeholders.items():
    if t_idx < len(doc.tables):
        doc.tables[t_idx].rows[0].cells[0].text = txt

# Check captions in text for any "Class/Sequence" combined words
for p in doc.paragraphs:
    if 'Class/Sequence' in p.text:
        p.text = p.text.replace('Class/Sequence Diagram', 'Sequence Diagram').replace('Class/Sequence', 'Sequence')

# Database descriptions: Ensure clear distinction between:
# - 25 Java Entity/Enum files
# - 21 JPA Entities directly mapping to 21 PostgreSQL Tables
# - 4 Enums: AuthProvider, DiscountType, Role, ShippingStatus
for p in doc.paragraphs:
    if '25 entity' in p.text.lower() or '21 bảng' in p.text.lower() or '25 bảng' in p.text.lower():
        if 'Tầng dữ liệu (Persistence Layer)' in p.text:
            p.text = "Tầng dữ liệu (Persistence Layer): Bao gồm 19 Spring Data JPA repositories và 25 Java Entity/Enum classes (gồm 21 thực thể JPA trực tiếp ánh xạ sang 21 bảng quan hệ trong cơ sở dữ liệu PostgreSQL và 4 Enums định nghĩa miền giá trị nghiệp vụ: AuthProvider, DiscountType, Role, ShippingStatus) [CẦN XÁC MINH DATABASE SCHEMA: số Entity Java và số Table PostgreSQL thực tế]."
        elif 'Cơ sở dữ liệu được mô hình hóa' in p.text:
            p.text = "Cơ sở dữ liệu được mô hình hóa dưới dạng quan hệ trên hệ quản trị cơ sở dữ liệu PostgreSQL. Mã nguồn hệ thống bao gồm 25 Java Entity/Enum classes (gồm 21 thực thể JPA ánh xạ tương ứng 21 bảng quan hệ PostgreSQL và 4 Enum giá trị) [CẦN XÁC MINH DATABASE SCHEMA: số Entity Java và số Table PostgreSQL thực tế]. Sơ đồ thực thể quan hệ (ERD) thể hiện đầy đủ các khóa chính, khóa ngoại và mối liên kết bảng."

# Port descriptions in Chapter IV & Table 72
t72 = doc.tables[72]
t72.rows[1].cells[2].text = "8081 (Local Dev & Docker) / 8082 (Docker Test Profile) [CẦN XÁC MINH PORT TỪ application.properties / application.yml / Docker Compose]"
t72.rows[2].cells[2].text = "5432 (Local & Docker) / 5434 (Docker Test DB)"
t72.rows[2].cells[3].text = "Lưu trữ 21 bảng thực thể PostgreSQL (tương ứng 21 JPA Entities và 4 Enums) và xử lý giao dịch nguyên tử ACID"

doc.save(doc_path)
print("Chapter 4 completed successfully.")
