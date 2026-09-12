import docx

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

# 1. Clean paragraphs in Chapter VI
# Remove duplicate instructions: "Chèn screenshot cho 3.x.x..."
# Remove repeated Figure 6.23 in troubleshooting
for p in list(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith('Chèn screenshot cho 3.'):
        p.text = "" # blank out instruction paragraph
    if 'Archive được cung cấp có cấu trúc page nhưng chưa có bộ screenshot cuối' in txt:
        p.text = "Hướng dẫn sử dụng này được tổ chức chi tiết theo từng đối tượng người dùng (Khách hàng và Quản trị viên), minh họa quy trình thao tác chuẩn trên giao diện hệ thống."

# Check repeated "Hình 6.23. Hướng dẫn Kiểm duyệt Bình luận..." around P937
for idx in range(930, min(945, len(doc.paragraphs))):
    p = doc.paragraphs[idx]
    if p.text.strip() == 'Hình 6.23. Hướng dẫn Kiểm duyệt Bình luận, Đánh giá và Quản lý Liên hệ.' and doc.paragraphs[idx-1].text.strip() == '3.4 Hỗ trợ và Troubleshooting':
        print(f"Removing duplicate caption at P{idx}")
        p.text = ""

# 2. Clean Tables 102 - 117
# Table 102 to 116 are placeholders for 3.2.1 to 3.3.7
# Each must have EXACTLY ONE clean placeholder corresponding to its section

clean_workflow_placeholders = {
    102: "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Luồng Đăng ký & Đăng nhập Khách hàng]",
    103: "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Luồng Khám phá Trang chủ, Danh mục & Tìm kiếm Sách]",
    104: "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Chi tiết sản phẩm & Quản lý Danh sách yêu thích (Wishlist)]",
    105: "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Thao tác Giỏ hàng & Điền thông tin Đặt hàng]",
    106: "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Cổng thanh toán Trực tuyến & Lịch sử Theo dõi Đơn hàng]",
    107: "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Giao diện Viết đánh giá sản phẩm và Xem bình luận]",
    108: "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Tích điểm thưởng Y-Point & Áp dụng Kho Voucher]",
    109: "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Cửa sổ Chatbot YiYi AI Assistant]",
    110: "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Đăng nhập Quản trị & Bảng điều khiển Dashboard Doanh thu]",
    111: "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Giao diện Quản trị Kho Sách và Tồn kho]",
    112: "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Quản trị Danh mục Thể loại & Banner Marketing]",
    113: "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Quản lý Danh sách Đơn hàng & Xử lý Vận chuyển]",
    114: "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Quản lý Danh sách Người dùng & Phân quyền RBAC]",
    115: "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Quản lý Mã Coupon Khuyến mãi & Điểm Y-Point]",
    116: "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Kiểm duyệt Đánh giá Bình luận & Quản lý Thư Liên hệ]"
}

for t_idx, clean_text in clean_workflow_placeholders.items():
    if t_idx < len(doc.tables):
        table = doc.tables[t_idx]
        for row in table.rows:
            for cell in row.cells:
                cell.text = clean_text

# 3. Fix Table 117 (Troubleshooting) - Header had placeholder inside cell 0!
# Row 0: ['Vấn đề[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Kiểm duyệt Đánh giá & Hòm thư Liên hệ]', 'Nguyên nhân có thể', 'Cách xử lý']
t117 = doc.tables[117]
t117.rows[0].cells[0].text = "Vấn đề phát sinh"
t117.rows[0].cells[1].text = "Nguyên nhân có thể"
t117.rows[0].cells[2].text = "Cách xử lý / Khắc phục"

# Fix Table 99: Remove overlapping text in cell 0
t99 = doc.tables[99]
t99.rows[0].cells[0].text = "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình trang chủ với trạng thái đã đăng nhập tài khoản]"

doc.save(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')
print("Chapter VI cleaned up successfully.")
