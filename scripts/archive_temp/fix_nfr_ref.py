import docx

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

# 1. Update NFR Table 44
t44 = doc.tables[44]
# Row 0: ID | Thuộc tính | Yêu cầu | Minh chứng
for row in t44.rows[1:]:
    attr = row.cells[1].text.strip()
    if 'Performance' in attr:
        row.cells[2].text = "Target response time: [CẦN XÁC MINH / BENCHMARK THỰC TẾ] (Thời gian phản hồi trang catalogue và thao tác giỏ hàng trong điều kiện mạng bình thường)."
        row.cells[3].text = "[CẦN XÁC MINH / CẦN BỔ SUNG EVIDENCE]"
    elif 'Security' in attr:
        row.cells[3].text = "[CẦN XÁC MINH / CẦN BỔ SUNG EVIDENCE: Báo cáo quét lỗ hổng OWASP / Pentest]"
    elif 'Usability' in attr:
        row.cells[3].text = "[CẦN XÁC MINH / CẦN BỔ SUNG EVIDENCE: Đánh giá khảo sát trải nghiệm người dùng]"

# 2. Update Appendix A
# P939: Note checklist
for p in doc.paragraphs:
    if 'Các mục sau được để dưới dạng placeholder rõ ràng để hai thành viên' in p.text:
        p.text = "Ghi chú: Phụ lục checklist này chỉ dùng trong quá trình hoàn thiện tài liệu và phải được loại bỏ hoặc chuyển thành Evidence Index trước khi nộp bản cuối cùng."
    if 'Đề xuất bàn giao cho hai người' in p.text:
        p.text = "Kế hoạch hoàn thiện minh chứng đồ án"

# Table 118: Update UAT reference, 17 tables -> 21 tables
t118 = doc.tables[118]
for row in t118.rows:
    for cell in row.cells:
        if '17 bảng' in cell.text:
            cell.text = cell.text.replace('17 bảng', '21 bảng thực thể [CẦN XÁC MINH DATABASE SCHEMA]')
        if 'Mục 5.8' in cell.text:
            cell.text = cell.text.replace('Mục 5.8', 'Chương V (Mục 3.4 & Mục 12)')

# Table 119: Update instructions and references
t119 = doc.tables[119]
for row in t119.rows:
    for cell in row.cells:
        if 'Mục 5.8' in cell.text:
            cell.text = cell.text.replace('Mục 5.8', 'Chương V (Kế hoạch UAT)')

# Table 121: Remove instructional tone, update references
t121 = doc.tables[121]
t121.rows[0].cells[3].text = "Mức độ ưu tiên"
t121.rows[0].cells[4].text = "Thành viên thực hiện"
t121.rows[0].cells[5].text = "Hạng mục công việc chi tiết"
for row in t121.rows[1:]:
    for cell in row.cells:
        if 'Mục 5.8' in cell.text:
            cell.text = cell.text.replace('Mục 5.8', 'Chương V')
        if '17 bảng' in cell.text:
            cell.text = cell.text.replace('17 bảng', '21 bảng thực thể [CẦN XÁC MINH DATABASE SCHEMA]')

# 3. Add References Section (Tài liệu tham khảo) at the end of the document
heading_ref = doc.add_paragraph("TÀI LIỆU THAM KHẢO", style='Heading 1')
ref_items = [
    "1. Oracle Corporation, \"Java Platform, Standard Edition 17 Documentation\", [CẦN XÁC MINH LINK THAM KHẢO].",
    "2. VMware, Inc., \"Spring Boot Reference Documentation (v3.2.x)\", [CẦN XÁC MINH LINK THAM KHẢO].",
    "3. VMware, Inc., \"Spring Security Reference Documentation\", [CẦN XÁC MINH LINK THAM KHẢO].",
    "4. Meta Platforms, Inc., \"React Official Documentation (v18.x)\", [CẦN XÁC MINH LINK THAM KHẢO].",
    "5. The PostgreSQL Global Development Group, \"PostgreSQL 16 Documentation\", [CẦN XÁC MINH LINK THAM KHẢO].",
    "6. JUnit Team, \"JUnit 5 User Guide\", [CẦN XÁC MINH LINK THAM KHẢO].",
    "7. JaCoCo Team, \"JaCoCo Java Code Coverage Library Documentation\", [CẦN XÁC MINH LINK THAM KHẢO].",
    "8. Postman, Inc., \"Postman API Platform & Newman CLI Documentation\", [CẦN XÁC MINH LINK THAM KHẢO].",
    "9. CodeceptJS Team, \"CodeceptJS Modern End-to-End Testing Documentation\", [CẦN XÁC MINH LINK THAM KHẢO].",
    "10. Microsoft Corporation, \"Playwright Node.js Automation Documentation\", [CẦN XÁC MINH LINK THAM KHẢO].",
    "11. SonarSource S.A., \"SonarQube Code Quality & Security Documentation\", [CẦN XÁC MINH LINK THAM KHẢO].",
    "12. OWASP Foundation, \"OWASP Top 10 Web Application Security Risks\", [CẦN XÁC MINH LINK THAM KHẢO]."
]
for item in ref_items:
    doc.add_paragraph(item)

doc.save(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')
print("NFR, Appendix A, and References updated successfully.")
