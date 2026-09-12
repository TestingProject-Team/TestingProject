# -*- coding: utf-8 -*-
import docx
import sys
import io

if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

def verify_chap5_figures(file_path):
    print(f"=== VERIFYING {file_path} ===")
    doc = docx.Document(file_path)
    
    figures = []
    placeholders = []
    
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("Hình 5."):
            figures.append(t)
            
    for t in doc.tables:
        txt = "".join(t._tbl.itertext()).strip()
        if "[CẦN BỔ SUNG EVIDENCE:" in txt:
            placeholders.append(txt[:80])
            
    print(f"Total Figures in Chapter 5: {len(figures)}")
    for f in figures:
        print(f"  {f}")
        
    print(f"\nTotal Evidence Placeholders in Chapter 5: {len(placeholders)}")
    for idx, pl in enumerate(placeholders, 1):
        print(f"  {idx}. {pl}")

if __name__ == "__main__":
    verify_chap5_figures("YiYi_Book_Capstone_Project_Report_Tieng_Viet (1).docx")
