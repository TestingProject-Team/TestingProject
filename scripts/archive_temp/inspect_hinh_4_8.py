import os, sys
import docx

doc_path = r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Content_Final.docx'
doc = docx.Document(doc_path)

with open(r'E:\TestingProject\hinh_4_8_context.txt', 'w', encoding='utf-8') as out:
    for i, p in enumerate(doc.paragraphs):
        if 'Hình 4.8' in p.text:
            out.write(f"P{i}: {p.text}\n")
            for j in range(max(0, i-5), min(len(doc.paragraphs), i+6)):
                out.write(f"  [{j}] {doc.paragraphs[j].text}\n")

    for t_idx, t in enumerate(doc.tables):
        for r_idx, r in enumerate(t.rows):
            for c_idx, c in enumerate(r.cells):
                if 'PRODUCT LIST' in c.text or 'Hình 4.8' in c.text or '4.8' in c.text:
                    out.write(f"Table {t_idx} [R{r_idx}C{c_idx}]: {c.text}\n")
print("Context written successfully.")
