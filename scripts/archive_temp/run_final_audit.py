import docx

doc_path = r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Content_Final.docx'
doc = docx.Document(doc_path)

targets = [
    'Hãy', 'Nhóm nên', 'archive', 'source archive', '1. 1.', '2. 2.', '3. 3.', '4. 4.',
    'REQ-', 'UC-01 (Quy trình mua sách', 'Response Time trung bình dưới 2000ms',
    '100% Toàn thời gian', 'sẵn sàng phục vụ người dùng thực tế', 'Class/Sequence Diagram',
    'UCT-01', 'UCT-02', 'UCT-03', 'CẦN XÁC MINH', 'CẦN BỔ SUNG'
]

results = {t: [] for t in targets}

# Paragraphs
for idx, p in enumerate(doc.paragraphs):
    txt = p.text
    for t in targets:
        if t.lower() in txt.lower():
            results[t].append((f'P{idx}', txt[:100].strip()))

# Tables
for tidx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell_txt = cell.text
            for t in targets:
                if t.lower() in cell_txt.lower():
                    results[t].append((f'T{tidx}[R{r_idx}C{c_idx}]', cell_txt[:100].strip()))

with open(r'E:\TestingProject\final_verification_audit.txt', 'w', encoding='utf-8') as f:
    f.write("=== AUDIT VERIFICATION REPORT ===\n\n")
    for t, matches in results.items():
        f.write(f"Target: [{t}] -> {len(matches)} occurrences\n")
        for loc, snip in matches[:5]:
            f.write(f"  [{loc}] {snip}\n")
        if len(matches) > 5:
            f.write(f"  ... and {len(matches) - 5} more\n")
        f.write("\n")

print("Audit written to final_verification_audit.txt")
