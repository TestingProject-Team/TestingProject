import sys
sys.path.append(r'E:\TestingProject\scripts')
from optimize_docx_helper import optimize_docx
import docx, os

target = r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx'
temp = target + '.tmp'
optimize_docx(target, temp)
doc = docx.Document(temp)
print(f"Validated (1): Paragraphs={len(doc.paragraphs)}, Tables={len(doc.tables)}")
os.replace(temp, target)
print("Optimized and replaced (1) docx successfully.")
