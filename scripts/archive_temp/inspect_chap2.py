import docx

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

with open(r'E:\TestingProject\chap2_inspection.txt', 'w', encoding='utf-8') as f:
    f.write("=== PARAGRAPHS IN CHAPTER II ===\n")
    for idx in range(107, 162):
        f.write(f"[P{idx}][{doc.paragraphs[idx].style.name}] {doc.paragraphs[idx].text}\n")
    
    f.write("\n=== TABLES IN CHAPTER II ===\n")
    for t_idx in range(6, 25):
        if t_idx >= len(doc.tables):
            break
        table = doc.tables[t_idx]
        f.write(f"\n--- TABLE {t_idx} (rows={len(table.rows)}, cols={len(table.columns)}) ---\n")
        for r_idx in range(min(5, len(table.rows))):
            cells_text = [c.text.strip().replace('\n', ' ') for c in table.rows[r_idx].cells]
            f.write(f"Row {r_idx}: {cells_text}\n")

print("Dumped Chap 2 inspection.")
