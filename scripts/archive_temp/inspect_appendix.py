import docx

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

with open(r'E:\TestingProject\appendix_inspection.txt', 'w', encoding='utf-8') as f:
    for idx in range(938, len(doc.paragraphs)):
        p = doc.paragraphs[idx]
        f.write(f"[P{idx}][{p.style.name}] {p.text}\n")
    
    for t_idx in range(118, len(doc.tables)):
        table = doc.tables[t_idx]
        f.write(f"\n=== TABLE {t_idx} (rows={len(table.rows)}, cols={len(table.columns)}) ===\n")
        for r_idx, row in enumerate(table.rows):
            f.write(f"Row {r_idx}: {[c.text.strip().replace('\n', ' ') for c in row.cells]}\n")

print("Dumped Appendix.")
