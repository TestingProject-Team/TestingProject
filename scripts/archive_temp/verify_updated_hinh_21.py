# -*- coding: utf-8 -*-
import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document("YiYi_Book_Capstone_Project_Report_Tieng_Viet_Updated.docx")
print(f"Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}")

for idx in range(135, min(155, len(doc._body._element))):
    elem = doc._body._element[idx]
    tag = elem.tag.split("}")[-1]
    text = "".join(elem.itertext())
    print(f"[{idx}] <{tag}>: {text[:100]}")
