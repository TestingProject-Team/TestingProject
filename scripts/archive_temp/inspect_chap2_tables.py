import docx

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

with open(r'E:\TestingProject\chap2_tables_detail.txt', 'w', encoding='utf-8') as f:
    for t_idx in range(15, 26):
        if t_idx >= len(doc.tables):
            break
        table = doc.tables[t_idx]
        f.write(f"\n=================== TABLE {t_idx} (rows={len(table.rows)}, cols={len(table.columns)}) ===================\n")
        for r_idx, row in enumerate(table.rows):
            cells_text = [c.text.strip().replace('\n', ' ') for c in row.cells]
            f.write(f"Row {r_idx}: {cells_text}\n")

print("Dumped tables 15-25.")
