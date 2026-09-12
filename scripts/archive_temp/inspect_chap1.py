import docx

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

with open(r'E:\TestingProject\chap1_inspection.txt', 'w', encoding='utf-8') as f:
    f.write("=== PARAGRAPHS IN CHAPTER I ===\n")
    for idx in range(60, 108):
        f.write(f"[P{idx}][{doc.paragraphs[idx].style.name}] {doc.paragraphs[idx].text}\n")
    
    f.write("\n=== TABLES IN CHAPTER I ===\n")
    # check tables 0 to 5
    for t_idx in range(0, 6):
        f.write(f"\n--- TABLE {t_idx} ---\n")
        table = doc.tables[t_idx]
        for r_idx, row in enumerate(table.rows):
            cells_text = [c.text.strip().replace('\n', ' ') for c in row.cells]
            f.write(f"Row {r_idx}: {cells_text}\n")

print("Dumped Chap 1 inspection.")
