# -*- coding: utf-8 -*-
"""
refine_all_chapters_comprehensive.py
Thực hiện rà soát và chuẩn hóa toàn bộ các placeholder hình ảnh, sơ đồ UML, screenshot và evidence
từ Chương I đến Chương VI theo đúng quy chuẩn báo cáo PDCMS.
"""

import os
import sys
import io
import shutil
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.table import Table
from xml.sax.saxutils import escape as escape_xml

if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

def set_p_text(p_elem, text):
    """Xóa các run cũ và thêm text mới cho một CT_P xml element."""
    for r in p_elem.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r'):
        p_elem.remove(r)
    esc_text = escape_xml(text)
    r_elem = parse_xml(
        f'<w:r {nsdecls("w")}>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:b/><w:i/>'
        f'<w:sz w:val="20"/>'
        f'<w:color w:val="1E293B"/>'
        f'</w:rPr>'
        f'<w:t>{esc_text}</w:t>'
        f'</w:r>'
    )
    p_elem.append(r_elem)

def format_placeholder_table(table, placeholder_text):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:bottom w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:left w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:right w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    cell = table.rows[0].cells[0]
    cell.width = Inches(6.2)
    
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
    tcPr.append(shd)
    
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="360" w:type="dxa"/><w:bottom w:w="360" w:type="dxa"/><w:left w:w="240" w:type="dxa"/><w:right w:w="240" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.2
    
    run = p.add_run(placeholder_text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(71, 85, 105)

def make_caption_p(caption_text):
    esc_cap = escape_xml(caption_text)
    cap_elem = parse_xml(
        f'<w:p {nsdecls("w")}>'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="120" w:after="240"/>'
        f'</w:pPr>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:b/><w:i/>'
        f'<w:sz w:val="20"/>'
        f'<w:color w:val="1E293B"/>'
        f'</w:rPr>'
        f'<w:t>{esc_cap}</w:t>'
        f'</w:r>'
        f'</w:p>'
    )
    return cap_elem

def make_placeholder_tbl():
    tbl_elem = parse_xml(
        f'<w:tbl {nsdecls("w")}>'
        f'<w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
        f'<w:tblGrid><w:gridCol w:w="8928"/></w:tblGrid>'
        f'<w:tr><w:tc><w:p/></w:tc></w:tr>'
        f'</w:tbl>'
    )
    return tbl_elem

def make_heading_p(text, level=2):
    esc_t = escape_xml(text)
    sz = "28" if level == 2 else "24"
    h_elem = parse_xml(
        f'<w:p {nsdecls("w")}>'
        f'<w:pPr>'
        f'<w:pStyle w:val="Heading{level}"/>'
        f'<w:spacing w:before="240" w:after="120"/>'
        f'</w:pPr>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:b/>'
        f'<w:sz w:val="{sz}"/>'
        f'<w:color w:val="0F172A"/>'
        f'</w:rPr>'
        f'<w:t>{esc_t}</w:t>'
        f'</w:r>'
        f'</w:p>'
    )
    return h_elem

def process_full_document(input_path, output_path):
    print(f"Loading document: {input_path}")
    doc = docx.Document(input_path)
    
    # ---------------------------------------------------------------------------
    # STEP 1: Format lại toàn bộ các bảng placeholder 1x1 hiện có thành dashed box
    # ---------------------------------------------------------------------------
    print("Formatting existing 1x1 placeholder tables...")
    for t_idx, table in enumerate(doc.tables):
        if len(table.rows) == 1 and len(table.columns) == 1:
            p_text = table.rows[0].cells[0].text.strip()
            # Giữ nguyên nội dung nếu đã có CẦN BỔ SUNG, nếu chưa có thì gán chuẩn
            if not p_text.startswith("[CẦN BỔ SUNG"):
                if "Sơ đồ" in p_text or "Kiến trúc" in p_text or "Giao diện" in p_text:
                    p_text = f"[CẦN BỔ SUNG DIAGRAM / SCREENSHOT: {p_text}]"
                else:
                    p_text = f"[CẦN BỔ SUNG EVIDENCE / DIAGRAM: {p_text}]"
            table.rows[0].cells[0].text = "" # Xóa text cũ
            format_placeholder_table(table, p_text)

    # ---------------------------------------------------------------------------
    # STEP 2: Xử lý chi tiết từng Chương
    # ---------------------------------------------------------------------------
    
    # ---------------------------------------------------------------------------
    # CHƯƠNG I: Khảo sát & Đặt bài toán
    # ---------------------------------------------------------------------------
    print("Processing Chapter I...")
    body_elements = list(doc._body._element)
    for idx, elem in enumerate(body_elements):
        tag = elem.tag.split('}')[-1]
        text = "".join(elem.itertext()).strip() if tag == "p" else " ".join("".join(elem.itertext()).split()[:20]).strip()
        
        # 1.1 Chuẩn hóa Caption Hình 1.1 Trang chủ YiYi Book
        if tag == "p" and ("Hình 1.1. Trang chủ YiYi Book" in text or "Hình 1.1. Giao diện trang chủ YiYi Book" in text):
            set_p_text(elem, "Hình 1.1. Giao diện Trang chủ Nền tảng Thương mại Điện tử YiYi Book.")
            
        # 1.2 Bổ sung Benchmark Evidence sau bảng khảo sát (Bảng Elem 0054)
        if tag == "tbl" and "Tiki" in text and "Fahasa" in text:
            print(f"  Found Benchmark Table at Elem {idx}")
            bench_items = [
                ("[CẦN BỔ SUNG SCREENSHOT: Ảnh chụp giao diện và tính năng tham khảo Fahasa]", "Hình 1.2. Minh chứng khảo sát tính năng trên Fahasa."),
                ("[CẦN BỔ SUNG SCREENSHOT: Ảnh chụp giao diện và tính năng tham khảo Tiki]", "Hình 1.3. Minh chứng khảo sát tính năng trên Tiki."),
                ("[CẦN BỔ SUNG SCREENSHOT: Ảnh chụp giao diện và luồng tư vấn sách tham khảo Amazon Books]", "Hình 1.4. Minh chứng khảo sát tính năng trên Amazon Books.")
            ]
            curr = elem
            for p_text, cap_text in bench_items:
                t_elem = make_placeholder_tbl()
                c_elem = make_caption_p(cap_text)
                curr.addnext(c_elem)
                curr.addnext(t_elem)
                tbl_obj = Table(t_elem, doc._body)
                format_placeholder_table(tbl_obj, p_text)
                curr = c_elem

    # ---------------------------------------------------------------------------
    # CHƯƠNG II: Quy trình phát triển & Quản lý dự án
    # ---------------------------------------------------------------------------
    print("Processing Chapter II...")
    body_elements = list(doc._body._element)
    for idx, elem in enumerate(body_elements):
        tag = elem.tag.split('}')[-1]
        text = "".join(elem.itertext()).strip() if tag == "p" else " ".join("".join(elem.itertext()).split()[:20]).strip()
        
        # 2.1 Chuẩn hóa Caption Quy trình phát triển Agile
        if tag == "p" and "Hình 2.1. Quy trình phát triển theo mô hình Agile" in text:
            set_p_text(elem, "Hình 2.1. Quy trình phát triển phần mềm theo mô hình Agile/Scrum của nhóm.")
            
        # 2.2 Bổ sung các minh chứng Jira, Git, CI/CD, Deployment sau Bảng phân công (Elem 0100)
        if tag == "tbl" and "Product Owner" in text and "Scrum Master" in text:
            print(f"  Found Team Table at Elem {idx}")
            pm_items = [
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh chụp Jira Scrum Board / Sprint Backlog]", "Hình 2.2. Minh chứng quản trị công việc trên Jira Scrum Board."),
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh GitHub Network Graph / Commit History / Pull Request Template]", "Hình 2.3. Minh chứng quản lý nhánh và lịch sử Commit trên GitHub."),
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh kết quả chạy GitHub Actions CI/CD Pipeline]", "Hình 2.4. Minh chứng thiết lập Pipeline CI/CD tự động hóa."),
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh kết quả chạy Smoke Test và Health Check môi trường Staging/Production]", "Hình 2.5. Minh chứng quy trình triển khai và kiểm thử chấp nhận (Smoke Test).")
            ]
            curr = elem
            for p_text, cap_text in pm_items:
                t_elem = make_placeholder_tbl()
                c_elem = make_caption_p(cap_text)
                curr.addnext(c_elem)
                curr.addnext(t_elem)
                tbl_obj = Table(t_elem, doc._body)
                format_placeholder_table(tbl_obj, p_text)
                curr = c_elem

        # 2.3 Bổ sung minh chứng đào tạo sau Bảng kế hoạch đào tạo (Elem 0122)
        if tag == "tbl" and "Git, GitHub Flow" in text and "Spring Boot RESTful" in text:
            print(f"  Found Training Table at Elem {idx}")
            t_elem = make_placeholder_tbl()
            c_elem = make_caption_p("Hình 2.6. Minh chứng tài liệu đào tạo nội bộ và biên bản họp kỹ thuật của nhóm.")
            elem.addnext(c_elem)
            elem.addnext(t_elem)
            tbl_obj = Table(t_elem, doc._body)
            format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp tài liệu đào tạo nội bộ / slide trình bày / biên bản meeting kỹ thuật]")

    # ---------------------------------------------------------------------------
    # CHƯƠNG III: Phân tích yêu cầu (System Context, Use Case, FR UI, NFR Evidence)
    # ---------------------------------------------------------------------------
    print("Processing Chapter III...")
    body_elements = list(doc._body._element)
    for idx, elem in enumerate(body_elements):
        tag = elem.tag.split('}')[-1]
        text = "".join(elem.itertext()).strip() if tag == "p" else " ".join("".join(elem.itertext()).split()[:20]).strip()
        
        # 3.1 Tách System Context và Use Case tổng quát
        if tag == "p" and "Hình 3.1. System Context và Use Case tổng quát" in text:
            print(f"  Found Combined Context/UseCase at Elem {idx}")
            set_p_text(elem, "Hình 3.1. Sơ đồ ngữ cảnh hệ thống (System Context Diagram).")
            # Chèn Use Case Diagram ngay sau đó
            t_elem = make_placeholder_tbl()
            c_elem = make_caption_p("Hình 3.2. Sơ đồ Use Case tổng quát của hệ thống YiYi Book.")
            elem.addnext(c_elem)
            elem.addnext(t_elem)
            tbl_obj = Table(t_elem, doc._body)
            format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG DIAGRAM: Sơ đồ Use Case tổng quát hệ thống YiYi Book]")

        # 3.2 Chuẩn hóa Caption cho 13 Functional Requirements UI Placeholders
        if tag == "p" and "Hình FR-" in text:
            fr_num = text.split('.')[0].replace("Hình ", "").strip()
            fr_map = {
                "FR-01": "Hình 3.3. Giao diện Xác thực và Đăng ký người dùng (Authentication & Register).",
                "FR-03": "Hình 3.4. Giao diện Duyệt danh mục và Danh sách Sách (Product Catalogue).",
                "FR-05": "Hình 3.5. Giao diện Tìm kiếm và Bộ lọc nâng cao (Search & Filter).",
                "FR-07": "Hình 3.6. Giao diện Chi tiết Sách và Đánh giá (Product Detail & Review).",
                "FR-09": "Hình 3.7. Giao diện Giỏ hàng và Quản lý Địa chỉ nhận hàng (Cart & Address).",
                "FR-11": "Hình 3.8. Giao diện Đặt hàng và Cổng thanh toán trực tuyến (Checkout & Payment).",
                "FR-13": "Hình 3.9. Giao diện Theo dõi Đơn hàng và Hậu mãi (Order Tracking & After-sales).",
                "FR-21": "Hình 3.10. Giao diện Tương tác Trợ lý ảo AI (YiYi AI Assistant).",
                "FR-22": "Hình 3.11. Giao diện Bảng điều khiển Quản trị (Admin Dashboard & Operations).",
                "FR-23": "Hình 3.12. Giao diện Quản trị Sách, Danh mục và Banner (Admin Catalogue & Banners).",
                "FR-24": "Hình 3.13. Giao diện Quản lý Đơn hàng và Vận chuyển (Admin Orders).",
                "FR-25": "Hình 3.14. Giao diện Quản trị Người dùng và Phân quyền (Admin Users & RBAC).",
                "FR-26": "Hình 3.15. Giao diện Quản trị Khuyến mãi, Đánh giá và Cài đặt (Admin Promotions & Settings)."
            }
            if fr_num in fr_map:
                set_p_text(elem, fr_map[fr_num])
                
        # 3.3 Thêm subsection 4.3 Minh chứng NFR sau Bảng NFR
        full_tbl_text = "".join(elem.itertext())
        if tag == "tbl" and ("NFR-01" in full_tbl_text or ("Thuộc tính" in full_tbl_text and "Minh chứng" in full_tbl_text and "NFR" in full_tbl_text)):
            print(f"  Found NFR Table at Elem {idx}")
            h_elem = make_heading_p("4.3 Minh chứng yêu cầu phi chức năng (Non-Functional Requirements Evidence)", level=2)
            nfr_items = [
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh đo thời gian phản hồi API và UI (Response Time < 500ms)]", "Hình 3.16. Kết quả đo thời gian phản hồi của hệ thống."),
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh kiểm thử kiểm soát quyền truy cập RBAC / HTTP 401 Unauthorized / HTTP 403 Forbidden]", "Hình 3.17. Minh chứng kiểm soát truy cập trái phép."),
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh giao diện đáp ứng đa thiết bị (Responsive UI) trên Desktop, Tablet và Mobile]", "Hình 3.18. Minh chứng tính tương thích và Responsive UI."),
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh kiểm thử tính toàn vẹn dữ liệu và Transaction Rollback khi đặt hàng thất bại]", "Hình 3.19. Minh chứng đảm bảo tính toàn vẹn dữ liệu trong xử lý Order.")
            ]
            elem.addnext(h_elem)
            curr = h_elem
            for p_text, cap_text in nfr_items:
                t_elem = make_placeholder_tbl()
                c_elem = make_caption_p(cap_text)
                curr.addnext(c_elem)
                curr.addnext(t_elem)
                tbl_obj = Table(t_elem, doc._body)
                format_placeholder_table(tbl_obj, p_text)
                curr = c_elem

    # ---------------------------------------------------------------------------
    # CHƯƠNG IV: Thiết kế hệ thống (Kiến trúc, Package, ERD, Class & Sequence Diagrams)
    # ---------------------------------------------------------------------------
    print("Processing Chapter IV...")
    body_elements = list(doc._body._element)
    for idx, elem in enumerate(body_elements):
        tag = elem.tag.split('}')[-1]
        text = "".join(elem.itertext()).strip() if tag == "p" else " ".join("".join(elem.itertext()).split()[:20]).strip()
        
        # Kiến trúc tổng thể & Kiến trúc Frontend
        if tag == "p" and "Hình 4.1. System Architecture của YiYi Book" in text:
            set_p_text(elem, "Hình 4.1. Kiến trúc tổng thể hệ thống YiYi Book.")
            
        if tag == "p" and "Hình 4.2. Cấu trúc Frontend component và route" in text:
            set_p_text(elem, "Hình 4.2. Kiến trúc luồng Component và Routing Frontend.")
            
        if tag == "p" and "Hình 4.3. Architecture controller-service-repository của Spring Boot" in text:
            set_p_text(elem, "Hình 4.3. Kiến trúc phân tầng Backend Spring Boot.")
            
        if tag == "p" and "Hình 4.4. Luồng retrieval, intent detection và streaming của YiYi AI" in text:
            set_p_text(elem, "Hình 4.4. Kiến trúc luồng tích hợp Trợ lý AI và xử lý ngôn ngữ tự nhiên.")

        if tag == "p" and "Hình 4.5. Package/module diagram" in text:
            set_p_text(elem, "Hình 4.5. Sơ đồ phân rã gói phần mềm hệ thống (Package Diagram).")

        if tag == "p" and "Hình 4.6. Database ERD" in text:
            set_p_text(elem, "Hình 4.6. Sơ đồ thực thể quan hệ cơ sở dữ liệu (ERD).")

        # Chuẩn hóa các sơ đồ chi tiết chức năng (Chương 4 bị ghi nhầm Hình 3.x)
        if tag == "p" and "Hình 3.1. Login và JWT Authentication" in text:
            set_p_text(elem, "Hình 4.7. Sơ đồ lớp (Class Diagram) và Tuần tự luồng Xác thực JWT.")
            
        if tag == "p" and "Hình 3.2. Xem Product List" in text:
            set_p_text(elem, "Hình 4.8. Sơ đồ tuần tự (Sequence Diagram) luồng Xem danh mục và Danh sách Sách.")

        if tag == "p" and "Hình 3.3. Search và Filter" in text:
            set_p_text(elem, "Hình 4.9. Sơ đồ tuần tự (Sequence Diagram) luồng Tìm kiếm và Lọc sách nâng cao.")

        if tag == "p" and "Hình 3.4. Product Detail và Review" in text:
            set_p_text(elem, "Hình 4.10. Sơ đồ lớp và Tuần tự (Class/Sequence Diagram) Chi tiết Sách & Đánh giá.")

        if tag == "p" and "Hình 3.5. Quản lý Cart" in text:
            set_p_text(elem, "Hình 4.11. Sơ đồ tuần tự (Sequence Diagram) luồng Thao tác Giỏ hàng (Cart Management).")

        if tag == "p" and "Hình 3.6. Checkout và COD" in text:
            set_p_text(elem, "Hình 4.12. Sơ đồ lớp và Tuần tự (Class/Sequence Diagram) Đặt hàng COD & Giảm giá.")

        if tag == "p" and "Hình 3.7. Payment Online" in text:
            set_p_text(elem, "Hình 4.13. Sơ đồ lớp và Tuần tự (Class/Sequence Diagram) Tích hợp Cổng thanh toán Online.")

        if tag == "p" and "Hình 3.8. Order và Return" in text:
            set_p_text(elem, "Hình 4.14. Sơ đồ trạng thái (State Diagram) & Tuần tự luồng Đơn hàng và Đổi trả.")

        if tag == "p" and "Hình 3.9. Y-Point và Membership" in text:
            set_p_text(elem, "Hình 4.15. Sơ đồ tuần tự (Sequence Diagram) luồng Tích lũy và Tiêu điểm Y-Point.")

        if tag == "p" and "Hình 3.10. Chat YiYi AI" in text:
            set_p_text(elem, "Hình 4.16. Sơ đồ thành phần và Tuần tự (Component/Sequence Diagram) Trợ lý AI.")

        if tag == "p" and "Hình 3.11. Catalogue Admin" in text:
            set_p_text(elem, "Hình 4.17. Sơ đồ tuần tự (Sequence Diagram) luồng Quản trị Danh mục & Sách (Admin).")

        if tag == "p" and "Hình 3.12. Xử lý Admin Order" in text:
            set_p_text(elem, "Hình 4.18. Sơ đồ tuần tự (Sequence Diagram) luồng Xử lý và Cập nhật Đơn hàng (Admin).")

        if tag == "p" and "Hình 3.13. Admin Promotion và Content" in text:
            set_p_text(elem, "Hình 4.19. Sơ đồ tuần tự (Sequence Diagram) luồng Quản trị Khuyến mãi & Nội dung (Admin).")

    # ---------------------------------------------------------------------------
    # CHƯƠNG V: Kiểm thử hệ thống (Đã có 12 test evidence -> bổ sung Jira Defect Hình 5.13)
    # ---------------------------------------------------------------------------
    print("Processing Chapter V...")
    body_elements = list(doc._body._element)
    for idx, elem in enumerate(body_elements):
        tag = elem.tag.split('}')[-1]
        text = "".join(elem.itertext()).strip() if tag == "p" else " ".join("".join(elem.itertext()).split()[:20]).strip()
        
        # Bảng Bug Tracking: Elem có BUG-01 và BUG-05
        full_tbl_text = "".join(elem.itertext())
        if tag == "tbl" and ("BUG-01" in full_tbl_text or ("Mã Defect" in full_tbl_text and "Mức độ" in full_tbl_text)):
            print(f"  Found Bug Tracking Table at Elem {idx}")
            t_elem = make_placeholder_tbl()
            c_elem = make_caption_p("Hình 5.13. Minh chứng quản lý và đóng lỗi (Defect Status: Closed/Verified) trên Jira.")
            elem.addnext(c_elem)
            elem.addnext(t_elem)
            tbl_obj = Table(t_elem, doc._body)
            format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp Jira Bug Board / Danh sách Defect trạng thái Closed/Verified]")

    # ---------------------------------------------------------------------------
    # CHƯƠNG VI: Triển khai (8 evidence cài đặt) & Hướng dẫn sử dụng (15 caption Hình 6.9 -> 6.23)
    # ---------------------------------------------------------------------------
    print("Processing Chapter VI...")
    body_elements = list(doc._body._element)
    for idx, elem in enumerate(body_elements):
        tag = elem.tag.split('}')[-1]
        text = "".join(elem.itertext()).strip() if tag == "p" else " ".join("".join(elem.itertext()).split()[:20]).strip()
        
        # 8 Evidence Triển khai / Cài đặt
        if tag == "p" and "Hình 6.1. Ứng dụng local chạy thành công" in text:
            print(f"  Found Installation Section at Elem {idx}")
            set_p_text(elem, "Hình 6.6. Minh chứng đăng nhập thành công vào hệ thống cục bộ (Localhost:5173).")
            # Trước Hình 6.6, ta thêm 5 bước cài đặt môi trường
            prev_tbl = elem.getprevious()
            if prev_tbl is not None and prev_tbl.tag.endswith("tbl"):
                tbl_obj = Table(prev_tbl, doc._body)
                format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình trang chủ với trạng thái đã đăng nhập tài khoản]")
                
                install_items = [
                    ("[CẦN BỔ SUNG EVIDENCE: Ảnh cấu hình file .env và application-local.properties]", "Hình 6.1. Minh chứng cấu hình biến môi trường hệ thống (.env và application.properties)."),
                    ("[CẦN BỔ SUNG EVIDENCE: Ảnh terminal lệnh docker compose up -d và Docker Desktop]", "Hình 6.2. Minh chứng khởi chạy PostgreSQL thành công qua Docker Container."),
                    ("[CẦN BỔ SUNG EVIDENCE: Ảnh console log Spring Boot Started Application in X seconds]", "Hình 6.3. Minh chứng Backend Spring Boot khởi động thành công trên cổng 8080."),
                    ("[CẦN BỔ SUNG EVIDENCE: Ảnh Postman / Browser kiểm tra Swagger UI hoặc API /actuator/health]", "Hình 6.4. Minh chứng kiểm tra API Health Check trả về kết quả UP."),
                    ("[CẦN BỔ SUNG EVIDENCE: Ảnh terminal lệnh npm run dev sẵn sàng truy cập Localhost]", "Hình 6.5. Minh chứng Frontend Vite React khởi động thành công trên cổng 5173.")
                ]
                curr = prev_tbl.getprevious()
                for p_text, cap_text in install_items:
                    t_elem = make_placeholder_tbl()
                    c_elem = make_caption_p(cap_text)
                    curr.addnext(c_elem)
                    curr.addnext(t_elem)
                    tbl_obj = Table(t_elem, doc._body)
                    format_placeholder_table(tbl_obj, p_text)
                    curr = c_elem

            # Sau Hình 6.6, thêm 2 bước kiểm tra Production
            prod_items = [
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình ứng dụng trên Domain Production]", "Hình 6.7. Minh chứng giao diện Frontend triển khai trên Production."),
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh Dashboard Cloud / Server Log / Swagger Production UP]", "Hình 6.8. Minh chứng Backend và Cơ sở dữ liệu Cloud hoạt động trên Production.")
            ]
            curr = elem
            for p_text, cap_text in prod_items:
                t_elem = make_placeholder_tbl()
                c_elem = make_caption_p(cap_text)
                curr.addnext(c_elem)
                curr.addnext(t_elem)
                tbl_obj = Table(t_elem, doc._body)
                format_placeholder_table(tbl_obj, p_text)
                curr = c_elem

    # Chuẩn hóa caption các mục hướng dẫn sử dụng Customer & Admin
    body_elements = list(doc._body._element)
    for idx, elem in enumerate(body_elements):
        tag = elem.tag.split('}')[-1]
        text = "".join(elem.itertext()).strip() if tag == "p" else " ".join("".join(elem.itertext()).split()[:20]).strip()

        hd_checks = [
            ("3.2.1 Register và Login", "Hình 6.9. Hướng dẫn Đăng ký tài khoản người dùng và Đăng nhập hệ thống.", "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Luồng Đăng ký & Đăng nhập Khách hàng]"),
            ("3.2.2 Browse và Search", "Hình 6.10. Hướng dẫn Khám phá Trang chủ và Tìm kiếm Sách nâng cao.", "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Luồng Duyệt danh mục & Tìm kiếm]"),
            ("3.2.3 Product và Wishlist", "Hình 6.11. Hướng dẫn Xem Chi tiết Sách và Quản lý Danh sách yêu thích.", "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Chi tiết sản phẩm & Wishlist]"),
            ("3.2.4 Cart và Checkout", "Hình 6.12. Hướng dẫn Thao tác Giỏ hàng và Nhập địa chỉ Giao hàng.", "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Giỏ hàng & Điền thông tin Đặt hàng]"),
            ("3.2.5 Payment và Order Tracking", "Hình 6.13. Hướng dẫn Thanh toán Trực tuyến và Theo dõi Đơn hàng.", "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Cổng thanh toán & Lịch sử Đơn hàng]"),
            ("3.2.6 Review và Community", "Hình 6.14. Hướng dẫn Đánh giá Sách và Bình luận chia sẻ.", "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Giao diện Viết đánh giá và Xem bình luận]"),
            ("3.2.7 Y-Point và Coupon", "Hình 6.15. Hướng dẫn Tích điểm Thưởng Y-Points và Áp dụng Mã Voucher.", "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Đổi điểm thưởng & Kho Voucher]"),
            ("3.2.8 YiYi AI", "Hình 6.16. Hướng dẫn Tương tác và Tư vấn Sách thông minh qua Trợ lý AI.", "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Cửa sổ Chatbot YiYi AI Assistant]"),
            ("3.3.1 Admin Login và Dashboard", "Hình 6.17. Hướng dẫn Đăng nhập Quản trị và Theo dõi Bảng điều khiển Tổng quan.", "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Admin Login & Dashboard Doanh thu]"),
            ("3.3.2 Quản lý Book", "Hình 6.18. Hướng dẫn Thêm mới, Chỉnh sửa và Quản trị Kho Sách.", "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Giao diện Quản lý Sách và Tồn kho]"),
            ("3.3.3 Quản lý Category và Banner", "Hình 6.19. Hướng dẫn Cấu hình Danh mục Thể loại và Banner Khuyến mãi.", "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Quản trị Thể loại & Banner Marketing]"),
            ("3.3.4 Quản lý Order", "Hình 6.20. Hướng dẫn Duyệt, Xử lý và Chuyển trạng thái Đơn hàng.", "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Quản lý Danh sách Đơn hàng & Vận chuyển]"),
            ("3.3.5 Quản lý User và Role", "Hình 6.21. Hướng dẫn Quản lý Tài khoản Khách hàng và Phân quyền Quản trị.", "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Quản lý Người dùng & RBAC Role]"),
            ("3.3.6 Quản lý Promotion và Reward", "Hình 6.22. Hướng dẫn Tạo Mã giảm giá và Cấu hình Chương trình Điểm thưởng.", "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Quản lý Mã Coupon & Điểm Y-Point]"),
            ("3.3.7 Kiểm duyệt Review và Communication", "Hình 6.23. Hướng dẫn Kiểm duyệt Bình luận, Đánh giá và Quản lý Liên hệ.", "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình Kiểm duyệt Đánh giá & Hòm thư Liên hệ]")
        ]
        
        for heading_text, cap_title, ph_text in hd_checks:
            if tag == "p" and heading_text in text:
                nxt = elem.getnext()
                while nxt is not None and not nxt.tag.endswith("tbl"):
                    nxt = nxt.getnext()
                if nxt is not None and nxt.tag.endswith("tbl"):
                    tbl_obj = Table(nxt, doc._body)
                    format_placeholder_table(tbl_obj, ph_text)
                    nxt_cap = nxt.getnext()
                    if nxt_cap is not None and nxt_cap.tag.endswith("p") and ("Hình " in "".join(nxt_cap.itertext())):
                        set_p_text(nxt_cap, cap_title)
                    else:
                        cap_p = make_caption_p(cap_title)
                        nxt.addnext(cap_p)

    # ---------------------------------------------------------------------------
    # STEP 3: Lưu ra file kết quả
    # ---------------------------------------------------------------------------
    output_target = "YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full.docx"
    print(f"Saving refined document to: {output_target}")
    doc.save(output_target)
    
    # Thử copy sang Final_Refined nếu không bị lock
    try:
        shutil.copy2(output_target, output_path)
        print(f"Successfully synced to {output_path}")
    except Exception as e:
        print(f"Notice: Could not copy directly to {output_path} (File locked by Word). Output is saved at {output_target}")
        
    print("SUCCESS: Refined all chapters placeholders completely!")

if __name__ == "__main__":
    src = "YiYi_Book_Capstone_Project_Report_Tieng_Viet (1).docx"
    dst = "YiYi_Book_Capstone_Project_Report_Tieng_Viet_Final_Refined.docx"
    process_full_document(src, dst)
