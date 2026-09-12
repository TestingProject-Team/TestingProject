import docx
import os

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

targets = [
    'Hãy', 'Nhóm nên', 'Archive', 'archive', 'chèn minh chứng', 'Chèn screenshot',
    'Mục 5.8', '100% Toàn thời gian', 'REQ-AUTH', 'REQ-CAT', 'REQ-CART', 'REQ-ORD', 'REQ-REW',
    '17 bảng', '25 bảng', '8080', '8081', '8082', 'sẵn sàng phục vụ người dùng thực tế', 'hoàn thành xuất sắc', 'CẦN BỔ SUNG'
]

results = {t: [] for t in targets}

# Search paragraphs
for idx, p in enumerate(doc.paragraphs):
    txt = p.text
    for t in targets:
        if t.lower() in txt.lower():
            results[t].append((f'P{idx}', txt[:120].strip()))

# Search tables
for tidx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell_txt = cell.text
            for t in targets:
                if t.lower() in cell_txt.lower():
                    results[t].append((f'T{tidx}[R{r_idx}C{c_idx}]', cell_txt[:120].strip()))

with open(r'E:\TestingProject\audit_scan_initial.txt', 'w', encoding='utf-8') as f:
    for t, matches in results.items():
        f.write(f"=== Target: [{t}] ({len(matches)} occurrences) ===\n")
        for loc, snip in matches[:10]:
            f.write(f"  [{loc}] {snip}\n")
        if len(matches) > 10:
            f.write(f"  ... and {len(matches) - 10} more\n")
        f.write("\n")

print("Scan saved to audit_scan_initial.txt")
