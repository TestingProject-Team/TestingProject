import docx
import re
import sys

doc_path = r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx'
doc = docx.Document(doc_path)

with open('E:\\TestingProject\\structure_full_dump.txt', 'w', encoding='utf-8') as f:
    f.write(f"Total paragraphs: {len(doc.paragraphs)}\n")
    f.write(f"Total tables: {len(doc.tables)}\n\n")
    
    f.write("=== HEADINGS & STRUCTURE ===\n")
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if not t:
            continue
        # Check headings or significant markers
        if p.style.name.startswith('Heading') or any(t.startswith(x) for x in ['CHƯƠNG', 'Chương', 'PHẦN', 'Phần', 'BẢNG', 'Bảng', 'HÌNH', 'Hình', 'PHỤ LỤC', 'Phụ lục']):
            f.write(f"[P{i}][Style:{p.style.name}] {t}\n")

print("Dumped structure to structure_full_dump.txt successfully.")
