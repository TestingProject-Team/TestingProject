import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
docx_path = "YiYi_Book_Capstone_Project_Report_Tieng_Viet (1).docx"
doc = docx.Document(docx_path)

print(f"Total tables: {len(doc.tables)}")

for idx, tbl in enumerate(doc.tables):
    rows = len(tbl.rows)
    cols = len(tbl.columns)
    first_cell = tbl.cell(0,0).text.strip().replace('\n', ' ')[:40] if rows > 0 and cols > 0 else "EMPTY"
    print(f"Table [{idx}]: {rows}x{cols} - Header/First cell: {first_cell}")
