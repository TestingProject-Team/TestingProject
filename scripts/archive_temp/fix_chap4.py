import docx

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

# 1. Update Section 2: Data Model & Port statements
for p in doc.paragraphs:
    # Check data model paragraph (P385, P404)
    if 'Persistence layer: 19 repository và 25 entity class mô hình hóa relational data' in p.text:
        p.text = "Tầng dữ liệu (Persistence Layer): Bao gồm 19 Spring Data JPA repositories và 25 Entity Java classes (trong đó có 21 thực thể ánh xạ bảng cơ sở dữ liệu PostgreSQL và 4 Enums nghiệp vụ) [CẦN XÁC MINH DATABASE SCHEMA: số Entity Java và số Table PostgreSQL thực tế]."
    
    if 'Database được mô hình hóa dưới dạng relational PostgreSQL schema' in p.text and 'Source archive chứa 25 entity class' in p.text:
        p.text = "Cơ sở dữ liệu được mô hình hóa dưới dạng quan hệ trên hệ quản trị cơ sở dữ liệu PostgreSQL. Mã nguồn hệ thống hiện chứa 25 Entity Java classes (gồm 21 thực thể JPA trực tiếp ánh xạ sang 21 bảng quan hệ trong cơ sở dữ liệu và 4 Enum định nghĩa miền giá trị nghiệp vụ) [CẦN XÁC MINH DATABASE SCHEMA: số Entity Java và số Table PostgreSQL thực tế]. Sơ đồ thực thể quan hệ (ERD) hoàn chỉnh thể hiện cấu trúc các bảng thực tế và khóa ngoại liên kết được trình bày dưới đây."

    # Update any combined captions or instructions
    if 'Hình 4.7. Sơ đồ lớp (Class Diagram) và Tuần tự luồng Xác thực JWT.' in p.text:
        p.text = "Hình 4.7. Sơ đồ tuần tự (Sequence Diagram) luồng Xác thực và cấp phát JWT Token."
    
    if 'Hình 4.10. Sơ đồ lớp và Tuần tự (Class/Sequence Diagram) Chi tiết Sách & Đánh giá.' in p.text:
        p.text = "Hình 4.10. Sơ đồ tuần tự (Sequence Diagram) luồng Xem chi tiết Sách & Đánh giá."
        
    if 'Hình 4.12. Sơ đồ lớp và Tuần tự (Class/Sequence Diagram) Đặt hàng COD & Giảm giá.' in p.text:
        p.text = "Hình 4.12. Sơ đồ tuần tự (Sequence Diagram) luồng Đặt hàng COD & Áp dụng Giảm giá."
        
    if 'Hình 4.13. Sơ đồ lớp và Tuần tự (Class/Sequence Diagram) Tích hợp Cổng thanh toán Online.' in p.text:
        p.text = "Hình 4.13. Sơ đồ tuần tự (Sequence Diagram) luồng Tích hợp Cổng thanh toán Online VNPAY."
        
    if 'Hình 4.14. Sơ đồ trạng thái (State Diagram) & Tuần tự luồng Đơn hàng và Đổi trả.' in p.text:
        p.text = "Hình 4.14. Sơ đồ tuần tự (Sequence Diagram) luồng Vận chuyển Đơn hàng và Xử lý Đổi trả."
        
    if 'Hình 4.16. Sơ đồ thành phần và Tuần tự (Component/Sequence Diagram) Trợ lý AI.' in p.text:
        p.text = "Hình 4.16. Sơ đồ tuần tự (Sequence Diagram) luồng Tương tác Trợ lý AI và RAG Engine."

# 2. Update Table 72 (Environment Ports)
# Port chuẩn xác từ application.properties và docker-compose:
# Backend: 8081 (Local / Docker mặc định), 8082 (Docker Test)
# PostgreSQL: 5432 (Local / Docker Compose), 5434 (Test DB)
# Frontend: 5173 (Vite Local Dev)
t72 = doc.tables[72]
t72.rows[1].cells[2].text = "8081 (Local Dev & Docker) / 8082 (Docker Test Profile) [CẦN XÁC MINH PORT TỪ application.properties / application.yml / Docker Compose]"
t72.rows[2].cells[2].text = "5432 (Local & Docker) / 5434 (Docker Test DB)"
t72.rows[2].cells[3].text = "Lưu trữ 21 bảng thực thể PostgreSQL (tương ứng 25 Java Entity/Enum classes) và xử lý giao dịch nguyên tử ACID [CẦN XÁC MINH DATABASE SCHEMA: số Entity Java và số Table PostgreSQL thực tế]"

doc.save(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')
print("Chapter IV updated successfully.")
