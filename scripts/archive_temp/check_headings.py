import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
docx_path = "YiYi_Book_Capstone_Project_Report_Tieng_Viet (1).docx"
doc = docx.Document(docx_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if p.style.name.startswith("Heading 1") or txt.startswith("I.") or txt.startswith("II.") or txt.startswith("III.") or txt.startswith("IV.") or txt.startswith("V.") or txt.startswith("VI.") or txt.startswith("VII."):
        print(f"P[{i}] ({p.style.name}): {txt}")
