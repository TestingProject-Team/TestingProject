import docx

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

# 1. Normalize Paragraphs in Chapter II
for p in doc.paragraphs:
    # P110: Scope and estimation
    if 'Archive cung cấp nhiều minh chứng triển khai' in p.text or 'Hãy hoàn thiện các cột ngày tháng' in p.text:
        p.text = "Phạm vi công việc được tổ chức thành các nhóm: khám phá sản phẩm, giao diện cửa hàng, quản trị, Backend services, tích hợp, đảm bảo chất lượng, triển khai và tài liệu. Lịch trình chi tiết và ước lượng person-days được theo dõi trên hệ thống quản trị dự án Jira/Sprint Plan."
    
    # P123: Process
    if 'Quy trình được đề xuất là Agile workflow' in p.text or 'Quy trình được đề xuất' in p.text:
        p.text = "Dự án áp dụng quy trình Agile/Scrum theo từng vòng lặp (Sprint). Mỗi iteration bắt đầu bằng việc lựa chọn nhóm backlog hướng đến người dùng, xác định acceptance criteria, triển khai các phần Frontend và Backend, kiểm thử API và UI tự động, sau đó kết thúc bằng việc nghiệm thu và đánh giá rủi ro."

# 2. Table 12: Scope & Estimation
t12 = doc.tables[12]
for row in t12.rows[1:]:
    # Col 2: Schedule -> [CẦN XÁC MINH TỪ JIRA/SPRINT PLAN]
    row.cells[2].text = "[CẦN XÁC MINH TỪ JIRA/SPRINT PLAN]"
    # Col 3: Estimation -> [CẦN XÁC MINH TỪ JIRA/SPRINT PLAN]
    row.cells[3].text = "[CẦN XÁC MINH TỪ JIRA/SPRINT PLAN]"

# 3. Table 21: Deliverables
t21 = doc.tables[21]
for row in t21.rows[1:]:
    for cell in row.cells:
        if 'Archive được cung cấp' in cell.text:
            cell.text = cell.text.replace('Archive được cung cấp', 'Đã tích hợp trong Repository mã nguồn')

# 4. Table 22: Responsibility Matrix
# Trạng thái phù hợp: Hoàn thành nội dung / Chờ evidence / Đang hoàn thiện / Hoàn thành
t22 = doc.tables[22]
status_map = {
    'Phan Văn Đỉnh': 'Hoàn thành nội dung (Chờ evidence UAT)',
    'Huỳnh Anh Phú': 'Hoàn thành nội dung (Chờ evidence UI)',
    'Lê Minh Tài': 'Hoàn thành nội dung (Chờ evidence DB/Deploy)',
    'Tạ Huy Thiên Văn': 'Hoàn thành nội dung (Chờ evidence E2E/Jira)',
    'Võ Ngọc Vân Anh': 'Đang hoàn thiện (Chờ evidence Static/Review)'
}
for row in t22.rows[1:]:
    owner = row.cells[0].text.strip()
    for name, st in status_map.items():
        if name in owner:
            row.cells[3].text = st

# 5. Table 71 in Chapter V: Replace "100% Toàn thời gian" with Primary / Supporting
t71 = doc.tables[71]
t71.rows[0].cells[3].text = "Vai trò phụ trách"
roles_commit = {
    'Phan Văn Đỉnh': 'Primary (Điều phối & RTM)',
    'Lê Minh Tài': 'Primary (Scrum & Backend)',
    'Huỳnh Anh Phú': 'Primary (API & Catalogue)',
    'Tạ Huy Thiên Văn': 'Primary (E2E & Checkout)',
    'Võ Ngọc Vân Anh': 'Supporting (Quality & Docs)'
}
for row in t71.rows[1:]:
    mem = row.cells[0].text.strip()
    for name, role_val in roles_commit.items():
        if name in mem:
            row.cells[3].text = role_val

doc.save(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')
print("Chapter II and Table 71 normalized successfully.")
