# -*- coding: utf-8 -*-
import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document("YiYi_Book_Capstone_Project_Report_Tieng_Viet (1).docx")

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

print("\n=== MỤC LỤC TRONG FILE MỚI ===")
for p in doc.paragraphs[20:55]:
    if p.text.strip():
        print(f"[{p.style.name}] {p.text}")

print("\n=== CÁC HEADING CHƯƠNG V MỚI ===")
for p in doc.paragraphs:
    if p.style.name.startswith("Heading"):
        t = p.text.strip()
        if "V. " in t or (len(t) > 0 and t[0].isdigit() and ("Scope" in t or "Strategy" in t or "Plan" in t or "Black-box" in t or "White-box" in t or "API" in t or "End-to-End" in t or "Static" in t or "Execution" in t or "Defect" in t or "Traceability" in t or "Conclusion" in t)):
            print(f"[{p.style.name}] {t}")
