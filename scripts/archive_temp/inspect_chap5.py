import docx

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

with open(r'E:\TestingProject\chap5_inspection.txt', 'w', encoding='utf-8') as f:
    f.write("=== CHAPTER V PARAGRAPHS (592 to 791) ===\n")
    for idx in range(592, 791):
        t = doc.paragraphs[idx].text.strip()
        if any(x in t for x in ['UAT', 'uat', '100%', 'RTM', 'REQ-', '750', 'sẵn sàng', 'hoàn thành xuất sắc', 'Mục 5.8', 'Conclusion', 'Kết luận']) or doc.paragraphs[idx].style.name.startswith('Heading'):
            f.write(f"[P{idx}][{doc.paragraphs[idx].style.name}] {t}\n")

    f.write("\n=== CHAPTER V TABLES ===\n")
    # tables around RTM and defect
    for t_idx in [70, 71, 72, 73, 74, 90, 91, 92, 93, 118, 119, 120, 121]:
        if t_idx < len(doc.tables):
            table = doc.tables[t_idx]
            f.write(f"\n--- TABLE {t_idx} (rows={len(table.rows)}, cols={len(table.columns)}) ---\n")
            for r_idx in range(min(4, len(table.rows))):
                cells_text = [c.text.strip().replace('\n', ' ') for c in table.rows[r_idx].cells]
                f.write(f"Row {r_idx}: {cells_text}\n")

print("Dumped Chap 5 inspection.")
