import docx

doc_path = r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Content_Final.docx'
doc = docx.Document(doc_path)

# 1. Update Chapter I: Target Users in Table 3 (Row 4)
t3 = doc.tables[3]
t3.rows[4].cells[0].text = "Người dùng mục tiêu"
t3.rows[4].cells[1].text = "Khách Vãng Lai (Guest), Khách Hàng Thành Viên (Member), Quản Trị Viên (Administrator)"

# 2. Check and fix any paragraph in Chapter I with old user descriptions
for p in doc.paragraphs[:100]:
    if 'Customer, Administrator, content operator' in p.text:
        p.text = p.text.replace(
            'Customer, Administrator, content operator và thành viên QA/dự án',
            'Khách Vãng Lai (Guest), Khách Hàng Thành Viên (Member), Quản Trị Viên (Administrator)'
        )
    if 'Xem chi tiết sách & Đánh giá' in p.text:
        p.text = p.text.replace('Xem chi tiết sách & Đánh giá', 'Xem chi tiết sách & Xem đánh giá')

# 3. Update Table 28 row for UC-03
t28 = doc.tables[28]
for row in t28.rows:
    if row.cells[0].text.strip() == 'UC-03':
        row.cells[1].text = "Xem chi tiết sách & Xem đánh giá"
        row.cells[3].text = "Xem thông tin xuất bản, tồn kho, tóm tắt nội dung và các bài đánh giá từ cộng đồng"

# 4. Keep benchmark placeholders in Table 6 as requested
t6 = doc.tables[6]
t6.rows[1].cells[2].text = "[CẦN BỔ SUNG: Benchmark system / URL / Evidence]"
t6.rows[2].cells[2].text = "[CẦN BỔ SUNG: Benchmark system / URL / Evidence]"
t6.rows[3].cells[2].text = "[CẦN BỔ SUNG: Benchmark payment gateway / Evidence]"
t6.rows[4].cells[2].text = "[CẦN BỔ SUNG: Benchmark AI assistant / Evidence]"

doc.save(doc_path)
print("Chapter 1 completed successfully.")
