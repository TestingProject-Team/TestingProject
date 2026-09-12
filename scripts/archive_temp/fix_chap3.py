import docx

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

# 1. Update Table 27 (Actors)
t27 = doc.tables[27]
# Existing rows in t27: header + ACT-01 to ACT-07
# Business Actors:
# - Khách Vãng Lai / Guest
# - Khách Hàng Thành Viên / Member
# - Administrator
# External Actors:
# - Cổng Thanh Toán / Payment Gateway
# - Groq AI Service
# - SMTP / Email Service
# REMOVE: Hosting/Storage, QA Member

actors_data = [
    ('ACT-01', 'Khách Vãng Lai (Guest)', 'Duyệt danh mục sách, tìm kiếm, xem chi tiết, đánh giá của sản phẩm, tương tác tư vấn cùng Trợ lý AI và đăng ký/đăng nhập tài khoản.'),
    ('ACT-02', 'Khách Hàng Thành Viên (Member)', 'Quản lý giỏ hàng, danh sách yêu thích, áp dụng coupon/điểm thưởng Y-Point, tạo đơn hàng, thanh toán trực tuyến, theo dõi trạng thái đơn, hủy/đổi trả đơn, viết đánh giá và quản lý hồ sơ cá nhân.'),
    ('ACT-03', 'Quản Trị Viên (Administrator)', 'Truy cập Admin Portal, theo dõi Dashboard doanh thu, quản trị kho sách, danh mục, đơn hàng, người dùng (RBAC), mã giảm giá, đánh giá, banner và cấu hình tham số hệ thống.'),
    ('ACT-04', 'Cổng Thanh Toán (Payment Gateway)', 'External Actor cung cấp dịch vụ thanh toán trực tuyến (VNPAY Sandbox / MoMo / VietQR), sinh URL thanh toán và gửi phản hồi trạng thái giao dịch (IPN/Callback) về Backend.'),
    ('ACT-05', 'Dịch Vụ AI (Groq AI Service)', 'External Actor cung cấp nền tảng suy luận LLM (Llama 3.3 70B qua Groq Cloud) để xử lý ngôn ngữ tự nhiên và trả về phản hồi streaming tư vấn sách dựa trên ngữ cảnh sản phẩm.'),
    ('ACT-06', 'Dịch Vụ Email (SMTP / Email Service)', 'External Actor hỗ trợ gửi email tự động xác nhận đơn hàng, thông báo tài khoản hoặc bản tin tin tức (Newsletter) định kỳ.')
]

# Clear excess rows or re-populate
while len(t27.rows) > 1:
    # remove last row
    tr = t27.rows[-1]._tr
    tr.getparent().remove(tr)

for aid, aname, aresp in actors_data:
    row = t27.add_row()
    row.cells[0].text = aid
    row.cells[1].text = aname
    row.cells[2].text = aresp

print(f"Table 27 updated with {len(t27.rows)-1} actors.")

# 2. Update Table 28 (Use Cases - 26 UCs)
use_cases_26 = [
    # GUEST
    ('UC-01', 'Duyệt danh mục & Trang chủ', 'Khách Vãng Lai / Guest', 'Khám phá sách mới, sách bán chạy, flash sale và cây danh mục sản phẩm'),
    ('UC-02', 'Tìm kiếm & Lọc sách', 'Khách Vãng Lai / Guest', 'Tìm kiếm theo từ khóa tên sách, tác giả và lọc đa tiêu chí (thể loại, giá)'),
    ('UC-03', 'Xem chi tiết sách & Đánh giá', 'Khách Vãng Lai / Guest', 'Xem thông tin xuất bản, tồn kho, tóm tắt nội dung và các bài đánh giá cộng đồng'),
    ('UC-04', 'Tư vấn qua Trợ lý AI', 'Khách Vãng Lai / Guest', 'Hỏi đáp, nhận diện nhu cầu đọc và nhận gợi ý danh mục sách phù hợp từ YiYi AI'),
    ('UC-05', 'Đăng ký tài khoản mới', 'Khách Vãng Lai / Guest', 'Tạo tài khoản khách hàng bằng email và nhận điểm thưởng khởi đầu'),
    ('UC-06', 'Đăng nhập & Quên mật khẩu', 'Khách Vãng Lai / Guest', 'Xác thực tài khoản qua email/mật khẩu để nhận JWT session hoặc khôi phục quyền truy cập'),
    
    # MEMBER
    ('UC-07', 'Quản lý Giỏ hàng', 'Khách Hàng Thành Viên / Member', 'Thêm sách, cập nhật số lượng, kiểm tra tồn kho và xóa mục khỏi giỏ hàng'),
    ('UC-08', 'Quản lý Sách yêu thích', 'Khách Hàng Thành Viên / Member', 'Lưu trữ các tựa sách quan tâm vào Wishlist cá nhân để xem lại thuận tiện'),
    ('UC-09', 'Áp dụng Mã giảm giá & Điểm', 'Khách Hàng Thành Viên / Member', 'Nhập coupon ưu đãi, voucher freeship và đổi điểm thưởng Y-Point trừ trực tiếp vào đơn hàng'),
    ('UC-10', 'Tạo Đơn hàng & COD', 'Khách Hàng Thành Viên / Member', 'Xác nhận thông tin giao hàng và đặt hàng thành công theo phương thức thanh toán tiền mặt (COD)'),
    ('UC-11', 'Thanh toán Online VNPay/MoMo', 'Khách Hàng Thành Viên / Member', 'Chuyển hướng cổng thanh toán VNPAY/VietQR an toàn và nhận kết quả giao dịch tức thời'),
    ('UC-12', 'Theo dõi Trạng thái Đơn hàng', 'Khách Hàng Thành Viên / Member', 'Tra cứu hành trình đơn qua 5 trạng thái (Chờ xác nhận, Đang xử lý, Đang giao, Đã giao, Đã hủy)'),
    ('UC-13', 'Hủy đơn & Yêu cầu hoàn tiền', 'Khách Hàng Thành Viên / Member', 'Hủy đơn hàng chưa xử lý hoặc gửi yêu cầu trả hàng/hoàn tiền kèm minh chứng hình ảnh'),
    ('UC-14', 'Viết Đánh giá & Bình luận', 'Khách Hàng Thành Viên / Member', 'Chấm điểm sao, viết nhận xét sản phẩm đã mua và thảo luận trong các chuỗi bình luận lồng nhau'),
    ('UC-15', 'Quản lý Hồ sơ & Địa chỉ', 'Khách Hàng Thành Viên / Member', 'Cập nhật thông tin cá nhân, sổ địa chỉ nhận hàng, hạng thành viên và sở thích huấn luyện AI'),
    ('UC-16', 'Nhận Thông báo & Liên hệ', 'Khách Hàng Thành Viên / Member', 'Xem thông báo trạng thái đơn hàng, gửi thư liên hệ CSKH và đăng ký nhận bản tin'),
    
    # ADMIN
    ('UC-17', 'Dashboard & Thống kê doanh thu', 'Quản Trị Viên / Administrator', 'Theo dõi tổng quan biểu đồ kinh doanh, doanh số bán lẻ và danh mục đơn chờ xử lý'),
    ('UC-18', 'Quản lý Danh mục & Tác giả', 'Quản Trị Viên / Administrator', 'Thêm mới, chỉnh sửa thông tin, cấu hình cấu trúc phân cấp thể loại và tác giả'),
    ('UC-19', 'Quản trị Kho sách / CRUD & Stock', 'Quản Trị Viên / Administrator', 'Thêm mới, nhập kho hàng loạt, cập nhật giá, tồn kho và trạng thái bày bán của sách'),
    ('UC-20', 'Quản lý Đơn hàng & Vận chuyển', 'Quản Trị Viên / Administrator', 'Duyệt đơn, cập nhật trạng thái đóng gói/vận chuyển và xử lý các yêu cầu đổi trả hàng'),
    ('UC-21', 'Quản lý Người dùng & RBAC', 'Quản Trị Viên / Administrator', 'Quản lý danh sách tài khoản khách hàng, kích hoạt/vô hiệu hóa và phân quyền quản trị'),
    ('UC-22', 'Quản lý Mã giảm giá & Reward', 'Quản Trị Viên / Administrator', 'Thiết lập mã khuyến mại, tỷ lệ chiết khấu, hạn sử dụng và kho quà tặng đổi điểm Y-Point'),
    ('UC-23', 'Quản lý Đánh giá & Phản hồi', 'Quản Trị Viên / Administrator', 'Kiểm duyệt, ẩn hoặc xóa các đánh giá/bình luận vi phạm tiêu chuẩn cộng đồng'),
    ('UC-24', 'Quản lý Banner quảng cáo', 'Quản Trị Viên / Administrator', 'Cấu hình hình ảnh slide banner trang chủ, liên kết chiến dịch và thứ tự hiển thị'),
    ('UC-25', 'Quản lý Yêu cầu liên hệ / CSKH', 'Quản Trị Viên / Administrator', 'Tiếp nhận hòm thư phản hồi, xử lý thắc mắc khách hàng và quản lý danh sách Newsletter'),
    ('UC-26', 'Cấu hình Thông số hệ thống', 'Quản Trị Viên / Administrator', 'Điều chỉnh các cài đặt chung về phí vận chuyển, tỷ lệ quy đổi điểm và thông tin liên hệ')
]

t28 = doc.tables[28]
while len(t28.rows) > 1:
    tr = t28.rows[-1]._tr
    tr.getparent().remove(tr)

for ucid, ucname, ucact, ucgoal in use_cases_26:
    row = t28.add_row()
    row.cells[0].text = ucid
    row.cells[1].text = ucname
    row.cells[2].text = ucact
    row.cells[3].text = ucgoal

print(f"Table 28 updated with {len(t28.rows)-1} use cases.")

# 3. Clean up instructional wording in P180
for p in doc.paragraphs:
    if 'Module map sau đây chuyển archive triển khai' in p.text or 'Nhóm nên sử dụng các ID này' in p.text:
        p.text = "Bảng đối chiếu module hệ thống sau đây liên kết các thành phần kiến trúc với danh mục yêu cầu chức năng, phục vụ công tác kiểm thử và truy vết thiết kế chi tiết."

# 4. Refine Key Business Flows in Chapter III
# Login Flow (around P188 - P192)
for idx, p in enumerate(doc.paragraphs):
    if p.text.strip() == 'FR-01 Register, Login và Session':
        # Inspect following paragraphs
        print("Found FR-01 at P", idx)
        # Main flow starts at idx+4 (Luồng chính)
        doc.paragraphs[idx+5].text = "1. Người dùng truy cập biểu mẫu Đăng nhập và nhập email cùng mật khẩu."
        doc.paragraphs[idx+6].text = "2. Frontend kiểm tra tính hợp lệ dữ liệu nhập (đúng định dạng email, mật khẩu không để trống)."
        doc.paragraphs[idx+7].text = "3. Frontend gửi HTTP POST kèm payload chứng thực đến Backend API (/api/auth/login)."
        doc.paragraphs[idx+8].text = "4. Backend xác thực thông tin tài khoản qua BCrypt, tạo JWT Bearer Token chứa thông tin User ID và Role, trả về phản hồi HTTP 200."
        # Add next line for storing session
        p_next = doc.paragraphs[idx+9]
        if p_next.text.strip() == 'Luồng thay thế và ngoại lệ':
            pass
        else:
            p_next.text = "5. Frontend lưu JWT vào AuthContext và trình duyệt, tự động đính kèm Token qua Axios Interceptors và chuyển hướng đến trang tương ứng."

# Cart Flow (around P240)
for idx, p in enumerate(doc.paragraphs):
    if p.text.strip() == 'FR-09 Cart và Address':
        print("Found FR-09 at P", idx)
        doc.paragraphs[idx+5].text = "1. Người dùng chọn thêm hoặc cập nhật số lượng tựa sách từ trang chi tiết hoặc danh mục."
        doc.paragraphs[idx+6].text = "2. Frontend gửi yêu cầu cập nhật mục giỏ hàng lên Backend API (/api/cart)."
        doc.paragraphs[idx+7].text = "3. Backend kiểm tra tính khả dụng của sản phẩm và số lượng tồn kho thực tế trong cơ sở dữ liệu."
        doc.paragraphs[idx+8].text = "4. Backend đồng bộ bản ghi CartItem, tính toán lại tổng tiền giỏ hàng và trả về dữ liệu giỏ hàng mới nhất."
        doc.paragraphs[idx+9].text = "5. Giao diện người dùng render trạng thái giỏ hàng cập nhật, số lượng badge trên Header và tổng giá trị đơn."

# Checkout Flow (around P253)
for idx, p in enumerate(doc.paragraphs):
    if p.text.strip() == 'FR-11 Checkout và Payment':
        print("Found FR-11 at P", idx)
        doc.paragraphs[idx+5].text = "1. Khách hàng xác nhận danh sách sách trong Giỏ hàng và chọn địa chỉ giao hàng hợp lệ."
        doc.paragraphs[idx+6].text = "2. Khách hàng chọn áp dụng mã giảm giá Coupon và điểm thưởng Y-Point (nếu có)."
        doc.paragraphs[idx+7].text = "3. Khách hàng chọn phương thức thanh toán: COD (tiền mặt khi nhận hàng) hoặc Chuyển khoản trực tuyến (VNPAY/VietQR)."
        doc.paragraphs[idx+8].text = "4. Backend xác thực lại số lượng tồn kho thực tế, kiểm tra tính hợp lệ của mã giảm giá/điểm thưởng, tính toán tổng chi phí và tạo đơn hàng mới ở trạng thái chờ xử lý."
        doc.paragraphs[idx+9].text = "5. Nếu chọn thanh toán Online, Backend sinh URL chuyển hướng cổng VNPAY; nếu COD, hệ thống hiển thị xác nhận đặt hàng thành công và gửi thông báo."

doc.save(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')
print("Chapter III updated successfully.")
