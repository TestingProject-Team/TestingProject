# -*- coding: utf-8 -*-
import os
import sys
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

SRC_DOCX = "YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full.docx"
OUT_DOCX = "YiYi_Book_Capstone_Project_Report_Tieng_Viet_Updated.docx"
IMAGE_PATH = "Hinh_2.1_Agile_Scrum_Workflow_YiYi_Book.png"

def insert_hinh_21(src_path, out_path, image_path):
    print(f"Loading {src_path}...")
    doc = docx.Document(src_path)
    
    body = doc._body._element
    elements = list(body)
    
    target_idx = None
    target_table_elem = None
    
    # Tìm bảng placeholder của Hình 2.1
    for idx, elem in enumerate(elements):
        tag = elem.tag.split("}")[-1]
        text = "".join(elem.itertext())
        if tag == "tbl" and "Sơ đồ quy trình Agile/Scrum và Sprint Workflow" in text:
            # Kiểm tra xem đoạn văn ngay sau có phải Hình 2.1 không
            if idx + 1 < len(elements):
                next_text = "".join(elements[idx+1].itertext())
                if "Hình 2.1" in next_text:
                    target_idx = idx
                    target_table_elem = elem
                    print(f"Found target placeholder table at index {idx}")
                    break
                    
    if target_table_elem is None:
        print("ERROR: Target placeholder table for Hình 2.1 not found!")
        return False
        
    # Tạo paragraph mới để chứa ảnh
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(image_path, width=Inches(6.4))
    
    # Lấy element của p_img vừa tạo
    p_img_elem = p_img._element
    
    # Thay thế target_table_elem bằng p_img_elem
    target_table_elem.addprevious(p_img_elem)
    body.remove(target_table_elem)
    
    # Chuẩn hóa caption của Hình 2.1 (đoạn văn liền kề sau)
    for sub in range(target_idx, min(target_idx + 4, len(body))):
        c_elem = body[sub]
        c_tag = c_elem.tag.split("}")[-1]
        if c_tag == "p":
            c_text = "".join(c_elem.itertext())
            if "Hình 2.1" in c_text:
                for child in list(c_elem):
                    c_elem.remove(child)
                p_obj = docx.text.paragraph.Paragraph(c_elem, doc)
                p_obj.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_obj.paragraph_format.space_before = Pt(2)
                p_obj.paragraph_format.space_after = Pt(8)
                run = p_obj.add_run("Hình 2.1. Sơ đồ quy trình phát triển phần mềm Agile/Scrum của dự án YiYi Book.")
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)
                run.font.italic = True
                run.font.color.rgb = RGBColor(68, 68, 68)
                print("Updated caption formatting successfully.")
                break
                
    doc.save(out_path)
    print(f"Successfully saved to {out_path}")
    
    # Cố gắng ghi đè lên file nguồn nếu không bị lock
    try:
        shutil.copyfile(out_path, src_path)
        print(f"Also synchronized directly to {src_path}")
    except Exception as e:
        print(f"Note: {src_path} is currently opened in Word by user, output is safely written in {out_path}")
        
    return True

if __name__ == "__main__":
    insert_hinh_21(SRC_DOCX, OUT_DOCX, IMAGE_PATH)
