# -*- coding: utf-8 -*-
import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

DOCX_FILES = [
    "YiYi_Book_Capstone_Project_Report_Tieng_Viet_Updated.docx",
    "YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full.docx"
]

IMG_32 = "Hinh_3.2_Overall_Use_Case_Diagram_YiYi_Book.png"

for doc_path in DOCX_FILES:
    if not os.path.exists(doc_path):
        continue
    try:
        doc = docx.Document(doc_path)
        body = doc._body._element
        elements = list(body)
        
        found = False
        for idx, elem in enumerate(elements):
            tag = elem.tag.split("}")[-1]
            text = "".join(elem.itertext())
            if tag == "p" and "Hình 3.2. Sơ đồ Use Case tổng quát hệ thống" in text:
                prev_elem = elements[idx-1]
                
                # Tạo paragraph ảnh mới
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(4)
                run_img = p_img.add_run()
                run_img.add_picture(IMG_32, width=Inches(6.4))
                
                # Thay thế phần tử trước caption
                prev_elem.addprevious(p_img._element)
                body.remove(prev_elem)
                found = True
                print(f"Replaced image before caption in {doc_path}")
                break
                
        if found:
            doc.save(doc_path)
            print(f"Successfully re-saved {doc_path}")
    except Exception as e:
        print(f"Could not update {doc_path}: {e}")
