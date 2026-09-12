# -*- coding: utf-8 -*-
import docx
import sys
import io

if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

def inspect_remaining(file_path):
    doc = docx.Document(file_path)
    print(f"P103: {doc.paragraphs[103].text}")
    print(f"T6 R0 C2: {doc.tables[6].rows[0].cells[2].text}")
    print(f"T77 R0 C0: {doc.tables[77].rows[0].cells[0].text}")

if __name__ == "__main__":
    inspect_remaining("YiYi_Book_Capstone_Project_Report_Tieng_Viet_Final_Refined.docx")
