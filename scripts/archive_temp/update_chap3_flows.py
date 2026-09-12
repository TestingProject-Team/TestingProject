import docx, re

doc_path = r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx'
doc = docx.Document(doc_path)

# 1. Search and fix double numbering "1. 1.", "2. 2.", etc.
double_patterns = [
    (r'^\s*1\.\s+1\.\s*', '1. '),
    (r'^\s*2\.\s+2\.\s*', '2. '),
    (r'^\s*3\.\s+3\.\s*', '3. '),
    (r'^\s*4\.\s+4\.\s*', '4. '),
    (r'^\s*5\.\s+5\.\s*', '5. '),
    (r'^\s*•\s+1\.\s*', '1. '),
    (r'^\s*•\s+2\.\s*', '2. '),
    (r'^\s*•\s+3\.\s*', '3. '),
    (r'^\s*•\s+4\.\s*', '4. '),
]

count_num_fixes = 0
for p in doc.paragraphs:
    orig = p.text
    new_txt = orig
    for pat, rep in double_patterns:
        if re.search(pat, new_txt):
            new_txt = re.sub(pat, rep, new_txt)
            count_num_fixes += 1
    if new_txt != orig:
        p.text = new_txt

print(f"Fixed {count_num_fixes} double numbering instances.")

# 2. Fix NFR-05 Maintainability "archive"
for row in doc.tables[44].rows:
    for cell in row.cells:
        if 'archive' in cell.text.lower():
            cell.text = cell.text.replace('Được quan sát trong archive', 'Được xác nhận thông qua cấu trúc package và phân tầng trong mã nguồn hiện tại.')
            cell.text = cell.text.replace('archive', 'mã nguồn')

for p in doc.paragraphs:
    if 'archive' in p.text.lower():
        p.text = p.text.replace('Được quan sát trong archive', 'Được xác nhận thông qua cấu trúc package và phân tầng trong mã nguồn hiện tại.')
        p.text = p.text.replace('archive được cung cấp', 'mã nguồn dự án')
        p.text = p.text.replace('source archive', 'mã nguồn dự án')

# 3. Refine remaining generic flows in Chapter III (e.g. FR-04, FR-05, FR-07, FR-21, FR-22, FR-23)
# Let's inspect where generic flows are and replace them with concise business steps
# Generic template was:
# "Mở page hoặc widget liên quan."
# "Nhập data được yêu cầu hoặc chọn product/action mong muốn."
# "Frontend gọi Backend hoặc service đã cấu hình."
# "UI render state success, loading, empty hoặc error."

generic_replacements = [
    # FR-04 Home & Catalogue
    ("FR-04 Home và Catalogue Discovery", [
        "1. Khách hàng truy cập Trang chủ hệ thống YiYi Book.",
        "2. Frontend gửi yêu cầu tải dữ liệu banners khuyến mãi, danh sách sách bán chạy (Bestsellers), sách mới và sách Flash Sale từ API (/api/books, /api/banners).",
        "3. Backend truy vấn cơ sở dữ liệu và trả về danh sách sản phẩm cùng thông tin chiết khấu hiện hành.",
        "4. Giao diện hiển thị Hero Slider, bảng đếm ngược Flash Sale, tab luân chuyển danh mục sách nổi bật và đối tác xuất bản."
    ]),
    # FR-05 Search & Filter
    ("FR-05 Search và Filter", [
        "1. Khách hàng nhập từ khóa tìm kiếm (tựa sách, tác giả) hoặc chọn thể loại, khoảng giá mong muốn trên thanh điều hướng.",
        "2. Frontend kích hoạt truy vấn lọc và gửi HTTP GET đến API tìm kiếm (/api/books/search).",
        "3. Backend thực hiện truy vấn cơ sở dữ liệu theo các tiêu chí kết hợp và sắp xếp theo giá hoặc độ phổ biến.",
        "4. Giao diện hiển thị lưới kết quả tương ứng kèm tổng số lượng sách tìm thấy; hiển thị trạng thái gợi ý nếu không có kết quả."
    ]),
    # FR-07 Product Detail & Review
    ("FR-07 Product Detail và Review", [
        "1. Khách hàng chọn một tựa sách cụ thể từ danh mục hoặc kết quả tìm kiếm.",
        "2. Frontend gửi yêu cầu lấy thông tin chi tiết sách (/api/books/{id}) và danh sách đánh giá liên quan (/api/reviews/book/{id}).",
        "3. Backend tổng hợp thông tin xuất bản, tồn kho, điểm đánh giá trung bình và các bình luận phân cấp.",
        "4. Giao diện hiển thị thư viện ảnh sách, giá bán, nút Thêm vào giỏ / Mua ngay và khu vực thảo luận của độc giả."
    ]),
    # FR-21 Trợ lý YiYi AI
    ("FR-21 Trợ lý YiYi AI", [
        "1. Khách hàng mở cửa sổ Trợ lý ảo YiYi AI trên giao diện và nhập câu hỏi tư vấn hoặc nhu cầu tìm sách.",
        "2. Module Intent Detection tại Frontend phân tích mục đích người dùng và kích hoạt cơ chế Mini-RAG để lọc ngữ cảnh sách phù hợp từ kho dữ liệu.",
        "3. Frontend đóng gói ngữ cảnh cùng lịch sử hội thoại, gửi truy vấn an toàn đến Groq API (mô hình Llama 3.3 70B).",
        "4. Phản hồi được hiển thị dạng Server-Sent Events (SSE) streaming kèm thẻ sản phẩm gợi ý trực quan để khách hàng bấm mua ngay."
    ]),
    # FR-22 Dashboard & Operations
    ("FR-22 Dashboard và Operations", [
        "1. Quản trị viên đăng nhập vào Cổng Quản trị Admin Portal bằng tài khoản có quyền ADMIN.",
        "2. Hệ thống tải dữ liệu tổng hợp: tổng doanh thu, số đơn hàng mới cần xử lý, tồn kho cảnh báo và người dùng mới.",
        "3. Backend truy vấn và thống kê dữ liệu đa bảng từ cơ sở dữ liệu PostgreSQL.",
        "4. Dashboard hiển thị các thẻ chỉ số KPI, biểu đồ doanh số theo thời gian và bảng danh sách đơn hàng chờ duyệt."
    ]),
    # FR-23 Book, Category, Banner và Featured Content
    ("FR-23 Book, Category, Banner và Featured Content", [
        "1. Quản trị viên chọn phân hệ quản lý Sách, Danh mục hoặc Banner trên thanh menu quản trị.",
        "2. Quản trị viên thực hiện thao tác Thêm mới, Chỉnh sửa thông tin, cập nhật số lượng tồn kho hoặc tải ảnh bìa.",
        "3. Backend kiểm tra quyền hạn (Role ADMIN), hợp thức hóa dữ liệu và cập nhật bản ghi vào cơ sở dữ liệu.",
        "4. Hệ thống thông báo kết quả thao tác thành công và đồng bộ dữ liệu ngay lập tức lên giao diện cửa hàng."
    ])
]

for title, steps in generic_replacements:
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip() == title:
            # Look for "Luồng chính" in subsequent paragraphs
            for offset in range(1, 15):
                if doc.paragraphs[idx+offset].text.strip() == "Luồng chính":
                    # Replace next 4 paragraphs with steps
                    for step_idx, step_txt in enumerate(steps):
                        p_step = doc.paragraphs[idx+offset+1+step_idx]
                        p_step.text = step_txt
                    break

doc.save(doc_path)
print("Chapter III flows updated successfully.")
