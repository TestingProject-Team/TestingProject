import docx

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

with open(r'E:\TestingProject\chap6_inspection.txt', 'w', encoding='utf-8') as f:
    f.write("=== CHAPTER VI PARAGRAPHS (791 to 946) ===\n")
    for idx in range(791, len(doc.paragraphs)):
        p = doc.paragraphs[idx]
        t = p.text.strip()
        f.write(f"[P{idx}][{p.style.name}] {t}\n")

print("Dumped Chap 6 inspection.")
