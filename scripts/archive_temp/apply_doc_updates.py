# -*- coding: utf-8 -*-
"""
apply_doc_updates.py
Thực hiện cập nhật toàn diện và chính xác file Word.
Xử lý trường hợp file Word đang mở bởi ứng dụng khác.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import sys
import os
import shutil

from report_content_builder import CHAPTER_V_SECTIONS, DETAILED_DESIGN_EXTRA, CHAPTER_VI_SYSTEM_REQ

sys.stdout.reconfigure(encoding='utf-8')

DOCX_IN = "YiYi_Book_Capstone_Project_Report_Tieng_Viet (1)_BACKUP.docx"
DOCX_OUT = "YiYi_Book_Capstone_Project_Report_Tieng_Viet (1).docx"
DOCX_OUT_ALT = "YiYi_Book_Capstone_Project_Report_Tieng_Viet_Updated.docx"

print(f"Loading {DOCX_IN}...")
doc = docx.Document(DOCX_IN)
body = doc._body._element

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def format_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if len(table.rows) > 0:
        for cell in table.rows[0].cells:
            set_cell_background(cell, "E2E8F0")
            set_cell_margins(cell, 120, 120, 160, 160)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.bold = True
                    run.font.size = Pt(10)
                    run.font.name = "Times New Roman"
    for row in table.rows[1:]:
        for cell in row.cells:
            set_cell_margins(cell, 100, 100, 140, 140)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    run.font.name = "Times New Roman"

def insert_p_after(ref_elem, text, style_name="Normal", bold=False, italic=False, size_pt=11, color_rgb=None, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT):
    new_p_elem = parse_xml(f'<w:p {nsdecls("w")}/>')
    ref_elem.addnext(new_p_elem)
    new_p = docx.text.paragraph.Paragraph(new_p_elem, doc)
    new_p.paragraph_format.line_spacing = 1.15
    new_p.paragraph_format.space_after = Pt(space_after)
    new_p.alignment = align
    run = new_p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    return new_p

def insert_heading_after(ref_elem, text, level):
    new_p_elem = parse_xml(f'<w:p {nsdecls("w")}/>')
    ref_elem.addnext(new_p_elem)
    new_p = docx.text.paragraph.Paragraph(new_p_elem, doc)
    try:
        new_p.style = f"Heading {level}"
    except Exception:
        pass
    new_p.paragraph_format.space_before = Pt(8)
    new_p.paragraph_format.space_after = Pt(4)
    run = new_p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    elif level == 2:
        run.font.size = Pt(12.5)
        run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    elif level == 3:
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
    return new_p

def insert_table_after(ref_elem, title, headers, rows_data):
    # Thêm tiêu đề bảng (caption)
    caption_elem = parse_xml(f'<w:p {nsdecls("w")}/>')
    ref_elem.addnext(caption_elem)
    caption_p = docx.text.paragraph.Paragraph(caption_elem, doc)
    caption_p.paragraph_format.space_before = Pt(6)
    caption_p.paragraph_format.space_after = Pt(3)
    c_run = caption_p.add_run(title)
    c_run.font.name = "Times New Roman"
    c_run.font.size = Pt(10)
    c_run.font.bold = True
    c_run.font.italic = True
    
    # Tạo table element và chèn sau caption_elem
    temp_table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    for c_idx, h in enumerate(headers):
        temp_table.cell(0, c_idx).text = str(h)
    for r_idx, row in enumerate(rows_data):
        for c_idx, val in enumerate(row):
            temp_table.cell(r_idx + 1, c_idx).text = str(val)
    format_table(temp_table)
    
    # Di chuyển table xml element tới vị trí sau caption_elem
    caption_elem.addnext(temp_table._tbl)
    
    # Thêm 1 paragraph trống sau table để dãn dòng
    spacer_elem = parse_xml(f'<w:p {nsdecls("w")}/>')
    temp_table._tbl.addnext(spacer_elem)
    return spacer_elem

print("1. Cập nhật Mục lục (TOC)...")
toc_v_index = None
for idx, p in enumerate(doc.paragraphs[:60]):
    if "V. Tài liệu kiểm thử phần mềm" in p.text and idx < 50:
        toc_v_index = idx
        break

if toc_v_index:
    print(f"Found TOC Chapter V at paragraph index {toc_v_index}")
    new_toc_v_entries = [
        "1. Scope (Phạm vi kiểm thử)",
        "2. Strategy (Chiến lược kiểm thử)",
        "3. Plan (Kế hoạch kiểm thử)",
        "4. Black-box Testing (Thiết kế Test Case phương pháp Hộp đen)",
        "5. White-box / Unit Testing (Kiểm thử Hộp trắng & Đơn vị)",
        "6. API Testing (Kiểm thử Giao diện Lập trình Ứng dụng)",
        "7. End-to-End Testing (Kiểm thử Toàn trình Giao diện E2E)",
        "8. Static Analysis & Code Quality (Phân tích mã tĩnh & Chất lượng mã nguồn)",
        "9. Test Execution Report (Báo cáo thực thi kiểm thử)",
        "10. Defect Management (Quản lý và Khắc phục lỗi)",
        "11. Requirement Traceability Matrix (Ma trận truy vết yêu cầu)",
        "12. Test Summary & Conclusion (Tổng kết & Kết luận)"
    ]
    for i, title in enumerate(new_toc_v_entries[:5]):
        doc.paragraphs[toc_v_index + 1 + i].text = f"    {title}\t"
    
    last_toc_p = doc.paragraphs[toc_v_index + 5]
    for title in new_toc_v_entries[5:]:
        last_toc_p = insert_p_after(last_toc_p._element, f"    {title}\t", size_pt=10)

print("2. Bổ sung Class & Sequence Flow vào Chương IV...")
for p in list(doc.paragraphs):
    p_text = p.text.strip()
    if "3.1 Login và JWT Authentication" in p_text and p.style.name.startswith("Heading"):
        for extra in reversed(DETAILED_DESIGN_EXTRA["login"]):
            insert_p_after(p._element, f"• {extra}", italic=True, size_pt=10.5)
    elif "3.3 Search và Filter" in p_text and p.style.name.startswith("Heading"):
        for extra in reversed(DETAILED_DESIGN_EXTRA["browse_search"]):
            insert_p_after(p._element, f"• {extra}", italic=True, size_pt=10.5)
    elif "3.6 Checkout và COD" in p_text and p.style.name.startswith("Heading"):
        for extra in reversed(DETAILED_DESIGN_EXTRA["cart_checkout"]):
            insert_p_after(p._element, f"• {extra}", italic=True, size_pt=10.5)
    elif "3.7 Payment Online" in p_text and p.style.name.startswith("Heading"):
        for extra in reversed(DETAILED_DESIGN_EXTRA["payment_flow"]):
            insert_p_after(p._element, f"• {extra}", italic=True, size_pt=10.5)
    elif "3.10 Chat YiYi AI" in p_text and p.style.name.startswith("Heading"):
        for extra in reversed(DETAILED_DESIGN_EXTRA["yiyi_ai"]):
            insert_p_after(p._element, f"• {extra}", italic=True, size_pt=10.5)

print("3. Thay thế toàn bộ nội dung Chương V thành 12 mục hoàn chỉnh...")
chap5_start_idx = None
chap6_start_idx = None

for idx, child in enumerate(body):
    if child.tag.endswith('p'):
        p = docx.text.paragraph.Paragraph(child, doc)
        if p.text.strip() == "V. Tài liệu kiểm thử phần mềm" and idx > 100:
            chap5_start_idx = idx
        elif p.text.strip().startswith("VI. Gói phát hành") and idx > 100:
            chap6_start_idx = idx

print(f"Chapter V index: {chap5_start_idx}, Chapter VI index: {chap6_start_idx}")

if chap5_start_idx and chap6_start_idx:
    elems_to_remove_count = chap6_start_idx - (chap5_start_idx + 1)
    print(f"Removing {elems_to_remove_count} old elements from Chapter V...")
    for _ in range(elems_to_remove_count):
        elem = body[chap5_start_idx + 1]
        body.remove(elem)
    
    curr_ref = body[chap5_start_idx]
    
    for sec in CHAPTER_V_SECTIONS:
        sec_heading = sec["heading"]
        sec_level = sec["level"]
        heading_p = insert_heading_after(curr_ref, sec_heading, sec_level)
        curr_ref = heading_p._element
        
        for item in sec["content"]:
            if isinstance(item, tuple):
                sub_title, sub_level, sub_paras = item
                sub_h = insert_heading_after(curr_ref, sub_title, sub_level)
                curr_ref = sub_h._element
                for p_item in sub_paras:
                    if isinstance(p_item, str):
                        sub_p = insert_p_after(curr_ref, p_item, size_pt=10.5)
                        curr_ref = sub_p._element
                    elif isinstance(p_item, dict) and p_item.get("type") == "table":
                        title = p_item.get("title", "")
                        headers = p_item["headers"]
                        rows_data = p_item["rows"]
                        spacer_elem = insert_table_after(curr_ref, title, headers, rows_data)
                        curr_ref = spacer_elem
            elif isinstance(item, str):
                p_obj = insert_p_after(curr_ref, item, size_pt=10.5)
                curr_ref = p_obj._element
            elif isinstance(item, dict) and item.get("type") == "table":
                title = item.get("title", "")
                headers = item["headers"]
                rows_data = item["rows"]
                spacer_elem = insert_table_after(curr_ref, title, headers, rows_data)
                curr_ref = spacer_elem

print("4. Bổ sung Yêu cầu phần cứng (Hardware Requirements) vào Chương VI...")
for p in list(doc.paragraphs):
    p_text = p.text.strip()
    if "2.1 System Requirement" in p_text and p.style.name.startswith("Heading"):
        prev_elem = p._element
        for line in CHAPTER_VI_SYSTEM_REQ[1:]:
            p_res = insert_p_after(prev_elem, line, size_pt=10.5)
            prev_elem = p_res._element
        break

print(f"5. Saving updated document...")
try:
    doc.save(DOCX_OUT)
    print(f"SUCCESS: Saved to {DOCX_OUT}")
except PermissionError:
    print(f"Warning: {DOCX_OUT} is currently locked by Word. Saving to {DOCX_OUT_ALT}...")
    doc.save(DOCX_OUT_ALT)
    print(f"SUCCESS: Saved to {DOCX_OUT_ALT}")
