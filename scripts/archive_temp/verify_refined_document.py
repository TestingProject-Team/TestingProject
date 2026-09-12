# -*- coding: utf-8 -*-
import docx
import re
import os
import sys
import io

if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

def verify_refined_document(file_path):
    print(f"=== Comprehensive Audit of Refined Document: {file_path} ===")
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return False

    doc = docx.Document(file_path)
    
    # 1. Check member names
    wrong_names = [
        r"Nguyễn Hữu Phú", r"Nguyễn Tấn Thiện", r"Huỳnh Văn Anh", r"Phan Đình\b(?! Văn Đỉnh)",
        r"Tạ Huy Thiên\b(?! Văn)", r"Nguyễn Văn A", r"Trần Văn B"
    ]
    
    # 2. Check draft keywords
    draft_patterns = [
        r"Bản nháp", r"\[XÁC MINH[^\]]*\]", r"\[CONFIRM[^\]]*\]", r"\[INSERT[^\]]*\]",
        r"\[RETEST GAP[^\]]*\]", r"\[CHÈN[^\]]*\]", r"\[Chèn[^\]]*\]", r"hãy chèn", r"cần chèn",
        r"thay bằng", r"\bTODO\b", r"\bFIXME\b", r"\[GAP[^\]]*\]", r"\bTBD\b"
    ]
    
    # 3. Check credentials
    cred_patterns = [
        r"admin123", r"user123", r"password123", r"secret123"
    ]
    
    # 4. Check Markdown table leftovers
    md_table_pattern = r"^\|.*\|.*\|"

    print("\n--- 1. PARAGRAPHS SCAN ---")
    draft_matches = []
    wrong_name_matches = []
    cred_matches = []
    md_table_matches = []
    
    for p_idx, p in enumerate(doc.paragraphs):
        text = p.text
        if not text:
            continue
            
        for pat in draft_patterns:
            if re.search(pat, text, re.IGNORECASE):
                draft_matches.append((p_idx, text[:120]))
                break
                
        for pat in wrong_names:
            if re.search(pat, text):
                wrong_name_matches.append((p_idx, text[:120]))
                break
                
        for pat in cred_patterns:
            if re.search(pat, text, re.IGNORECASE):
                cred_matches.append((p_idx, text[:120]))
                break
                
        if re.search(md_table_pattern, text.strip()):
            md_table_matches.append((p_idx, text[:120]))

    print(f"Total paragraphs: {len(doc.paragraphs)}")
    print(f"Draft keywords in paragraphs: {len(draft_matches)}")
    for p_idx, text in draft_matches:
        print(f"  [P{p_idx}]: {text}")

    print(f"Wrong member names in paragraphs: {len(wrong_name_matches)}")
    for p_idx, text in wrong_name_matches:
        print(f"  [P{p_idx}]: {text}")

    print(f"Credentials/Secrets in paragraphs: {len(cred_matches)}")
    for p_idx, text in cred_matches:
        print(f"  [P{p_idx}]: {text}")

    print(f"Markdown table leftovers in paragraphs: {len(md_table_matches)}")
    for p_idx, text in md_table_matches:
        print(f"  [P{p_idx}]: {text}")

    print("\n--- 2. TABLES SCAN ---")
    table_draft_matches = []
    table_wrong_name_matches = []
    table_cred_matches = []
    
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                text = cell.text
                if not text:
                    continue
                for pat in draft_patterns:
                    if re.search(pat, text, re.IGNORECASE):
                        table_draft_matches.append((t_idx, r_idx, c_idx, text[:100]))
                        break
                for pat in wrong_names:
                    if re.search(pat, text):
                        table_wrong_name_matches.append((t_idx, r_idx, c_idx, text[:100]))
                        break
                for pat in cred_patterns:
                    if re.search(pat, text, re.IGNORECASE):
                        table_cred_matches.append((t_idx, r_idx, c_idx, text[:100]))
                        break

    print(f"Total tables: {len(doc.tables)}")
    print(f"Draft keywords in tables: {len(table_draft_matches)}")
    for t_idx, r_idx, c_idx, text in table_draft_matches[:10]:
        print(f"  [T{t_idx} R{r_idx} C{c_idx}]: {text}")

    print(f"Wrong member names in tables: {len(table_wrong_name_matches)}")
    for t_idx, r_idx, c_idx, text in table_wrong_name_matches:
        print(f"  [T{t_idx} R{r_idx} C{c_idx}]: {text}")

    print(f"Credentials/Secrets in tables: {len(table_cred_matches)}")
    for t_idx, r_idx, c_idx, text in table_cred_matches:
        print(f"  [T{t_idx} R{r_idx} C{c_idx}]: {text}")

    # Inspecting specific verified placeholders
    print("\n--- 3. VERIFYING STANDARDIZED PLACEHOLDERS ---")
    std_placeholders = []
    for p_idx, p in enumerate(doc.paragraphs):
        if "[CẦN BỔ SUNG" in p.text:
            std_placeholders.append((p_idx, p.text[:100]))
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                if "[CẦN BỔ SUNG" in cell.text:
                    std_placeholders.append((f"T{t_idx}R{r_idx}C{c_idx}", cell.text[:100]))

    print(f"Total standardized [CẦN BỔ SUNG...] placeholders found: {len(std_placeholders)}")
    for loc, text in std_placeholders[:8]:
        print(f"  [{loc}]: {text}")

if __name__ == "__main__":
    verify_refined_document("YiYi_Book_Capstone_Project_Report_Tieng_Viet_Final_Refined.docx")
