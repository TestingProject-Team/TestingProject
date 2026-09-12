# -*- coding: utf-8 -*-
import docx
import sys
import io

if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

def inspect_all_figures_and_placeholders(file_path):
    doc = docx.Document(file_path)
    body = doc._body._body
    
    print(f"Inspecting file: {file_path}")
    print(f"Total elements: {len(body)}")
    
    current_chapter = "FRONT"
    
    for idx, child in enumerate(body):
        text = "".join(child.itertext()).strip()
        tag = child.tag.split('}')[-1]
        
        # Track chapters
        if tag == "p":
            if text.startswith("I. ") or "I. Giới thiệu" in text:
                current_chapter = "CHƯƠNG I"
            elif text.startswith("II. ") or "II. Kế hoạch" in text:
                current_chapter = "CHƯƠNG II"
            elif text.startswith("III. ") or "III. Đặc tả" in text:
                current_chapter = "CHƯƠNG III"
            elif text.startswith("IV. ") or "IV. Mô tả" in text:
                current_chapter = "CHƯƠNG IV"
            elif text.startswith("V. ") or "V. Tài liệu" in text:
                current_chapter = "CHƯƠNG V"
            elif text.startswith("VI. ") or "VI. Gói" in text or "VI. Kết luận" in text:
                current_chapter = "CHƯƠNG VI"
                
            if text.startswith("Hình ") or "Hình " in text[:15]:
                print(f"[{current_chapter}] [Elem {idx}] FIGURE CAPTION: {text[:100]}")
                
        elif tag == "tbl":
            if "[CẦN BỔ SUNG" in text:
                # Get the first line or placeholder text
                snippet = text.split("\n")[0][:100]
                print(f"[{current_chapter}] [Elem {idx}] PLACEHOLDER BOX: {snippet}")

if __name__ == "__main__":
    inspect_all_figures_and_placeholders("YiYi_Book_Capstone_Project_Report_Tieng_Viet (1).docx")
