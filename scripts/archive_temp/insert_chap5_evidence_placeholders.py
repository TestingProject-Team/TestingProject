# -*- coding: utf-8 -*-
"""
insert_chap5_evidence_placeholders.py
Chèn chính xác 12 khung placeholder và caption Word (Hình 5.1 -> Hình 5.12) vào Chương V.
Giữ nguyên toàn bộ cấu trúc, số liệu, bảng biểu hiện có.
"""

import os
import sys
import io
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

def format_placeholder_table(table, placeholder_text):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:bottom w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:left w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:right w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    cell = table.rows[0].cells[0]
    cell.width = Inches(6.2)
    
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
    tcPr.append(shd)
    
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="360" w:type="dxa"/><w:bottom w:w="360" w:type="dxa"/><w:left w:w="240" w:type="dxa"/><w:right w:w="240" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.2
    
    run = p.add_run(placeholder_text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(71, 85, 105) # Slate 600

def insert_placeholder_after(ref_elem, placeholder_text, caption_text):
    """
    Chèn 1 bảng placeholder và 1 đoạn caption ngay sau ref_elem.
    Thứ tự trong DOM sau ref_elem: [tbl_elem, cap_elem]
    """
    tbl_elem = parse_xml(
        f'<w:tbl {nsdecls("w")}>'
        f'<w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
        f'<w:tr>'
        f'<w:tc><w:p><w:r><w:t></w:t></w:r></w:p></w:tc>'
        f'</w:tr>'
        f'</w:tbl>'
    )
    
    cap_elem = parse_xml(
        f'<w:p {nsdecls("w")}>'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="120" w:after="240"/>'
        f'</w:pPr>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:b/><w:i/>'
        f'<w:sz w:val="20"/>'
        f'<w:color w:val="1E293B"/>'
        f'</w:rPr>'
        f'<w:t>{caption_text}</w:t>'
        f'</w:r>'
        f'</w:p>'
    )
    
    ref_elem.addnext(cap_elem)
    ref_elem.addnext(tbl_elem)
    
    from docx.table import Table
    tbl_obj = Table(tbl_elem, ref_elem.getparent())
    format_placeholder_table(tbl_obj, placeholder_text)
    
    return cap_elem

def process_document(input_path, output_path):
    print(f"Reading: {input_path}")
    doc = docx.Document(input_path)
    body = doc._body._body
    
    # 12 vị trí cần chèn
    targets = []
    
    for idx, elem in enumerate(body):
        if idx < 640:
            continue
        text = "".join(elem.itertext()).strip()
        tag = elem.tag.split('}')[-1]
        
        # 1. Sau Bảng 5.9 (Sau Bảng Unit Tests) -> Hình 5.1
        if tag == "tbl" and "Lớp kiểm thử (Test Class)" in text and "Số lượng Test Cases" in text and idx < 760:
            targets.append((1, elem, "[CẦN BỔ SUNG EVIDENCE: Ảnh kết quả chạy 299 Unit Tests bằng lệnh `mvn test` trên Terminal/IntelliJ IDEA]", "Hình 5.1. Kết quả thực thi Unit Test bằng JUnit 5 và Mockito."))
            
        # 2. Sau 5.2 JaCoCo (đoạn Branch Coverage) -> Hình 5.2
        elif "Độ bao phủ nhánh rẽ (Branch Coverage): Đạt 41.27%" in text:
            targets.append((2, elem, "[CẦN BỔ SUNG EVIDENCE: Ảnh giao diện báo cáo HTML tổng quan của JaCoCo Code Coverage (index.html)]", "Hình 5.2. Báo cáo tổng quan Code Coverage bằng JaCoCo."))
            
        # 3. Sau 5.3 Statement & Branch Coverage (đoạn Transactional Rollback) -> Hình 5.3
        elif "Kiểm thử tính toàn vẹn giao dịch (Transactional Rollback)" in text:
            targets.append((3, elem, "[CẦN BỔ SUNG EVIDENCE: Ảnh chi tiết độ bao phủ mã nguồn (Statement & Branch Coverage) của các Service cốt lõi trên giao diện JaCoCo]", "Hình 5.3. Statement và Branch Coverage của các Service nghiệp vụ."))
            
        # 4. Sau 6.2 Newman Automation (đoạn 41.5 giây) -> Hình 5.4
        elif "Thời gian thực thi trung bình: 41.5 giây" in text:
            targets.append((4, elem, "[CẦN BỔ SUNG EVIDENCE: Ảnh màn hình Terminal thực thi bộ kiểm thử API toàn diện bằng Newman CLI]", "Hình 5.4. Kết quả chạy API Automation bằng Newman CLI."))
            
        # 5. Sau Bảng 5.10 (Bảng tổng hợp API) -> Hình 5.5
        elif tag == "tbl" and "Nhóm Module API" in text and "Số Assertions" in text:
            targets.append((5, elem, "[CẦN BỔ SUNG EVIDENCE: Báo cáo kết quả kiểm thử API xuất từ Newman HTML Reporter hoặc Postman Collection Runner]", "Hình 5.5. Báo cáo kết quả kiểm thử API bằng Postman và Newman."))
            
        # 6. Sau Bảng 5.11 (Bảng E2E scenarios) -> Hình 5.6
        elif tag == "tbl" and "Bộ kịch bản (Test Suite)" in text and "Mã kịch bản" in text:
            targets.append((6, elem, "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình kết quả chạy CodeceptJS trên Terminal hiển thị 20/20 kịch bản PASS]", "Hình 5.6. Kết quả thực thi E2E Test bằng CodeceptJS và Playwright."))
            
        # 7a. Trong 7.3 - Sau bullet 1 (plugin screenshotOnFail) -> Hình 5.7
        elif "Tích hợp cơ chế tự động chụp ảnh bằng chứng: Cấu hình plugin screenshotOnFail" in text:
            targets.append((7, elem, "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình giao diện ứng dụng thực tế trong quá trình robot tự động thao tác trình duyệt]", "Hình 5.7. Ảnh giao diện tự động hóa trong quá trình thực thi kịch bản E2E."))
            
        # 7b. Trong 7.3 - Sau bullet 3 (20/20 PASS) -> Hình 5.8
        elif "20/20 kịch bản đạt trạng thái PASS 100% với tổng thời gian" in text:
            targets.append((8, elem, "[CẦN BỔ SUNG EVIDENCE: Ảnh cấu trúc thư mục output/screenshots lưu trữ các ảnh chụp bằng chứng kiểm thử E2E]", "Hình 5.8. Thư mục lưu trữ bằng chứng và ảnh chụp màn hình (output/screenshots)."))
            
        # 8. Sau 8.1 Checkstyle (đoạn 0 vi phạm) -> Hình 5.9
        elif "0 vi phạm Checkstyle (0 violations) trên toàn bộ mã nguồn Java Backend (BUILD SUCCESS)" in text:
            targets.append((9, elem, "[CẦN BỔ SUNG EVIDENCE: Ảnh Terminal thực thi lệnh `mvn checkstyle:check` đạt BUILD SUCCESS và 0 violation]", "Hình 5.9. Báo cáo kiểm tra chuẩn mã nguồn Java bằng Checkstyle (0 violations)."))
            
        # 9. Sau 8.2 SpotBugs (đoạn 0 bugs) -> Hình 5.10
        elif "0 lỗi tiềm ẩn phát hiện (0 bugs found) trên toàn bộ hệ thống" in text:
            targets.append((10, elem, "[CẦN BỔ SUNG EVIDENCE: Ảnh Terminal thực thi lệnh `mvn spotbugs:check` đạt BUILD SUCCESS và 0 bugs found]", "Hình 5.10. Báo cáo phân tích bytecode bằng SpotBugs (0 bugs found)."))
            
        # 10. Sau 8.3 SonarQube & Flake8 (đoạn Flake8) -> Hình 5.11
        elif "Công cụ Flake8 (.flake8): Đảm bảo toàn bộ các mã nguồn kịch bản kiểm thử tự động" in text:
            targets.append((11, elem, "[CẦN BỔ SUNG EVIDENCE: Ảnh giao diện SonarQube Dashboard hiển thị trạng thái Quality Gate PASSED]", "Hình 5.11. Kết quả đánh giá chất lượng mã nguồn trên SonarQube Quality Gate."))
            
        # 11. Sau Bảng 5.12 (Bảng tổng hợp thực thi kiểm thử toàn diện) -> Hình 5.12
        elif tag == "tbl" and "Hạng mục kiểm thử" in text and "Công cụ / Framework" in text:
            targets.append((12, elem, "[CẦN BỔ SUNG EVIDENCE: Ảnh giao diện GitHub Actions CI/CD hiển thị toàn bộ workflow kiểm thử tự động đạt trạng thái SUCCESS]", "Hình 5.12. Kết quả thực thi toàn bộ luồng kiểm thử tự động trên CI/GitHub Actions."))

    print(f"Found {len(targets)} targets to insert placeholders:")
    for t_id, elem, p_text, cap in targets:
        print(f"  Target {t_id}: {cap}")
        
    if len(targets) != 12:
        print(f"ERROR: Expected 12 targets, but found {len(targets)}! Aborting.")
        return False
        
    # Chèn các khung placeholder
    for t_id, elem, p_text, cap in targets:
        insert_placeholder_after(elem, p_text, cap)
        print(f"  -> Inserted {cap}")
        
    temp_output = output_path + ".tmp.docx"
    print(f"\nSaving document to temporary file {temp_output}...")
    doc.save(temp_output)
    
    if os.path.exists(output_path):
        try:
            os.remove(output_path)
            os.rename(temp_output, output_path)
            print(f"Successfully saved to {output_path}!")
        except Exception as e:
            print(f"Could not overwrite {output_path} directly ({e}). Saved at {temp_output}.")
    else:
        os.rename(temp_output, output_path)
        print(f"Successfully saved to {output_path}!")

    return True

if __name__ == "__main__":
    src = "YiYi_Book_Capstone_Project_Report_Tieng_Viet_Final_Refined.docx"
    dst = "YiYi_Book_Capstone_Project_Report_Tieng_Viet_Final_Refined.docx"
    process_document(src, dst)
