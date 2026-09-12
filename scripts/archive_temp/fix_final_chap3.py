import docx
import re

doc_path = r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Content_Final.docx'
doc = docx.Document(doc_path)

# 1. Actors (Table 27)
t27 = doc.tables[27]
actors_data = [
    ('ACT-01', 'Khách Vãng Lai (Guest)', 'Duyệt danh mục sách, tìm kiếm, xem chi tiết, xem đánh giá của sản phẩm, tương tác tư vấn cùng Trợ lý AI và đăng ký/đăng nhập tài khoản.'),
    ('ACT-02', 'Khách Hàng Thành Viên (Member)', 'Quản lý giỏ hàng, danh sách yêu thích, áp dụng coupon/điểm thưởng Y-Point, tạo đơn hàng, thanh toán trực tuyến, theo dõi trạng thái đơn, hủy/đổi trả đơn, viết đánh giá và quản lý hồ sơ cá nhân.'),
    ('ACT-03', 'Quản Trị Viên (Administrator)', 'Truy cập Admin Portal, theo dõi Dashboard doanh thu, quản trị kho sách, danh mục, đơn hàng, người dùng (RBAC), mã giảm giá, đánh giá, banner và cấu hình tham số hệ thống.'),
    ('ACT-04', 'Cổng Thanh Toán (Payment Gateway)', 'External Actor cung cấp dịch vụ thanh toán trực tuyến (VNPAY Sandbox / MoMo / VietQR), sinh URL thanh toán và gửi phản hồi trạng thái giao dịch (IPN/Callback) về Backend.'),
    ('ACT-05', 'Dịch Vụ AI (Groq AI Service)', 'External Actor cung cấp nền tảng suy luận LLM (Llama 3.3 70B qua Groq Cloud) để xử lý ngôn ngữ tự nhiên và trả về phản hồi streaming tư vấn sách dựa trên ngữ cảnh sản phẩm.'),
    ('ACT-06', 'Dịch Vụ Email (SMTP / Email Service)', 'External Actor hỗ trợ gửi email tự động xác nhận đơn hàng, thông báo tài khoản hoặc bản tin tin tức (Newsletter) định kỳ.')
]
while len(t27.rows) > 1:
    tr = t27.rows[-1]._tr
    tr.getparent().remove(tr)

for aid, aname, aresp in actors_data:
    row = t27.add_row()
    row.cells[0].text = aid
    row.cells[1].text = aname
    row.cells[2].text = aresp

# 2. Use Cases (Table 28) - 26 Use cases with UC-03 as 'Xem chi tiết sách & Xem đánh giá'
use_cases_26 = [
    ('UC-01', 'Duyệt danh mục & Trang chủ', 'Khách Vãng Lai / Guest', 'Khám phá sách mới, sách bán chạy, flash sale và cây danh mục sản phẩm'),
    ('UC-02', 'Tìm kiếm & Lọc sách', 'Khách Vãng Lai / Guest', 'Tìm kiếm theo từ khóa tên sách, tác giả và lọc đa tiêu chí (thể loại, giá)'),
    ('UC-03', 'Xem chi tiết sách & Xem đánh giá', 'Khách Vãng Lai / Guest', 'Xem thông tin xuất bản, tồn kho, tóm tắt nội dung và các bài đánh giá cộng đồng'),
    ('UC-04', 'Tư vấn qua Trợ lý AI', 'Khách Vãng Lai / Guest', 'Hỏi đáp, nhận diện nhu cầu đọc và nhận gợi ý danh mục sách phù hợp từ YiYi AI'),
    ('UC-05', 'Đăng ký tài khoản mới', 'Khách Vãng Lai / Guest', 'Tạo tài khoản khách hàng bằng email và nhận điểm thưởng khởi đầu'),
    ('UC-06', 'Đăng nhập & Quên mật khẩu', 'Khách Vãng Lai / Guest', 'Xác thực tài khoản qua email/mật khẩu để nhận JWT session hoặc khôi phục quyền truy cập'),
    ('UC-07', 'Quản lý Giỏ hàng', 'Khách Hàng Thành Viên / Member', 'Thêm sách, cập nhật số lượng, kiểm tra tồn kho và xóa mục khỏi giỏ hàng'),
    ('UC-08', 'Quản lý Sách yêu thích', 'Khách Hàng Thành Viên / Member', 'Lưu trữ các tựa sách quan tâm vào Wishlist cá nhân để xem lại thuận tiện'),
    ('UC-09', 'Áp dụng Mã giảm giá & Điểm', 'Khách Hàng Thành Viên / Member', 'Nhập coupon ưu đãi, voucher freeship và đổi điểm thưởng Y-Point trừ trực tiếp vào đơn hàng'),
    ('UC-10', 'Tạo Đơn hàng & COD', 'Khách Hàng Thành Viên / Member', 'Xác nhận thông tin giao hàng và đặt hàng thành công theo phương thức thanh toán tiền mặt (COD)'),
    ('UC-11', 'Thanh toán Online VNPay/MoMo', 'Khách Hàng Thành Viên / Member', 'Chuyển hướng cổng thanh toán VNPAY/VietQR an toàn và nhận kết quả giao dịch tức thời'),
    ('UC-12', 'Theo dõi Trạng thái Đơn hàng', 'Khách Hàng Thành Viên / Member', 'Tra cứu hành trình đơn qua 5 trạng thái (Chờ xác nhận, Đang xử lý, Đang giao, Đã giao, Đã hủy)'),
    ('UC-13', 'Hủy đơn & Yêu cầu hoàn tiền', 'Khách Hàng Thành Viên / Member', 'Hủy đơn hàng chưa xử lý hoặc gửi yêu cầu trả hàng/hoàn tiền kèm minh chứng hình ảnh'),
    ('UC-14', 'Viết Đánh giá & Bình luận', 'Khách Hàng Thành Viên / Member', 'Chấm điểm sao, viết nhận xét sản phẩm đã mua và thảo luận trong các chuỗi bình luận lồng nhau'),
    ('UC-15', 'Quản lý Hồ sơ & Địa chỉ', 'Khách Hàng Thành Viên / Member', 'Cập nhật thông tin cá nhân, sổ địa chỉ nhận hàng, hạng thành viên và sở thích huấn luyện AI'),
    ('UC-16', 'Nhận Thông báo & Liên hệ', 'Khách Hàng Thành Viên / Member', 'Xem thông báo trạng thái đơn hàng, gửi thư liên hệ CSKH và đăng ký nhận bản tin'),
    ('UC-17', 'Dashboard & Thống kê doanh thu', 'Quản Trị Viên / Administrator', 'Theo dõi tổng quan biểu đồ kinh doanh, doanh số bán lẻ và danh mục đơn chờ xử lý'),
    ('UC-18', 'Quản lý Danh mục & Tác giả', 'Quản Trị Viên / Administrator', 'Thêm mới, chỉnh sửa thông tin, cấu hình cấu trúc phân cấp thể loại và tác giả'),
    ('UC-19', 'Quản trị Kho sách / CRUD & Stock', 'Quản Trị Viên / Administrator', 'Thêm mới, nhập kho hàng loạt, cập nhật giá, tồn kho và trạng thái bày bán của sách'),
    ('UC-20', 'Quản lý Đơn hàng & Vận chuyển', 'Quản Trị Viên / Administrator', 'Duyệt đơn, cập nhật trạng thái đóng gói/vận chuyển và xử lý các yêu cầu đổi trả hàng'),
    ('UC-21', 'Quản lý Người dùng & RBAC', 'Quản Trị Viên / Administrator', 'Quản lý danh sách tài khoản khách hàng, kích hoạt/vô hiệu hóa và phân quyền quản trị'),
    ('UC-22', 'Quản lý Mã giảm giá & Reward', 'Quản Trị Viên / Administrator', 'Thiết lập mã khuyến mại, tỷ lệ chiết khấu, hạn sử dụng và kho quà tặng đổi điểm Y-Point'),
    ('UC-23', 'Quản lý Đánh giá & Phản hồi', 'Quản Trị Viên / Administrator', 'Kiểm duyệt, ẩn hoặc xóa các đánh giá/bình luận vi phạm tiêu chuẩn cộng đồng'),
    ('UC-24', 'Quản lý Banner quảng cáo', 'Quản Trị Viên / Administrator', 'Cấu hình hình ảnh slide banner trang chủ, liên kết chiến dịch và thứ tự hiển thị'),
    ('UC-25', 'Quản lý Yêu cầu liên hệ / CSKH', 'Quản Trị Viên / Administrator', 'Tiếp nhận hòm thư phản hồi, xử lý thắc mắc khách hàng và quản lý danh sách Newsletter'),
    ('UC-26', 'Cấu hình Thông số hệ thống', 'Quản Trị Viên / Administrator', 'Điều chỉnh các cài đặt chung về phí vận chuyển, tỷ lệ quy đổi điểm và thông tin liên hệ')
]
t28 = doc.tables[28]
while len(t28.rows) > 1:
    tr = t28.rows[-1]._tr
    tr.getparent().remove(tr)

for ucid, ucname, ucact, ucgoal in use_cases_26:
    row = t28.add_row()
    row.cells[0].text = ucid
    row.cells[1].text = ucname
    row.cells[2].text = ucact
    row.cells[3].text = ucgoal

# 3. Check and clean double numbering across all paragraphs
for p in doc.paragraphs:
    # Fix double numbering patterns
    t = p.text
    t = re.sub(r'^\s*1\.\s+1\.\s*', '1. ', t)
    t = re.sub(r'^\s*2\.\s+2\.\s*', '2. ', t)
    t = re.sub(r'^\s*3\.\s+3\.\s*', '3. ', t)
    t = re.sub(r'^\s*4\.\s+4\.\s*', '4. ', t)
    t = re.sub(r'^\s*5\.\s+5\.\s*', '5. ', t)
    t = re.sub(r'^\s*•\s+1\.\s*', '1. ', t)
    t = re.sub(r'^\s*•\s+2\.\s*', '2. ', t)
    t = re.sub(r'^\s*•\s+3\.\s*', '3. ', t)
    t = re.sub(r'^\s*•\s+4\.\s*', '4. ', t)
    p.text = t

# 4. Check NFR-05 in Table 44 and all text for "archive"
t44 = doc.tables[44]
for row in t44.rows:
    for cell in row.cells:
        if 'archive' in cell.text.lower():
            cell.text = cell.text.replace('Được quan sát trong archive', 'Được xác nhận thông qua cấu trúc package và mô hình phân tầng trong mã nguồn hiện tại.')
            cell.text = cell.text.replace('archive được cung cấp', 'mã nguồn dự án')
            cell.text = cell.text.replace('archive', 'mã nguồn dự án')

for p in doc.paragraphs:
    if 'archive' in p.text.lower():
        p.text = p.text.replace('Được quan sát trong archive', 'Được xác nhận thông qua cấu trúc package và mô hình phân tầng trong mã nguồn hiện tại.')
        p.text = p.text.replace('archive được cung cấp', 'mã nguồn dự án')
        p.text = p.text.replace('source archive', 'mã nguồn dự án')
        p.text = p.text.replace('archive', 'mã nguồn dự án')

doc.save(doc_path)
print("Chapter 3 completed successfully.")
