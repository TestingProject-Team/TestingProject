# -*- coding: utf-8 -*-
import docx
import os
import re
import sys

src_path = os.path.abspath('docs/reports_docx/YiYi_Book_Project_Report.docx')
out_path = os.path.abspath('YiYi_Book_Project_Report_Revised_Final.docx')

print("Reading source docx from:", src_path)
doc = docx.Document(src_path)

def safe_replace_para(p, old_pat, new_str, flags=0):
    txt = p.text
    if re.search(old_pat, txt, flags):
        new_txt = re.sub(old_pat, new_str, txt, flags=flags)
        if len(p.runs) > 0:
            p.runs[0].text = new_txt
            for r in p.runs[1:]:
                r.text = ''
        else:
            p.text = new_txt

def safe_replace_cell(cell, old_pat, new_str, flags=0):
    for p in cell.paragraphs:
        safe_replace_para(p, old_pat, new_str, flags=flags)

# 1. Update Headings
heading_maps = [
    (r'^I\.\s+Giới thiệu dự án', 'CHƯƠNG 1. GIỚI THIỆU DỰ ÁN'),
    (r'^1\.\s+Tổng quan$', '1.1. Tổng quan'),
    (r'^1\.1\s+Thông tin dự án', '1.1.1. Thông tin dự án'),
    (r'^1\.2\s+Nhóm dự án', '1.1.2. Nhóm dự án'),
    (r'^2\.\s+Bối cảnh sản phẩm', '1.2. Bối cảnh sản phẩm'),
    (r'^3\.\s+Hệ thống hiện có và tham chiếu thị trường', '1.3. Khảo sát hệ thống hiện có và đối chuẩn thị trường'),
    (r'^4\.\s+Cơ hội kinh doanh', '1.4. Cơ hội kinh doanh'),
    (r'^5\.\s+Tầm nhìn sản phẩm phần mềm', '1.5. Tầm nhìn sản phẩm phần mềm'),
    (r'^6\.\s+Phạm vi và giới hạn dự án', '1.6. Phạm vi và giới hạn dự án'),
    (r'^6\.1\s+Major Features', '1.6.1. Các tính năng chính'),
    (r'^6\.2\s+Giới hạn và nội dung loại trừ', '1.6.2. Giới hạn và nội dung loại trừ'),

    (r'^II\.\s+Kế hoạch quản lý dự án', 'CHƯƠNG 2. KẾ HOẠCH QUẢN LÝ DỰ ÁN'),
    (r'^1\.\s+Tổng quan$', '2.1. Tổng quan quản lý'),
    (r'^1\.1\s+Phạm vi và ước lượng', '2.1.1. Phạm vi và ước lượng công việc'),
    (r'^1\.2\s+Mục tiêu dự án', '2.1.2. Mục tiêu dự án'),
    (r'^1\.3\s+Rủi ro dự án', '2.1.3. Quản lý rủi ro dự án'),
    (r'^2\.\s+Phương pháp quản lý', '2.2. Phương pháp quản lý và quy trình'),
    (r'^2\.1\s+Quy trình dự án', '2.2.1. Quy trình phát triển dự án'),
    (r'^2\.2\s+Quản lý chất lượng', '2.2.2. Quy trình quản lý chất lượng'),
    (r'^2\.3\s+Kế hoạch đào tạo', '2.2.3. Kế hoạch đào tạo và chuyển giao'),
    (r'^3\.\s+Sản phẩm bàn giao của dự án', '2.3. Sản phẩm bàn giao của dự án'),
    (r'^4\.\s+Phân công trách nhiệm', '2.4. Phân công trách nhiệm'),
    (r'^5\.\s+Trao đổi trong dự án', '2.5. Cơ chế trao đổi trong dự án'),
    (r'^6\.\s+Quản lý cấu hình', '2.6. Quản lý cấu hình'),
    (r'^6\.1\s+Quản lý tài liệu', '2.6.1. Quản lý tài liệu'),
    (r'^6\.2\s+Quản lý Source Code', '2.6.2. Quản lý mã nguồn'),
    (r'^6\.3\s+Công cụ và hạ tầng', '2.6.3. Công cụ và hạ tầng hỗ trợ'),

    (r'^III\.\s+Đặc tả yêu cầu phần mềm', 'CHƯƠNG 3. ĐẶC TẢ YÊU CẦU PHẦN MỀM'),
    (r'^1\.\s+Tổng quan sản phẩm', '3.1. Tổng quan sản phẩm'),
    (r'^2\.\s+Yêu cầu người dùng', '3.2. Yêu cầu người dùng'),
    (r'^2\.1\s+Actor của hệ thống', '3.2.1. Tác nhân hệ thống (Actors)'),
    (r'^2\.2\s+Use Case', '3.2.2. Danh mục trường hợp sử dụng (Use Cases)'),
    (r'^3\.\s+Yêu cầu chức năng', '3.3. Yêu cầu chức năng'),
    (r'^3\.1\s+Tổng quan chức năng hệ thống', '3.3.1. Phân rã chức năng hệ thống'),
    (r'^3\.2\s+Web Application', '3.3.2. Đặc tả luồng chức năng ứng dụng web'),
    (r'^3\.2\.1\s+Ứng dụng Customer.*', '3.3.2.1. Phân hệ dành cho Khách hàng'),
    (r'^3\.2\.2\s+Cổng Administration.*', '3.3.2.2. Phân hệ dành cho Quản trị viên'),
    (r'^4\.\s+Yêu cầu phi chức năng', '3.4. Yêu cầu phi chức năng'),
    (r'^4\.1\s+External Interface', '3.4.1. Giao diện bên ngoài (External Interfaces)'),
    (r'^4\.2\s+Thuộc tính chất lượng', '3.4.2. Thuộc tính chất lượng hệ thống'),
    (r'^5\.\s+Phụ lục yêu cầu', '3.5. Quy tắc nghiệp vụ và thông báo hệ thống'),
    (r'^5\.1\s+Business Rule', '3.5.1. Quy tắc nghiệp vụ (Business Rules)'),
    (r'^5\.2\s+Yêu cầu chung', '3.5.2. Các yêu cầu kỹ thuật chung'),
    (r'^5\.3\s+Danh sách Application Message', '3.5.3. Danh mục thông báo ứng dụng'),

    (r'^IV\.\s+Mô tả thiết kế phần mềm', 'CHƯƠNG 4. MÔ TẢ THIẾT KẾ PHẦN MỀM'),
    (r'^1\.\s+Thiết kế hệ thống', '4.1. Thiết kế kiến trúc hệ thống'),
    (r'^1\.1\s+System Architecture', '4.1.1. Kiến trúc tổng thể hệ thống'),
    (r'^1\.1\.1\s+Frontend Architecture', '4.1.1.1. Kiến trúc tầng giao diện (Frontend)'),
    (r'^1\.1\.2\s+Backend Architecture', '4.1.1.2. Kiến trúc tầng dịch vụ (Backend)'),
    (r'^1\.1\.3\s+AI Assistant Architecture', '4.1.1.3. Kiến trúc tích hợp trợ lý ảo AI'),
    (r'^1\.1\.4\s+Deployment Architecture', '4.1.1.4. Kiến trúc triển khai hệ thống'),
    (r'^1\.2\s+Package Diagram', '4.1.2. Sơ đồ phân rã gói phần mềm (Package Diagram)'),
    (r'^2\.\s+Thiết kế Database', '4.2. Thiết kế cơ sở dữ liệu'),
    (r'^2\.1\s+Ghi chú về Data Integrity và Transaction', '4.2.1. Ràng buộc toàn vẹn dữ liệu và giao dịch'),
    (r'^3\.\s+Thiết kế chi tiết', '4.3. Thiết kế chi tiết các luồng xử lý'),

    (r'^V\.\s+Tài liệu kiểm thử phần mềm', 'CHƯƠNG 5. TÀI LIỆU KIỂM THỬ PHẦN MỀM'),
    (r'^1\.\s+Phạm vi kiểm thử', '5.1. Phạm vi kiểm thử'),
    (r'^1\.1\s+In-Scope Testing.*', '5.1.1. Phạm vi trong kiểm thử (In-Scope Testing)'),
    (r'^1\.2\s+Out-of-Scope Testing.*', '5.1.2. Phạm vi ngoài kiểm thử (Out-of-Scope Testing)'),
    (r'^1\.3\s+Constraints & Assumptions.*', '5.1.3. Ràng buộc và giả định kiểm thử'),
    (r'^2\.\s+Chiến lược kiểm thử', '5.2. Chiến lược kiểm thử'),
    (r'^2\.1\s+Testing Types.*', '5.2.1. Các loại hình kiểm thử (Testing Types)'),
    (r'^2\.2\s+Test Levels.*', '5.2.2. Các cấp độ kiểm thử (Test Levels)'),
    (r'^2\.3\s+Testing Tools.*', '5.2.3. Công cụ kiểm thử phần mềm (Testing Tools)'),
    (r'^3\.\s+Kế hoạch kiểm thử', '5.3. Kế hoạch tổ chức kiểm thử'),
    (r'^3\.1\s+Human Resources.*', '5.3.1. Phân công nhân sự kiểm thử'),
    (r'^3\.2\s+Test Environment.*', '5.3.2. Cấu hình môi trường kiểm thử'),
    (r'^3\.3\s+Test Data.*', '5.3.3. Quản lý dữ liệu kiểm thử'),
    (r'^3\.4\s+Entry & Exit Criteria.*', '5.3.4. Tiêu chí bắt đầu và kết thúc kiểm thử'),
    (r'^3\.5\s+Test Milestones.*', '5.3.5. Các cột mốc triển khai kiểm thử'),
    (r'^4\.\s+Thiết kế Test Case bằng Black-box Testing', '5.4. Thiết kế ca kiểm thử hộp đen (Black-box Testing)'),
    (r'^4\.1\s+Equivalence Partitioning.*', '5.4.1. Phân vùng tương đương (Equivalence Partitioning)'),
    (r'^4\.2\s+Boundary Value Analysis.*', '5.4.2. Phân tích giá trị biên (Boundary Value Analysis)'),
    (r'^4\.3\s+Decision Table Testing.*', '5.4.3. Kiểm thử bảng quyết định (Decision Table Testing)'),
    (r'^4\.4\s+State Transition Testing.*', '5.4.4. Kiểm thử chuyển trạng thái (State Transition Testing)'),
    (r'^4\.5\s+Use Case Testing.*', '5.4.5. Kiểm thử theo trường hợp sử dụng (Use Case Testing)'),
    (r'^5\.\s+White-box / Unit Testing', '5.5. Kiểm thử hộp trắng và kiểm thử đơn vị (Unit Testing)'),
    (r'^5\.1\s+JUnit 5 & Mockito Test Implementation', '5.5.1. Triển khai kiểm thử với JUnit 5 và Mockito'),
    (r'^5\.2\s+JaCoCo Code Coverage Report', '5.5.2. Báo cáo độ bao phủ mã nguồn JaCoCo'),
    (r'^5\.3\s+Statement & Branch Coverage Analysis', '5.5.3. Phân tích độ bao phủ câu lệnh và nhánh rẽ'),
    (r'^6\.\s+API Testing', '5.6. Kiểm thử giao diện lập trình ứng dụng (API Testing)'),
    (r'^6\.1\s+Postman Collection Structure', '5.6.1. Cấu trúc bộ sưu tập Postman Collection'),
    (r'^6\.2\s+Newman Automation Execution', '5.6.2. Tự động hóa kiểm thử API với Newman CLI'),
    (r'^6\.3\s+API Test Result Summary', '5.6.3. Tổng hợp kết quả kiểm thử API'),
    (r'^7\.\s+End-to-End Testing', '5.7. Kiểm thử toàn trình giao diện (End-to-End Testing)'),
    (r'^7\.1\s+CodeceptJS Test Design.*', '5.7.1. Thiết kế kịch bản kiểm thử E2E theo mô hình POM'),
    (r'^7\.2\s+E2E Test Execution & Scenarios', '5.7.2. Danh sách kịch bản và kết quả kiểm thử E2E'),
    (r'^7\.3\s+Screenshots & Failure Evidence Plugin', '5.7.3. Thu thập hình ảnh và minh chứng kiểm thử E2E'),
    (r'^8\.\s+Static Analysis & Code Quality', '5.8. Phân tích mã nguồn tĩnh và chất lượng mã'),
    (r'^8\.1\s+Checkstyle Java Standards', '5.8.1. Kiểm tra chuẩn mã nguồn Java với Checkstyle'),
    (r'^8\.2\s+SpotBugs Bytecode Analysis', '5.8.2. Phân tích mã bytecode với SpotBugs'),
    (r'^8\.3\s+SonarQube & Flake8 Configuration', '5.8.3. Cấu hình đánh giá chất lượng SonarQube và Flake8'),
    (r'^9\.\s+Test Execution Report.*', '5.9. Báo cáo tổng hợp thực thi kiểm thử'),
    (r'^9\.1\s+Test Execution Summary', '5.9.1. Tổng kết số liệu thực thi kiểm thử'),
    (r'^9\.2\s+Pass / Fail / Blocked Statistics', '5.9.2. Thống kê tỷ lệ Đạt / Không đạt / Tồn đọng'),
    (r'^9\.3\s+Coverage Summary Metrics', '5.9.3. Đánh giá mức độ bao phủ kiểm thử'),
    (r'^10\.\s+Defect Management.*', '5.10. Quản lý lỗi và khiếm khuyết phần mềm'),
    (r'^10\.1\s+Bug Tracking & Defect Life Cycle', '5.10.1. Quy trình quản lý và vòng đời khiếm khuyết'),
    (r'^10\.2\s+Resolved Defects Log.*', '5.10.2. Nhật ký các khiếm khuyết đã được khắc phục'),
    (r'^10\.3\s+Defect Status & Evidence Index', '5.10.3. Hiện trạng khiếm khuyết và lưu trữ minh chứng'),
    (r'^11\.\s+Requirement Traceability Matrix.*', '5.11. Ma trận truy vết yêu cầu (RTM)'),
    (r'^11\.1\s+Chi tiết Ma trận Truy vết Yêu cầu', '5.11.1. Chi tiết ma trận truy vết yêu cầu phần mềm'),
    (r'^12\.\s+Test Summary & Conclusion.*', '5.12. Đánh giá và định hướng phát triển'),
    (r'^12\.1\s+Đánh giá mức độ sẵn sàng sản phẩm', '5.12.1. Đánh giá mức độ hoàn thiện sản phẩm'),
    (r'^12\.2\s+Đề xuất và Định hướng nâng cấp tiếp theo', '5.12.2. Đề xuất và định hướng nâng cấp tiếp theo'),

    (r'^VI\.\s+Gói phát hành và hướng dẫn người dùng', 'CHƯƠNG 6. GÓI PHÁT HÀNH VÀ HƯỚNG DẪN NGƯỜI DÙNG'),
    (r'^1\.\s+Gói sản phẩm bàn giao', '6.1. Danh mục sản phẩm bàn giao'),
    (r'^2\.\s+Hướng dẫn cài đặt', '6.2. Hướng dẫn cài đặt và cấu hình'),
    (r'^2\.1\s+System Requirement', '6.2.1. Yêu cầu cấu hình hệ thống'),
    (r'^2\.2\s+Hướng dẫn cài đặt', '6.2.2. Các bước cài đặt và vận hành hệ thống'),
    (r'^3\.\s+Hướng dẫn sử dụng', '6.3. Hướng dẫn sử dụng hệ thống'),
    (r'^3\.1\s+Tổng quan', '6.3.1. Tổng quan quy trình thao tác'),
    (r'^3\.2\s+Customer Workflow', '6.3.2. Quy trình thao tác dành cho Khách hàng'),
    (r'^3\.2\.1\s+Register và Login', '6.3.2.1. Đăng ký tài khoản và Đăng nhập'),
    (r'^3\.2\.2\s+Browse và Search', '6.3.2.2. Duyệt danh mục và Tìm kiếm sách'),
    (r'^3\.2\.3\s+Product và Wishlist', '6.3.2.3. Xem chi tiết sách và Quản lý danh sách yêu thích'),
    (r'^3\.2\.4\s+Cart và Checkout', '6.3.2.4. Quản lý giỏ hàng và Đặt hàng'),
    (r'^3\.2\.5\s+Payment và Order Tracking', '6.3.2.5. Thanh toán và Theo dõi đơn hàng'),
    (r'^3\.2\.6\s+Review và Community', '6.3.2.6. Đánh giá và Bình luận chia sẻ'),
    (r'^3\.2\.7\s+Y-Point và Coupon', '6.3.2.7. Tích lũy điểm Y-Point và Áp dụng mã giảm giá'),
    (r'^3\.2\.8\s+YiYi AI', '6.3.2.8. Tương tác với trợ lý ảo YiYi AI'),
    (r'^3\.3\s+Admin/Manager Workflow', '6.3.3. Quy trình thao tác dành cho Quản trị viên'),
    (r'^3\.3\.1\s+Admin Login và Dashboard', '6.3.3.1. Đăng nhập quản trị và Theo dõi bảng điều khiển'),
    (r'^3\.3\.2\s+Quản lý Book', '6.3.3.2. Quản lý danh mục sách và tồn kho'),
    (r'^3\.3\.3\s+Quản lý Category và Banner', '6.3.3.3. Quản lý thể loại và Banner truyền thông'),
    (r'^3\.3\.4\s+Quản lý Order', '6.3.3.4. Quản lý và Chuyển trạng thái đơn hàng'),
    (r'^3\.3\.5\s+Quản lý User và Role', '6.3.3.5. Quản lý tài khoản người dùng và Phân quyền'),
    (r'^3\.3\.6\s+Quản lý Promotion và Reward', '6.3.3.6. Quản lý mã khuyến mãi và Chương trình điểm thưởng'),
    (r'^3\.3\.7\s+Kiểm duyệt Review và Communication', '6.3.3.7. Kiểm duyệt đánh giá và Quản lý hòm thư liên hệ'),
    (r'^3\.4\s+Hỗ trợ và Troubleshooting', '6.3.4. Hướng dẫn xử lý sự cố thường gặp'),
    (r'^Phụ lục A\.\s+Minh chứng và Checklist hoàn thành', 'PHỤ LỤC A. MINH CHỨNG VÀ DANH MỤC THEO DÕI HẠNG MỤC')
]

for p in doc.paragraphs:
    txt = p.text.strip()
    for pat, rep in heading_maps:
        if re.search(pat, txt):
            safe_replace_para(p, pat, rep)
            break

# 2. Update Lời cảm ơn
ack_text = [
    'LỜI CẢM ƠN',
    'Để hoàn thành đồ án tốt nghiệp và xây dựng thành công “Hệ thống Nhà sách trực tuyến và Quản trị YiYi Book”, nhóm sinh viên chúng em đã nhận được sự quan tâm, hỗ trợ và hướng dẫn tận tình từ quý Thầy, Cô.',
    'Nhóm xin bày tỏ lòng cảm ơn chân thành đến giảng viên hướng dẫn Nguyễn Văn Chiến đã luôn theo sát, định hướng chuyên môn, đóng góp những nhận xét kỹ thuật xác đáng, giúp nhóm hoàn thiện yêu cầu hệ thống, quy trình nghiệp vụ, thiết kế kiến trúc, triển khai mã nguồn và xây dựng các bộ kiểm chứng chất lượng phần mềm.',
    'Nhóm cũng gửi lời cảm ơn đến tất cả các thành viên trong nhóm dự án vì tinh thần trách nhiệm và sự nỗ lực phối hợp trong việc phát triển giao diện người dùng, xây dựng dịch vụ nền tảng, thiết kế cơ sở dữ liệu, tích hợp thanh toán, tích hợp trợ lý AI và thực thi các hoạt động kiểm thử.',
    'Báo cáo đồ án được hoàn thiện dựa trên mã nguồn chính thức của hệ thống YiYi Book, tích hợp đặc tả kiến trúc, kết quả kiểm thử và hướng dẫn vận hành.'
]

for idx, p_idx in enumerate([55, 56, 57, 58, 59]):
    if p_idx < len(doc.paragraphs):
        p = doc.paragraphs[p_idx]
        if len(p.runs) > 0:
            p.runs[0].text = ack_text[idx]
            for r in p.runs[1:]:
                r.text = ''
        else:
            p.text = ack_text[idx]

# 3. Terminology and Academic Vietnamese replacements
REPLACEMENTS = [
    (r'Nhóm nên benchmark YiYi Book với trải nghiệm của các online bookstore và marketplace hiện có\..*?nghiệm đã được nhóm phê duyệt\.', 
     'Nhóm tham chiếu mô hình vận hành của một số nhà sách trực tuyến và sàn thương mại điện tử nhằm khảo sát cách tổ chức danh mục, tìm kiếm, đặt hàng và thanh toán. Báo cáo ghi nhận nội dung khảo sát chi tiết cần được bổ sung minh chứng đối chuẩn thực tế.'),
    
    (r'Archive được cung cấp không có tên sinh viên chính thức, thông tin supervisor, project code cuối cùng, yêu cầu đã phê duyệt hoặc deployment evidence cuối cùng\.',
     'Tài liệu báo cáo được xây dựng dựa trên mã nguồn hiện có của hệ thống; các thông tin hành chính và minh chứng triển khai môi trường thực tế được tiếp tục cập nhật và hoàn thiện.'),
     
    (r'Module map sau đây chuyển archive triển khai thành các nhóm yêu cầu\. Nhóm nên sử dụng các ID này khi đính kèm screenshot, API evidence, test case và các phần detailed design\.',
     'Danh mục phân rã dưới đây chuyển đổi các thành phần triển khai thành các nhóm yêu cầu chức năng. Các mã định danh này được sử dụng thống nhất xuyên suốt các phần đặc tả, thiết kế chi tiết và ma trận truy vết kiểm thử.'),
     
    (r'Hướng dẫn sử dụng này được tổ chức theo Actor\. Mỗi bước dưới đây nên đi kèm screenshot từ UI cuối cùng\. Archive được cung cấp có nhiều asset hữu ích, hãy đối chiếu các flow bên dưới với screenshot thực tế\.',
     'Hướng dẫn sử dụng được tổ chức theo từng tác nhân của hệ thống. Mỗi bước thực hiện mô tả tuần tự các thao tác người dùng tương ứng với các giao diện chức năng của ứng dụng.'),
     
    (r'Payment provider được thể hiện thông qua sandbox/configuration flow; production merchant approval, settlement reconciliation và fraud monitoring phải được verify riêng\.',
     'Tính năng thanh toán trực tuyến được triển khai trên môi trường thử nghiệm (Sandbox); các nghiệp vụ đối soát tài khoản thực tế và phê duyệt đối tác thương mại nằm ngoài phạm vi đồ án.'),
     
    (r'AI assistant yêu cầu environment variable được cấu hình an toàn và không được xem là nguồn dữ liệu chính thức về stock, price hoặc policy bên ngoài product context\.',
     'Trợ lý ảo AI yêu cầu cấu hình biến môi trường an toàn và hoạt động như công cụ hỗ trợ tư vấn, không thay thế dữ liệu tồn kho, giá bán chính thức trong cơ sở dữ liệu.'),
     
    (r'Tài liệu dự án xác định các gap về authorization và error handling cần được xử lý trước khi public production release; không sao chép secret hoặc seeded credential vào báo cáo công khai\.',
     'Các cấu hình bảo mật và xử lý ngoại lệ được thiết kế tuân thủ nguyên tắc không để lộ thông tin nhạy cảm hoặc mật khẩu thử nghiệm trong tài liệu báo cáo.'),

    (r'\bNhóm nên\b', 'Nhóm dự kiến'),
    (r'\bnhóm nên\b', 'nhóm dự kiến'),
    (r'\bHãy xác nhận\b', 'Cần xác minh'),
    (r'\bhãy xác nhận\b', 'cần xác minh'),
    (r'\bBáo cáo chủ động để trống\b', 'Báo cáo ghi nhận'),
    (r'\bbáo cáo chủ động để trống\b', 'báo cáo ghi nhận'),
    (r'\bArchive được cung cấp\b', 'Bộ mã nguồn và tài liệu của dự án'),
    (r'\barchive được cung cấp\b', 'bộ mã nguồn và tài liệu của dự án'),
    (r'\bSource archive cho thấy\b', 'Mã nguồn hệ thống thể hiện'),
    (r'\bsource archive cho thấy\b', 'mã nguồn hệ thống thể hiện'),
    (r'\bSource archive chứa\b', 'Mã nguồn hệ thống bao gồm'),
    (r'\bsource archive chứa\b', 'mã nguồn hệ thống bao gồm'),
    (r'\bProject cũng ghi nhận\b', 'Hệ thống đồng thời ghi nhận'),
    (r'\bproject cũng ghi nhận\b', 'hệ thống đồng thời ghi nhận'),
    (r'\bĐược đề xuất là\b', 'Được thiết kế theo'),
    (r'\bđược đề xuất là\b', 'được thiết kế theo'),
    (r'Hãy hoàn thiện các cột ngày tháng và công sức dự kiến dựa trên project board của nhóm\.', 'Các thông số về lịch trình chi tiết và ước lượng công sức được ghi nhận theo kế hoạch triển khai của nhóm.'),
    (r'Hãy xác nhận provider cuối cùng, URL, environment variable, domain và ngày deployment trước khi nộp báo cáo\.', 'Thông tin về nhà cung cấp dịch vụ, biến môi trường và tên miền triển khai được cấu hình theo tài liệu môi trường.'),

    # Exaggerations
    (r'\btoàn diện\b', 'đầy đủ'),
    (r'\bToàn diện\b', 'Đầy đủ'),
    (r'\btriệt để\b', 'hiệu quả'),
    (r'\bTriệt để\b', 'Hiệu quả'),
    (r'\bchuẩn quốc tế\b', 'tiêu chuẩn'),
    (r'\bChuẩn quốc tế\b', 'Tiêu chuẩn'),
    (r'\btối ưu tuyệt đối\b', 'phù hợp'),
    (r'\bhoàn hảo\b', 'đầy đủ'),
    (r'\bhoàn thành xuất sắc\b', 'thực hiện đầy đủ'),
    (r'\bthành công tuyệt đối\b', 'thành công'),

    # English phrases to Vietnamese
    (r'Customer có thể browse, search, filter, compare và purchase book\.', 'Khách hàng có thể duyệt danh mục, tìm kiếm, lọc, so sánh và mua sách.'),
    (r'Customer có thể lưu wishlist, quản lý address, review product và track order\.', 'Khách hàng có thể lưu danh sách yêu thích, quản lý địa chỉ nhận hàng, đánh giá sản phẩm và theo dõi đơn hàng.'),
    (r'Customer có thể sử dụng coupon, Y-Point và truy cập các benefit liên quan đến membership\.', 'Khách hàng có thể sử dụng mã giảm giá, điểm thưởng Y-Point và hưởng các quyền lợi thành viên.'),
    (r'Administrator có thể quản lý catalogue, banner, order, user, review, reward, contact, newsletter, notification và site setting\.', 'Quản trị viên có thể quản lý danh mục sách, banner truyền thông, đơn hàng, người dùng, đánh giá, điểm thưởng, liên hệ, bản tin, thông báo và cài đặt hệ thống.'),
    (r'YiYi AI có thể sử dụng product inventory context để cung cấp recommendation sách ngắn gọn và được cá nhân hóa\.', 'Trợ lý ảo YiYi AI sử dụng ngữ cảnh danh mục sách để gợi ý sách phù hợp theo nhu cầu của người dùng.')
]

for p in doc.paragraphs:
    for pat, rep in REPLACEMENTS:
        safe_replace_para(p, pat, rep)

for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for pat, rep in REPLACEMENTS:
                safe_replace_cell(cell, pat, rep)

# 4. Clean up corrupted placeholder tables
clean_placeholder_texts = {
    107: '[CẦN BỔ SUNG: Ảnh chụp màn hình Giao diện Viết đánh giá và Xem bình luận]',
    108: '[CẦN BỔ SUNG: Ảnh chụp màn hình Đổi điểm thưởng & Kho Voucher]',
    109: '[CẦN BỔ SUNG: Ảnh chụp màn hình Cửa sổ Chatbot YiYi AI Assistant]',
    110: '[CẦN BỔ SUNG: Ảnh chụp màn hình Admin Login & Dashboard Doanh thu]',
    111: '[CẦN BỔ SUNG: Ảnh chụp màn hình Giao diện Quản lý Sách và Tồn kho]',
    112: '[CẦN BỔ SUNG: Ảnh chụp màn hình Quản trị Thể loại & Banner Marketing]',
    113: '[CẦN BỔ SUNG: Ảnh chụp màn hình Quản lý Danh sách Đơn hàng & Vận chuyển]',
    114: '[CẦN BỔ SUNG: Ảnh chụp màn hình Quản lý Người dùng & Phân quyền RBAC]',
    115: '[CẦN BỔ SUNG: Ảnh chụp màn hình Quản lý Mã Coupon & Điểm Y-Point]',
    116: '[CẦN BỔ SUNG: Ảnh chụp màn hình Kiểm duyệt Đánh giá & Hòm thư Liên hệ]'
}

for t_idx, txt in clean_placeholder_texts.items():
    if t_idx < len(doc.tables):
        tbl = doc.tables[t_idx]
        if len(tbl.rows) > 0 and len(tbl.rows[0].cells) > 0:
            c = tbl.rows[0].cells[0]
            if len(c.paragraphs) > 0:
                c.paragraphs[0].text = txt
                for p in c.paragraphs[1:]:
                    p.text = ''

if len(doc.tables) > 117:
    tbl117 = doc.tables[117]
    if len(tbl117.rows) > 0 and len(tbl117.rows[0].cells) > 0:
        tbl117.rows[0].cells[0].paragraphs[0].text = 'Vấn đề phát sinh'

# 5. Clean up Table 3 text
if len(doc.tables) > 3:
    t3 = doc.tables[3]
    for r in t3.rows:
        if 'Minh chứng từ archive' in r.cells[0].text:
            r.cells[0].paragraphs[0].text = 'Minh chứng thành phần mã nguồn'
        if '109 API row được document' in r.cells[1].text:
            r.cells[1].paragraphs[0].text = r.cells[1].text.replace('109 API row được document', '109 mục API được tài liệu hóa')

# Save output file
doc.save(out_path)
print("Saved final document to:", out_path)
