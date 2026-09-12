import docx

doc = docx.Document(r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx')

with open(r'E:\TestingProject\chap3_actors_uc_detail.txt', 'w', encoding='utf-8') as f:
    f.write("=== TABLE 27: ACTORS ===\n")
    t27 = doc.tables[27]
    for r_idx, r in enumerate(t27.rows):
        f.write(f"Row {r_idx}: {[c.text.strip().replace('\n', ' ') for c in r.cells]}\n")

    f.write("\n=== TABLE 28: USE CASES ===\n")
    t28 = doc.tables[28]
    for r_idx, r in enumerate(t28.rows):
        f.write(f"Row {r_idx}: {[c.text.strip().replace('\n', ' ') for c in r.cells]}\n")

print("Dumped table 27 and 28.")
