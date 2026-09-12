# -*- coding: utf-8 -*-
import docx
import os
import sys
from docx.shared import Inches

sys.stdout.reconfigure(encoding='utf-8')

doc_path = os.path.abspath('YiYi_Book_Project_Report_Revised_Final_v2.docx')
doc = docx.Document(doc_path)

IMG_DIR = r"E:\TestingProject\ảnh file docx\01_infographic_revised"

# Map Table index to image filename
DIAGRAM_INSERTS = {
    14: "Hinh_2.1_Agile_Scrum_Workflow_YiYi_Book.png",
    48: "Hinh_4.3_Backend_Layered_Architecture.png",
    49: "Hinh_4.4_AI_RAG_Intent_Flow_Architecture.png",
    50: "Hinh_4.5_Package_Diagram_YiYi_Book.png",
    51: "Hinh_4.6_ERD_Database_YiYi_Book.png",
    57: "Hinh_4.7_Class_Diagram_Login_Auth.png",
    58: "Hinh_4.8_Sequence_Diagram_Product_List.png",
    59: "Hinh_4.9_Sequence_Diagram_Search_Filter.png",
    60: "Hinh_4.10a_Class_Diagram_BookDetail_Review.png",
    61: "Hinh_4.11_Sequence_Diagram_Cart_Management.png",
    62: "Hinh_4.12_Class_Diagram_Checkout_COD.png",
    63: "Hinh_4.13_Class_Diagram_Payment_Online.png",
    64: "Hinh_4.14_Order_State_Diagram.png",
    65: "Hinh_4.15_Sequence_Diagram_Reward_Points.png",
    66: "Hinh_4.16_Component_Diagram_YiYi_AI.png",
    67: "Hinh_4.17_Sequence_Diagram_Admin_Catalogue.png",
    68: "Hinh_4.18_Sequence_Diagram_Admin_Order.png",
    69: "Hinh_4.19_Sequence_Diagram_Admin_Promotion_Content.png"
}

print(f"Updating all diagrams in {doc_path}...")
for t_idx, img_name in DIAGRAM_INSERTS.items():
    if t_idx < len(doc.tables):
        img_path = os.path.join(IMG_DIR, img_name)
        if os.path.exists(img_path):
            tbl = doc.tables[t_idx]
            if len(tbl.rows) > 0 and len(tbl.rows[0].cells) > 0:
                cell = tbl.rows[0].cells[0]
                for p in cell.paragraphs:
                    p.text = ''
                run = cell.paragraphs[0].add_run()
                run.add_picture(img_path, width=Inches(6.2))
                print(f"  -> Table {t_idx:2d} updated with: {img_name}")
        else:
            print(f"  [MISSING FILE] {img_path}")

doc.save(doc_path)
print(f"Successfully updated all diagrams in {doc_path}!")
