from xml.sax.saxutils import escape as escape_xml
# -*- coding: utf-8 -*-
"""
refine_all_chapters_placeholders.py
Thực hiện rà soát và chuẩn hóa toàn bộ các placeholder hình ảnh, sơ đồ UML, screenshot và evidence
từ Chương I đến Chương VI theo đúng quy chuẩn báo cáo PDCMS.
"""

import os
import sys
import io
import shutil
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.table import Table

if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

def set_p_text(p_elem, text):
    """Xóa các run cũ và thêm text mới cho một CT_P xml element."""
    for r in p_elem.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r'):
        p_elem.remove(r)
    esc_text = escape_xml(text)
    r_elem = parse_xml(
        f'<w:r {nsdecls("w")}>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:b/><w:i/>'
        f'<w:sz w:val="20"/>'
        f'<w:color w:val="1E293B"/>'
        f'</w:rPr>'
        f'<w:t>{esc_text}</w:t>'
        f'</w:r>'
    )
    p_elem.append(r_elem)

def format_placeholder_table(table, placeholder_text):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:bottom w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:left w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:right w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    cell = table.rows[0].cells[0]
    cell.width = Inches(6.2)
    
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
    tcPr.append(shd)
    
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="360" w:type="dxa"/><w:bottom w:w="360" w:type="dxa"/><w:left w:w="240" w:type="dxa"/><w:right w:w="240" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.2
    
    run = p.add_run(placeholder_text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(71, 85, 105)

def make_caption_p(caption_text):
    esc_cap = escape_xml(caption_text)
    cap_elem = parse_xml(
        f'<w:p {nsdecls("w")}>'
        f'<w:pPr>'
        f'<w:jc w:val="center"/>'
        f'<w:spacing w:before="120" w:after="240"/>'
        f'</w:pPr>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:b/><w:i/>'
        f'<w:sz w:val="20"/>'
        f'<w:color w:val="1E293B"/>'
        f'</w:rPr>'
        f'<w:t>{esc_cap}</w:t>'
        f'</w:r>'
        f'</w:p>'
    )
    return cap_elem

def make_placeholder_tbl():
    tbl_elem = parse_xml(
        f'<w:tbl {nsdecls("w")}>'
        f'<w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
        f'<w:tr>'
        f'<w:tc><w:p><w:r><w:t></w:t></w:r></w:p></w:tc>'
        f'</w:tr>'
        f'</w:tbl>'
    )
    return tbl_elem

def make_heading_p(text, level=2):
    sz_val = "26" if level == 2 else "24"
    esc_text = escape_xml(text)
    h_elem = parse_xml(
        f'<w:p {nsdecls("w")}>'
        f'<w:pPr>'
        f'<w:spacing w:before="240" w:after="120"/>'
        f'</w:pPr>'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        f'<w:b/>'
        f'<w:sz w:val="{sz_val}"/>'
        f'<w:color w:val="0F172A"/>'
        f'</w:rPr>'
        f'<w:t>{esc_text}</w:t>'
        f'</w:r>'
        f'</w:p>'
    )
    return h_elem

def process_full_document(input_path, output_path):
    print(f"Reading document: {input_path}")
    doc = docx.Document(input_path)
    
    # ---------------------------------------------------------------------------
    # STEP 1: Scan all elements and categorize by Chapter
    # ---------------------------------------------------------------------------
    body_elements = list(doc._body._element)
    elem_info = []
    
    for idx, elem in enumerate(body_elements):
        tag = elem.tag.split('}')[-1]
        text = ""
        if tag == "p":
            text = "".join(elem.itertext()).strip()
        elif tag == "tbl":
            # get sample text from table
            text = " ".join("".join(elem.itertext()).split()[:20]).strip()
        elem_info.append((idx, elem, tag, text))

    print(f"Total root elements: {len(elem_info)}")

    # ---------------------------------------------------------------------------
    # CHƯƠNG I: Cập nhật caption và thêm 3 Benchmark evidence sau bảng Benchmark
    # ---------------------------------------------------------------------------
    print("Processing Chapter I...")
    for idx, elem, tag, text in elem_info:
        # Hình 1.1 Homepage
        if tag == "p" and "Hình 1.1." in text:
            set_p_text(elem, "Hình 1.1. Giao diện Trang chủ hệ thống YiYi Book.")
            
        # Bảng Benchmark: "Tiêu chí" và "Fahasa" và "Tiki"
        if tag == "tbl" and "Fahasa" in text and "Tiki" in text and "Shopee" in text:
            print(f"  Found Benchmark Table at Elem {idx}")
            bm_items = [
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh tham khảo giao diện và luồng nghiệp vụ mua sách tại Fahasa/Tiki]", "Hình 1.2. Minh chứng khảo sát thực tế mô hình Online Bookstore."),
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh tham khảo mô hình Sàn giao dịch & Gian hàng Shopee/Lazada]", "Hình 1.3. Minh chứng khảo sát thực tế mô hình sàn thương mại điện tử."),
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh giao diện cổng thanh toán trực tuyến mẫu (VNPAY/MoMo Sandbox)]", "Hình 1.4. Minh chứng khảo sát cổng thanh toán trực tuyến.")
            ]
            curr = elem
            for p_text, cap_text in bm_items:
                t_elem = make_placeholder_tbl()
                c_elem = make_caption_p(cap_text)
                curr.addnext(c_elem)
                curr.addnext(t_elem)
                tbl_obj = Table(t_elem, doc._body)
                format_placeholder_table(tbl_obj, p_text)
                curr = c_elem

    # ---------------------------------------------------------------------------
    # CHƯƠNG II: Chuẩn hóa Agile diagram và thêm 5 evidence quản lý & đào tạo
    # ---------------------------------------------------------------------------
    print("Processing Chapter II...")
    for idx, elem, tag, text in elem_info:
        # Hình quy trình Agile/Scrum cũ (Hình 2.1. Quy trình phát triển phần mềm Agile/Scrum)
        if tag == "p" and ("Quy trình phát triển phần mềm Agile/Scrum" in text or "Quy trình phát triển Agile/Scrum" in text):
            print(f"  Found Agile Diagram at Elem {idx}")
            set_p_text(elem, "Hình 2.1. Quy trình phát triển phần mềm Agile/Scrum của dự án YiYi Book.")
            
            # Thêm 4 evidence Jira, GitHub, CI, Deployment sau Hình 2.1
            eng_items = [
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh bảng Jira Scrum Board / Sprint Backlog thực tế của dự án]", "Hình 2.2. Minh chứng quản trị công việc và phân chia Sprint trên Jira."),
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh quản lý Git Branch, Commit History và Pull Request trên GitHub]", "Hình 2.3. Minh chứng quy trình Git Branching và Code Review trên GitHub."),
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh GitHub Actions CI Workflow chạy tự động Build & Test thành công]", "Hình 2.4. Minh chứng luồng tự động hóa kiểm thử liên tục (CI/CD Pipeline)."),
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh kết quả Smoke Test môi trường triển khai thực tế]", "Hình 2.5. Minh chứng kiểm thử chấp nhận và triển khai hệ thống.")
            ]
            curr = elem
            for p_text, cap_text in eng_items:
                t_elem = make_placeholder_tbl()
                c_elem = make_caption_p(cap_text)
                curr.addnext(c_elem)
                curr.addnext(t_elem)
                tbl_obj = Table(t_elem, doc._body)
                format_placeholder_table(tbl_obj, p_text)
                curr = c_elem

        # Bảng Kế hoạch chuyển giao công nghệ & Đào tạo: "Nội dung đào tạo"
        if tag == "tbl" and "Nội dung đào tạo" in text and "Đối tượng" in text:
            print(f"  Found Training Plan Table at Elem {idx}")
            t_elem = make_placeholder_tbl()
            c_elem = make_caption_p("Hình 2.6. Minh chứng buổi chuyển giao công nghệ và đào tạo vận hành hệ thống.")
            elem.addnext(c_elem)
            elem.addnext(t_elem)
            tbl_obj = Table(t_elem, doc._body)
            format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG EVIDENCE: Ảnh biên bản bàn giao / Ảnh chụp buổi đào tạo vận hành hệ thống]")

    # ---------------------------------------------------------------------------
    # CHƯƠNG III: Tách System Context & Use Case; Chuẩn hóa 13 FR captions; Thêm 4 NFR evidence
    # ---------------------------------------------------------------------------
    print("Processing Chapter III...")
    for idx, elem, tag, text in elem_info:
        # Tách System Context và Use Case
        if tag == "p" and "Sơ đồ ngữ cảnh và Use Case tổng quát hệ thống" in text:
            print(f"  Found System Context & Use Case at Elem {idx}")
            # elem hiện tại sẽ là caption Hình 3.1
            set_p_text(elem, "Hình 3.1. Sơ đồ ngữ cảnh hệ thống (System Context Diagram).")
            # Trước elem là bảng placeholder cũ, sửa text placeholder cho Hình 3.1
            prev_elem = elem.getprevious()
            if prev_elem is not None and prev_elem.tag.endswith("tbl"):
                tbl_obj = Table(prev_elem, doc._body)
                format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG DIAGRAM: Sơ đồ ngữ cảnh hệ thống (System Context Diagram) phân định ranh giới giữa YiYi Book và các External Services: Payment Gateway, Email Service, AI LLM]")
            
            # Thêm Sơ đồ Use Case tổng quát (Hình 3.2) ngay sau Hình 3.1
            t_elem = make_placeholder_tbl()
            c_elem = make_caption_p("Hình 3.2. Sơ đồ Use Case tổng quát hệ thống (Overall Use Case Diagram).")
            elem.addnext(c_elem)
            elem.addnext(t_elem)
            tbl_obj = Table(t_elem, doc._body)
            format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG DIAGRAM: Sơ đồ Use Case tổng quát (Overall Use Case Diagram) gồm 3 Actors chính: Khách vãng lai, Khách hàng thành viên, Quản trị viên]")

        # 3.2 Chuẩn hóa Caption FR-xx thành Hình 3.3 -> Hình 3.15
        if tag == "p" and text.startswith("Hình FR-"):
            fr_num = text.split(".")[0].replace("Hình ", "").strip()
            fr_map = {
                "FR-01": "Hình 3.3. Giao diện chức năng Đăng ký, Đăng nhập và Xác thực (Authentication).",
                "FR-04": "Hình 3.4. Giao diện Trang chủ và Khám phá Danh mục sách (Home & Catalogue).",
                "FR-05": "Hình 3.5. Giao diện Tìm kiếm và Bộ lọc nâng cao (Search & Filter).",
                "FR-07": "Hình 3.6. Giao diện Chi tiết Sách và Đánh giá (Product Detail & Review).",
                "FR-09": "Hình 3.7. Giao diện Giỏ hàng và Quản lý Địa chỉ nhận hàng (Cart & Address).",
                "FR-11": "Hình 3.8. Giao diện Đặt hàng và Cổng thanh toán trực tuyến (Checkout & Payment).",
                "FR-13": "Hình 3.9. Giao diện Theo dõi Đơn hàng và Hậu mãi (Order Tracking & After-sales).",
                "FR-21": "Hình 3.10. Giao diện Tương tác Trợ lý ảo AI (YiYi AI Assistant).",
                "FR-22": "Hình 3.11. Giao diện Bảng điều khiển Quản trị (Admin Dashboard & Operations).",
                "FR-23": "Hình 3.12. Giao diện Quản trị Sách, Danh mục và Banner (Admin Catalogue & Banners).",
                "FR-24": "Hình 3.13. Giao diện Quản lý Đơn hàng và Vận chuyển (Admin Orders).",
                "FR-25": "Hình 3.14. Giao diện Quản trị Người dùng và Phân quyền (Admin Users & RBAC).",
                "FR-26": "Hình 3.15. Giao diện Quản trị Khuyến mãi, Đánh giá và Cài đặt (Admin Promotions & Settings)."
            }
            if fr_num in fr_map:
                set_p_text(elem, fr_map[fr_num])
                
        # 3.3 Thêm subsection 4.x Minh chứng yêu cầu phi chức năng sau Bảng NFR
        if tag == "tbl" and "ID" in text and "Thuộc tính" in text and "NFR-01" in text:
            print(f"  Found NFR Table at Elem {idx}")
            h_elem = make_heading_p("4.3 Minh chứng yêu cầu phi chức năng (Non-Functional Requirements Evidence)", level=2)
            
            nfr_items = [
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh đo thời gian phản hồi API và UI (Response Time < 500ms)]", "Hình 3.16. Kết quả đo thời gian phản hồi của hệ thống."),
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh kiểm thử kiểm soát quyền truy cập RBAC / HTTP 401 Unauthorized / HTTP 403 Forbidden]", "Hình 3.17. Minh chứng kiểm soát truy cập trái phép."),
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh giao diện đáp ứng đa thiết bị (Responsive UI) trên Desktop, Tablet và Mobile]", "Hình 3.18. Minh chứng tính tương thích và Responsive UI."),
                ("[CẦN BỔ SUNG EVIDENCE: Ảnh kiểm thử tính toàn vẹn dữ liệu và Transaction Rollback khi đặt hàng thất bại]", "Hình 3.19. Minh chứng đảm bảo tính toàn vẹn dữ liệu trong xử lý Order.")
            ]
            
            elem.addnext(h_elem)
            curr = h_elem
            for p_text, cap_text in nfr_items:
                t_elem = make_placeholder_tbl()
                c_elem = make_caption_p(cap_text)
                curr.addnext(c_elem)
                curr.addnext(t_elem)
                tbl_obj = Table(t_elem, doc._body)
                format_placeholder_table(tbl_obj, p_text)
                curr = c_elem

    # ---------------------------------------------------------------------------
    # CHƯƠNG IV: Thiết kế hệ thống (Kiến trúc, Package, ERD, Class & Sequence Diagrams)
    # ---------------------------------------------------------------------------
    print("Processing Chapter IV...")
    
    # Reload elements sau các thao tác chèn
    body_elements = list(doc._body._element)
    elem_info = []
    for idx, elem in enumerate(body_elements):
        tag = elem.tag.split('}')[-1]
        text = ""
        if tag == "p":
            text = "".join(elem.itertext()).strip()
        elif tag == "tbl":
            text = " ".join("".join(elem.itertext()).split()[:20]).strip()
        elem_info.append((idx, elem, tag, text))

    for idx, elem, tag, text in elem_info:
        # Chuẩn hóa caption Kiến trúc tổng thể & Kiến trúc Frontend
        if tag == "p" and "Kiến trúc tổng thể hệ thống YiYi Book" in text:
            set_p_text(elem, "Hình 4.1. Kiến trúc tổng thể hệ thống YiYi Book.")
            
        if tag == "p" and "Kiến trúc luồng Component và Routing Frontend" in text:
            set_p_text(elem, "Hình 4.2. Kiến trúc luồng Component và Routing Frontend.")
            
            # Thêm Backend Layered Arch, AI RAG Arch, Package Diagram, ERD sau Hình 4.2
            arch_items = [
                ("[CẦN BỔ SUNG DIAGRAM: Sơ đồ kiến trúc phân tầng Backend (Controller - Service - Repository - Entity)]", "Hình 4.3. Kiến trúc phân tầng Backend Spring Boot."),
                ("[CẦN BỔ SUNG DIAGRAM: Sơ đồ luồng xử lý AI Assistant & RAG Pipeline (Embedding, Vector Search, LLM Context)]", "Hình 4.4. Kiến trúc luồng tích hợp Trợ lý AI và xử lý ngôn ngữ tự nhiên."),
                ("[CẦN BỔ SUNG DIAGRAM: Sơ đồ phân rã gói phần mềm hệ thống (Package Diagram)]", "Hình 4.5. Sơ đồ phân rã gói phần mềm hệ thống (Package Diagram)."),
                ("[CẦN BỔ SUNG DIAGRAM: Sơ đồ thực thể quan hệ cơ sở dữ liệu chi tiết (Entity Relationship Diagram - ERD) 17 bảng]", "Hình 4.6. Sơ đồ thực thể quan hệ cơ sở dữ liệu (ERD).")
            ]
            curr = elem
            for p_text, cap_text in arch_items:
                t_elem = make_placeholder_tbl()
                c_elem = make_caption_p(cap_text)
                curr.addnext(c_elem)
                curr.addnext(t_elem)
                tbl_obj = Table(t_elem, doc._body)
                format_placeholder_table(tbl_obj, p_text)
                curr = c_elem

        # Sửa các caption cũ bị sai đánh số từ Chương 3 trong Chương 4
        # 4.1 Đăng ký, Đăng nhập
        if tag == "p" and "Hình 3.1. Sơ đồ lớp chức năng Đăng ký và Đăng nhập" in text:
            set_p_text(elem, "Hình 4.7. Sơ đồ lớp (Class Diagram) chức năng Xác thực và Phân quyền.")
        if tag == "p" and "Hình 3.2. Sơ đồ tuần tự chức năng Đăng nhập" in text:
            set_p_text(elem, "Hình 4.8. Sơ đồ tuần tự (Sequence Diagram) luồng Đăng nhập và Cấp phát JWT.")

        # 4.2 Khám phá sách (Sản phẩm)
        if tag == "p" and "Hình 3.3. Sơ đồ tuần tự luồng Danh sách và Bộ lọc Sách" in text:
            set_p_text(elem, "Hình 4.9. Sơ đồ tuần tự (Sequence Diagram) luồng Xem danh mục và Lọc sách.")
        if tag == "p" and "Hình 3.4. Sơ đồ tuần tự luồng Tìm kiếm" in text:
            set_p_text(elem, "Hình 4.10. Sơ đồ tuần tự (Sequence Diagram) luồng Tìm kiếm sách nâng cao.")
        if tag == "p" and "Hình 3.5. Sơ đồ lớp chức năng Chi tiết sách và Đánh giá" in text:
            set_p_text(elem, "Hình 4.11. Sơ đồ lớp (Class Diagram) chức năng Chi tiết Sách và Đánh giá.")
        if tag == "p" and "Hình 3.6. Sơ đồ tuần tự luồng Đánh giá sản phẩm" in text:
            set_p_text(elem, "Hình 4.12. Sơ đồ tuần tự (Sequence Diagram) luồng Gửi đánh giá và Nhận điểm thưởng.")

        # 4.3 Giỏ hàng & Đặt hàng
        if tag == "p" and "Hình 3.7. Sơ đồ tuần tự luồng Thêm vào giỏ" in text:
            set_p_text(elem, "Hình 4.13. Sơ đồ tuần tự (Sequence Diagram) luồng Thao tác Giỏ hàng (Thêm/Sửa/Xóa).")
        if tag == "p" and "Hình 3.8. Sơ đồ lớp chức năng Giỏ hàng và Thanh toán" in text:
            set_p_text(elem, "Hình 4.14. Sơ đồ lớp (Class Diagram) chức năng Giỏ hàng, Đơn hàng và Thanh toán.")
        if tag == "p" and "Hình 3.9. Sơ đồ tuần tự luồng Đặt hàng (Checkout)" in text:
            set_p_text(elem, "Hình 4.15. Sơ đồ tuần tự (Sequence Diagram) luồng Đặt hàng COD và Áp dụng Khuyến mãi.")
        if tag == "p" and "Hình 3.10. Sơ đồ lớp tích hợp Cổng thanh toán trực tuyến" in text:
            set_p_text(elem, "Hình 4.16. Sơ đồ lớp (Class Diagram) tích hợp Cổng thanh toán trực tuyến.")
        if tag == "p" and "Hình 3.11. Sơ đồ tuần tự luồng Thanh toán Online" in text:
            set_p_text(elem, "Hình 4.17. Sơ đồ tuần tự (Sequence Diagram) luồng Thanh toán trực tuyến (VNPAY/MoMo Sandbox).")

        # 4.4 Trạng thái đơn hàng, Y-Points & AI & Admin
        if tag == "p" and "Hình 3.12. Sơ đồ trạng thái vòng đời Đơn hàng" in text:
            set_p_text(elem, "Hình 4.18. Sơ đồ chuyển trạng thái (State Machine Diagram) vòng đời Đơn hàng.")
        if tag == "p" and "Hình 3.13. Sơ đồ tuần tự luồng Yêu cầu trả hàng" in text:
            set_p_text(elem, "Hình 4.19. Sơ đồ tuần tự (Sequence Diagram) luồng Yêu cầu Trả hàng / Hoàn tiền.")
        if tag == "p" and "Hình 3.14. Sơ đồ tuần tự luồng Tích lũy và Tiêu điểm Y-Point" in text:
            set_p_text(elem, "Hình 4.20. Sơ đồ tuần tự (Sequence Diagram) luồng Tích lũy và Sử dụng điểm thưởng Y-Point.")
        if tag == "p" and "Hình 3.15. Sơ đồ thành phần mô-đun Trợ lý ảo AI" in text:
            set_p_text(elem, "Hình 4.21. Sơ đồ thành phần (Component Diagram) mô-đun Trợ lý AI và Vector Store.")
        if tag == "p" and "Hình 3.16. Sơ đồ tuần tự luồng Tư vấn sách với AI" in text:
            set_p_text(elem, "Hình 4.22. Sơ đồ tuần tự (Sequence Diagram) luồng Tư vấn sách thông minh qua AI Assistant.")
        if tag == "p" and "Hình 3.17. Sơ đồ tuần tự luồng Quản trị Danh mục và Sách" in text:
            set_p_text(elem, "Hình 4.23. Sơ đồ tuần tự (Sequence Diagram) luồng Quản trị Danh mục và Sản phẩm (Admin).")
        if tag == "p" and "Hình 3.18. Sơ đồ tuần tự luồng Cập nhật trạng thái đơn hàng" in text:
            set_p_text(elem, "Hình 4.24. Sơ đồ tuần tự (Sequence Diagram) luồng Xử lý và Cập nhật trạng thái Đơn hàng (Admin).")
        if tag == "p" and "Hình 3.19. Sơ đồ tuần tự luồng Quản lý Banner và Khuyến mãi" in text:
            set_p_text(elem, "Hình 4.25. Sơ đồ tuần tự (Sequence Diagram) luồng Quản trị Khuyến mãi và Banner (Admin).")

    # ---------------------------------------------------------------------------
    # CHƯƠNG V: Kiểm thử hệ thống (Đã có 12 test evidence -> bổ sung Jira Defect Hình 5.13)
    # ---------------------------------------------------------------------------
    print("Processing Chapter V...")
    body_elements = list(doc._body._element)
    elem_info = []
    for idx, elem in enumerate(body_elements):
        tag = elem.tag.split('}')[-1]
        text = ""
        if tag == "p":
            text = "".join(elem.itertext()).strip()
        elif tag == "tbl":
            text = " ".join("".join(elem.itertext()).split()[:20]).strip()
        elem_info.append((idx, elem, tag, text))

    for idx, elem, tag, text in elem_info:
        # Kiểm tra bảng Defect Tracking (Bảng 5.8: Báo cáo Bug & Defect Tracking)
        if tag == "tbl" and "BUG-01" in text and "BUG-05" in text:
            print(f"  Found Bug Tracking Table at Elem {idx}")
            t_elem = make_placeholder_tbl()
            c_elem = make_caption_p("Hình 5.13. Minh chứng quản lý và đóng lỗi (Defect Status: Closed/Verified) trên Jira.")
            elem.addnext(c_elem)
            elem.addnext(t_elem)
            tbl_obj = Table(t_elem, doc._body)
            format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp Jira Bug Board / Danh sách Defect trạng thái Closed/Verified]")

    # ---------------------------------------------------------------------------
    # CHƯƠNG VI: Triển khai & Hướng dẫn sử dụng
    # ---------------------------------------------------------------------------
    print("Processing Chapter VI...")
    body_elements = list(doc._body._element)
    elem_info = []
    for idx, elem in enumerate(body_elements):
        tag = elem.tag.split('}')[-1]
        text = ""
        if tag == "p":
            text = "".join(elem.itertext()).strip()
        elif tag == "tbl":
            text = " ".join("".join(elem.itertext()).split()[:20]).strip()
        elem_info.append((idx, elem, tag, text))

    for idx, elem, tag, text in elem_info:
        # Bổ sung 8 evidence cài đặt & triển khai theo 8 bước
        # Bước 1: Chuẩn bị cấu hình môi trường
        if tag == "p" and "Bước 1: Chuẩn bị cấu hình môi trường" in text:
            t_elem = make_placeholder_tbl()
            c_elem = make_caption_p("Hình 6.1. Minh chứng cấu hình biến môi trường hệ thống (.env và application.properties).")
            elem.addnext(c_elem)
            elem.addnext(t_elem)
            tbl_obj = Table(t_elem, doc._body)
            format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG EVIDENCE: Ảnh cấu hình file .env và application-local.properties]")
            
        # Bước 2: Khởi chạy cơ sở dữ liệu
        if tag == "p" and "Bước 2: Khởi chạy cơ sở dữ liệu" in text:
            t_elem = make_placeholder_tbl()
            c_elem = make_caption_p("Hình 6.2. Minh chứng khởi chạy PostgreSQL thành công qua Docker Container.")
            elem.addnext(c_elem)
            elem.addnext(t_elem)
            tbl_obj = Table(t_elem, doc._body)
            format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG EVIDENCE: Ảnh terminal lệnh docker compose up -d và Docker Desktop]")

        # Bước 3: Khởi động Backend
        if tag == "p" and "Bước 3: Khởi động Backend Service" in text:
            t_elem = make_placeholder_tbl()
            c_elem = make_caption_p("Hình 6.3. Minh chứng Backend Spring Boot khởi động thành công trên cổng 8080.")
            elem.addnext(c_elem)
            elem.addnext(t_elem)
            tbl_obj = Table(t_elem, doc._body)
            format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG EVIDENCE: Ảnh console log Spring Boot Started Application in X seconds]")

        # Bước 4: Kiểm tra Backend
        if tag == "p" and "Bước 4: Kiểm tra Backend khả dụng" in text:
            t_elem = make_placeholder_tbl()
            c_elem = make_caption_p("Hình 6.4. Minh chứng kiểm tra API Health Check trả về kết quả UP.")
            elem.addnext(c_elem)
            elem.addnext(t_elem)
            tbl_obj = Table(t_elem, doc._body)
            format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG EVIDENCE: Ảnh Postman / Browser kiểm tra Swagger UI hoặc API /actuator/health]")

        # Bước 5: Khởi động Frontend
        if tag == "p" and "Bước 5: Khởi động Frontend Service" in text:
            t_elem = make_placeholder_tbl()
            c_elem = make_caption_p("Hình 6.5. Minh chứng Frontend Vite React khởi động thành công trên cổng 5173.")
            elem.addnext(c_elem)
            elem.addnext(t_elem)
            tbl_obj = Table(t_elem, doc._body)
            format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG EVIDENCE: Ảnh terminal lệnh npm run dev sẵn sàng truy cập Localhost]")

        # Bước 6: Đăng nhập tài khoản mẫu
        if tag == "p" and "Bước 6: Đăng nhập tài khoản mẫu" in text:
            t_elem = make_placeholder_tbl()
            c_elem = make_caption_p("Hình 6.6. Minh chứng đăng nhập thành công vào hệ thống cục bộ (Localhost:5173).")
            elem.addnext(c_elem)
            elem.addnext(t_elem)
            tbl_obj = Table(t_elem, doc._body)
            format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình trang chủ với trạng thái đã đăng nhập tài khoản]")

        # Bước 7: Kiểm tra triển khai Frontend
        if tag == "p" and "Bước 7: Kiểm tra triển khai Frontend trên môi trường Production" in text:
            t_elem = make_placeholder_tbl()
            c_elem = make_caption_p("Hình 6.7. Minh chứng giao diện Frontend triển khai trên Production.")
            elem.addnext(c_elem)
            elem.addnext(t_elem)
            tbl_obj = Table(t_elem, doc._body)
            format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG EVIDENCE: Ảnh chụp màn hình ứng dụng trên Domain Production]")

        # Bước 8: Kiểm tra triển khai Backend
        if tag == "p" and "Bước 8: Kiểm tra triển khai Backend trên môi trường Production" in text:
            t_elem = make_placeholder_tbl()
            c_elem = make_caption_p("Hình 6.8. Minh chứng Backend và Cơ sở dữ liệu Cloud hoạt động trên Production.")
            elem.addnext(c_elem)
            elem.addnext(t_elem)
            tbl_obj = Table(t_elem, doc._body)
            format_placeholder_table(tbl_obj, "[CẦN BỔ SUNG EVIDENCE: Ảnh Dashboard Cloud / Server Log / Swagger Production UP]")

        # 6.2 Chuẩn hóa Caption Hướng dẫn sử dụng (Hình HD-xx -> Hình 6.9 -> Hình 6.23)
        if tag == "p" and text.startswith("Hình HD-"):
            hd_num = text.split(".")[0].replace("Hình ", "").strip()
            hd_map = {
                "HD-01": "Hình 6.9. Hướng dẫn Đăng ký tài khoản người dùng mới.",
                "HD-02": "Hình 6.10. Hướng dẫn Đăng nhập hệ thống và cấp phiên làm việc.",
                "HD-03": "Hình 6.11. Hướng dẫn Khám phá Trang chủ và Xem các phân loại sách.",
                "HD-04": "Hình 6.12. Hướng dẫn Tìm kiếm và Lọc sách theo tiêu chí nâng cao.",
                "HD-05": "Hình 6.13. Hướng dẫn Xem thông tin Chi tiết Sách và Đọc đánh giá.",
                "HD-06": "Hình 6.14. Hướng dẫn Thao tác Quản lý Giỏ hàng.",
                "HD-07": "Hình 6.15. Hướng dẫn Nhập thông tin Nhận hàng và Áp dụng Mã giảm giá.",
                "HD-08": "Hình 6.16. Hướng dẫn Thực hiện Thanh toán Trực tuyến qua Cổng thanh toán.",
                "HD-09": "Hình 6.17. Hướng dẫn Theo dõi Chi tiết Đơn hàng và Lịch sử mua hàng.",
                "HD-10": "Hình 6.18. Hướng dẫn Gửi Yêu cầu Đổi trả / Hoàn tiền đơn hàng.",
                "HD-11": "Hình 6.19. Hướng dẫn Đánh giá sản phẩm và Tích lũy điểm Y-Points.",
                "HD-21": "Hình 6.20. Hướng dẫn Tương tác và Nhận tư vấn từ Trợ lý ảo AI.",
                "HD-22": "Hình 6.21. Hướng dẫn Quản trị Danh mục và Sản phẩm trên Admin Portal.",
                "HD-24": "Hình 6.22. Hướng dẫn Xử lý và Chuyển trạng thái Đơn hàng trên Admin Portal.",
                "HD-26": "Hình 6.23. Hướng dẫn Quản lý Chương trình Khuyến mãi và Cài đặt hệ thống."
            }
            if hd_num in hd_map:
                set_p_text(elem, hd_map[hd_num])

    # ---------------------------------------------------------------------------
    # STEP 3: Lưu ra file kết quả
    # ---------------------------------------------------------------------------
    tmp_output = output_path + ".tmp.docx"
    print(f"Saving refined document to temporary path: {tmp_output}")
    doc.save(tmp_output)
    
    print(f"Copying to final destination: {output_path}")
    shutil.copy2(tmp_output, output_path)
    if os.path.exists(tmp_output):
        os.remove(tmp_output)
        
    print("SUCCESS: Refined all chapters placeholders successfully!")

if __name__ == "__main__":
    src = "YiYi_Book_Capstone_Project_Report_Tieng_Viet (1).docx"
    dst = "YiYi_Book_Capstone_Project_Report_Tieng_Viet_Final_Refined.docx"
    process_full_document(src, dst)