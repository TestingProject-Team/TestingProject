# -*- coding: utf-8 -*-
import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document("YiYi_Book_Capstone_Project_Report_Tieng_Viet (1)_BACKUP.docx")

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

for idx, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text.startswith("CHƯƠNG") or text.startswith("Chương") or text.startswith("MỤC LỤC") or text.startswith("V. ") or text.startswith("IV. ") or text.startswith("VI. "):
        print(f"P[{idx}]: (style: {p.style.name}) {text[:100]}")
    elif p.style.name.startswith("Heading"):
        print(f"P[{idx}] [{p.style.name}]: {text[:80]}")
