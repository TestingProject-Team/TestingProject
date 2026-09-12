import docx
import os
import sys

# Set encoding for stdout
sys.stdout.reconfigure(encoding='utf-8')

docx_path = "YiYi_Book_Capstone_Project_Report_Tieng_Viet (1).docx"
doc = docx.Document(docx_path)

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

with open("docs_headings_summary.txt", "w", encoding="utf-8") as f:
    for idx, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt:
            if p.style.name.startswith("Heading") or any(txt.startswith(x) for x in ["I.", "II.", "III.", "IV.", "V.", "VI.", "VII.", "Chương", "CHƯƠNG", "1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.", "10."]):
                f.write(f"[{idx}] ({p.style.name}) {txt}\n")

print("Exported docs_headings_summary.txt successfully.")
