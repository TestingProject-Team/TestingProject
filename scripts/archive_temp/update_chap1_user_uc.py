import docx

doc_path = r'E:\TestingProject\YiYi_Book_Capstone_Project_Report_Tieng_Viet_Refined_Full(1).docx'
doc = docx.Document(doc_path)

# 1. Update Target Users in Table 3 (Row 4)
t3 = doc.tables[3]
# Row 4 is 'Người dùng mục tiêu'
t3.rows[4].cells[1].text = "Khách Vãng Lai (Guest), Khách Hàng Thành Viên (Member), Quản Trị Viên (Administrator)"

# 2. Update UC-03 in Table 28 (Row 3)
t28 = doc.tables[28]
# Check row 3 (UC-03)
for row in t28.rows:
    if row.cells[0].text.strip() == 'UC-03':
        row.cells[1].text = "Xem chi tiết sách & Xem đánh giá"
        row.cells[3].text = "Xem thông tin xuất bản, tồn kho, tóm tắt nội dung và các bài đánh giá từ cộng đồng"

# Also check any paragraph mentioning UC-03
for p in doc.paragraphs:
    if 'UC-03' in p.text and 'Xem chi tiết sách & Đánh giá' in p.text:
        p.text = p.text.replace('Xem chi tiết sách & Đánh giá', 'Xem chi tiết sách & Xem đánh giá')

doc.save(doc_path)
print("Updated Chapter I Target User & UC-03 successfully.")
